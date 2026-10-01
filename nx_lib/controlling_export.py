"""Sydoc Controlling: the month as one .xlsx (#433).

Pure and Flask-free like nx_lib/finance_export.py: it takes the month payload
of nx_lib/controlling.py (month_payload, plus the view's label/state) and the
translated labels, and returns the workbook bytes. Four sheets, the shape of
the old Projektcontrolling workbook:

* Overview -- the margin table: one row per stream, a total row;
* Detail -- per stream its hours by task x rate, its invoices with their
  lines, the difference and the per-document KPIs;
* Hours by task -- task x stream, with totals and the FTE line;
* Volumes -- documents per stream against the previous month.

A figure the page cannot compute is written as its state ("no invoice",
"draft only", ...), never as 0.
"""

from __future__ import annotations

import io
from typing import Any

from .reporting.export import _safe_cell

MONEY = "#,##0.00"
HOURS = "#,##0.00"
PCT = "0.0%"
COUNT = "#,##0"


def _or_state(value, cell, labels):
    if value is not None:
        return value
    return labels["states"].get(cell.get("state"), "")


def workbook(payload: dict, labels: dict[str, Any], month_label: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEF4")
    wb = Workbook()

    def title(ws, text):
        ws.append([_safe_cell(labels["title"])])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([_safe_cell(text)])
        ws["A2"].font = Font(bold=True, size=12)
        ws.append([_safe_cell(month_label)])
        ws.append([])

    def header(ws, values):
        ws.append([_safe_cell(v) for v in values])
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = head_fill

    def fmt(ws, row, formats):
        for col, f in formats.items():
            ws.cell(row=row, column=col).number_format = f

    def widths(ws, first=34, rest=16):
        ws.column_dimensions["A"].width = first
        for i in range(2, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = rest

    def flags(cell):
        out = [labels[f] for f in cell.get("flags") or [] if f in labels]
        return ", ".join(out)

    # ---- Overview -------------------------------------------------------
    ws = wb.active
    ws.title = labels["overview"][:31]
    title(ws, labels["overview"])
    header(
        ws,
        [
            labels["stream"],
            labels["hours"],
            labels["cost"],
            labels["invoiced"],
            labels["invoiced_incl"],
            labels["margin"],
            labels["margin_pct"],
            labels["delta_margin"],
            labels["documents"],
            labels["state"],
        ],
    )
    money_cols = {2: HOURS, 3: MONEY, 4: MONEY, 5: MONEY, 6: MONEY, 7: PCT, 8: MONEY, 9: COUNT}
    for s in payload["streams"]:
        c = s["cur"]
        ws.append(
            [
                _safe_cell(s["label"]),
                c["hours"],
                c["cost"],
                _or_state(c["invoiced"], c, labels),
                c["invoicedIncl"],
                c["margin"],
                c["marginPct"],
                (s.get("delta") or {}).get("margin"),
                c["documents"],
                _safe_cell(
                    " · ".join(x for x in (labels["states"].get(c["state"], ""), flags(c)) if x)
                ),
            ]
        )
        fmt(ws, ws.max_row, money_cols)
    t = payload["totals"]["cur"]
    ws.append(
        [
            labels["total"],
            t["hours"],
            t["cost"],
            t["invoiced"],
            t["invoicedIncl"],
            t["margin"],
            t["marginPct"],
            None,
            t["documents"],
            f"{t['streamsWithMargin']}/{t['streams']}",
        ]
    )
    for c in ws[ws.max_row]:
        c.font = bold
    fmt(ws, ws.max_row, money_cols)
    if payload.get("unassigned"):
        ws.append([])
        ws.append([labels["unassigned"]])
        ws.cell(row=ws.max_row, column=1).font = bold
        for inv in payload["unassigned"]:
            ws.append([_safe_cell(f"{inv['nr']} · {inv['title']}"), inv["date"], None, inv["excl"]])
            fmt(ws, ws.max_row, {4: MONEY})
    if payload.get("unmapped"):
        ws.append([])
        ws.append([labels["unmapped"]])
        ws.cell(row=ws.max_row, column=1).font = bold
        for u in payload["unmapped"]:
            ws.append([_safe_cell(f"{u['customer']} · {u['package']} · {u['task']}"), u["hours"]])
            fmt(ws, ws.max_row, {2: HOURS})
    widths(ws)

    # ---- Detail ---------------------------------------------------------
    ws = wb.create_sheet(labels["detail"][:31])
    title(ws, labels["detail"])
    for s in payload["streams"]:
        c, p = s["cur"], s["prev"] or {}
        ws.append([_safe_cell(s["label"])])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        header(
            ws,
            [labels["task"], labels["hours"], labels["rate"], labels["cost"], labels["previous"]],
        )
        prev_tasks = p.get("tasks") or {}
        for task, hours in c["tasks"].items():
            cost = round(hours * c["rate"], 2) if c["rate"] is not None else None
            ws.append([_safe_cell(task), hours, c["rate"], cost, prev_tasks.get(task)])
            fmt(ws, ws.max_row, {2: HOURS, 3: MONEY, 4: MONEY, 5: HOURS})
        for ext in c.get("externalCosts") or []:
            ws.append(
                [_safe_cell(f"{labels['external']}: {ext['label']}"), None, None, ext["amount"]]
            )
            fmt(ws, ws.max_row, {4: MONEY})
        ws.append([labels["total"], c["hours"], c["rate"], c["cost"], p.get("hours")])
        for cell in ws[ws.max_row]:
            cell.font = bold
        fmt(ws, ws.max_row, {2: HOURS, 3: MONEY, 4: MONEY, 5: HOURS})
        ws.append([])
        header(
            ws,
            [
                labels["invoice"],
                labels["date"],
                labels["quantity"],
                labels["unit_price"],
                labels["amount"],
            ],
        )
        for inv in c["invoices"]:
            ws.append(
                [
                    _safe_cell(
                        f"{inv['nr']} · {inv['title']} ({labels['states'].get(inv['status'], inv['status'])})"
                    ),
                    inv["date"],
                    None,
                    None,
                    inv["excl"],
                ]
            )
            ws.cell(row=ws.max_row, column=1).font = bold
            fmt(ws, ws.max_row, {5: MONEY})
            for line in inv.get("positions") or []:
                if line.get("kind") != "line":
                    continue
                ws.append(
                    [
                        _safe_cell("  " + (line.get("text") or "").replace("\n", " · ")),
                        None,
                        line.get("amount"),
                        line.get("unitPrice"),
                        line.get("total"),
                    ]
                )
                fmt(ws, ws.max_row, {3: HOURS, 4: MONEY, 5: MONEY})
        ws.append([labels["invoiced"], None, None, None, _or_state(c["invoiced"], c, labels)])
        fmt(ws, ws.max_row, {5: MONEY})
        ws.append([labels["invoiced_incl"], None, None, None, c["invoicedIncl"]])
        fmt(ws, ws.max_row, {5: MONEY})
        ws.append([labels["margin"], None, None, None, c["margin"], c["marginPct"]])
        for cell in ws[ws.max_row]:
            cell.font = bold
        fmt(ws, ws.max_row, {5: MONEY, 6: PCT})
        k = c["kpis"]
        ws.append([labels["documents"], None, None, None, c["documents"]])
        ws.append([labels["seconds_per_doc"], None, None, None, k["secondsPerDocument"]])
        ws.append([labels["docs_per_hour"], None, None, None, k["documentsPerHour"]])
        ws.append([labels["chf_per_doc"], None, None, None, k["chfPerDocument"]])
        if c.get("incomplete"):
            ws.append([_safe_cell(f"{labels['incomplete']}: {c['incomplete']}")])
        ws.append([])
    widths(ws, first=60)

    # ---- Hours by task --------------------------------------------------
    ws = wb.create_sheet(labels["tasks"][:31])
    title(ws, labels["tasks"])
    m = payload["tasks"]
    names = {s["key"]: s["label"] for s in payload["streams"]}
    header(ws, [labels["task"], *[names.get(k, k) for k in m["streams"]], labels["total"]])
    for task, row, total in zip(m["tasks"], m["cells"], m["taskTotals"], strict=True):
        ws.append([_safe_cell(task), *row, total])
        fmt(ws, ws.max_row, {i: HOURS for i in range(2, len(row) + 3)})
    ws.append([labels["total"], *m["streamTotals"], m["total"]])
    for cell in ws[ws.max_row]:
        cell.font = bold
    fmt(ws, ws.max_row, {i: HOURS for i in range(2, len(m["streamTotals"]) + 3)})
    if t.get("fte") is not None:
        ws.append([labels["fte"], t["fte"]])
    widths(ws, rest=14)

    # ---- Volumes --------------------------------------------------------
    ws = wb.create_sheet(labels["volumes"][:31])
    title(ws, labels["volumes"])
    header(ws, [labels["stream"], labels["documents"], labels["previous"]])
    for s in payload["streams"]:
        ws.append(
            [_safe_cell(s["label"]), s["cur"]["documents"], (s["prev"] or {}).get("documents")]
        )
        fmt(ws, ws.max_row, {2: COUNT, 3: COUNT})
    prev_t = (payload["totals"].get("prev") or {}).get("documents")
    ws.append([labels["total"], t["documents"], prev_t])
    for cell in ws[ws.max_row]:
        cell.font = bold
    fmt(ws, ws.max_row, {2: COUNT, 3: COUNT})
    widths(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

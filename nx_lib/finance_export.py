"""Sydoc Finance: the billable BPS hours as Excel and PDF, one sheet per invoice (#408).

Pure and Flask-free like nx_lib/finance.py: it takes the BPS section payload
(live, or the snapshot of a closed month) and the section's Bookings spec,
and returns file bytes. The view supplies the translated labels.

A sheet is what accounting bills on one invoice: one per customer, except the
customers in ``Bookings.split`` (Privera), whose project packages are billed
separately -- "Tagesgeschäft Neuzugänge" and "Neuzugänge" are both the
Neuzugänge invoice. The billed hours are recomputed from the booked ones
(``bps.billed_hours``), so a snapshot taken before the rounding existed
exports the same way as a live month.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from . import bps
from .reporting.export import _safe_cell

#: The package prefix of Privera's day-to-day streams: "Tagesgeschäft Invoice"
#: is billed with "Invoice".
_DAILY_PREFIX = "tagesgeschäft "
#: Characters Excel refuses in a sheet name.
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
_SLUG_BAD = re.compile(r"[^a-z0-9]+")


@dataclass
class Row:
    date: dt.date | None
    package: str
    task: str
    person: str
    hours: Decimal
    billed: Decimal
    comment: str


@dataclass
class Sheet:
    """One invoice's bookings."""

    key: str
    title: str
    client: str
    rows: list[Row] = field(default_factory=list)

    @property
    def hours(self) -> Decimal:
        return sum((r.hours for r in self.rows), Decimal(0))

    @property
    def billed(self) -> Decimal:
        return sum((r.billed for r in self.rows), Decimal(0))


def slug(text: str) -> str:
    """A sheet's key in a URL and a file name: 'Privera Neuzugänge' -> 'privera-neuzugaenge'."""
    text = text.casefold()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("é", "e"), ("è", "e"), ("à", "a")):
        text = text.replace(a, b)
    return _SLUG_BAD.sub("-", text).strip("-") or "sheet"


def _package(package: str) -> str:
    """A split customer's package as its invoice: the daily-stream prefix dropped."""
    text = (package or "").strip()
    if text.casefold().startswith(_DAILY_PREFIX):
        text = text[len(_DAILY_PREFIX) :].strip()
    return text


def _decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    return Decimal(str(value))


def _date(value: Any) -> dt.date | None:
    if isinstance(value, dt.date):
        return value
    try:
        return dt.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def sheets(payload: dict, bookings: Any) -> list[Sheet]:
    """The BPS section payload as invoice sheets, in the order of `bookings.customers`.

    A customer the spec does not name (a snapshot from before the customer
    filter) goes after the named ones, under its own BPS name.
    """
    bk = (payload or {}).get("bookings") or {}
    idx = {c["field"]: i for i, c in enumerate(bk.get("columns") or [])}

    def cell(row: list, name: str) -> Any:
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    client_of = dict(bookings.customers)
    order = [c for c, _ in bookings.customers]
    split = {c.casefold(): [p.casefold() for p in packages] for c, packages in bookings.split}
    by_key: dict[str, Sheet] = {}
    rank: dict[str, tuple] = {}
    for group in bk.get("groups") or []:
        for r in group.get("rows") or []:
            customer = str(cell(r, bookings.group_by) or group.get("key") or "")
            client = client_of.get(customer, customer)
            package = str(cell(r, "Projektpaket") or "")
            title = client
            if customer.casefold() in split:
                title = " ".join(filter(None, (client, _package(package))))
            key = slug(title)
            if key not in by_key:
                by_key[key] = Sheet(key=key, title=title, client=client)
                pos = order.index(customer) if customer in order else len(order)
                packages = split.get(customer.casefold(), [])
                sub = _package(package).casefold()
                sub_pos = packages.index(sub) if sub in packages else len(packages)
                rank[key] = (pos, sub_pos, title.casefold())
            hours = _decimal(cell(r, bookings.hours))
            by_key[key].rows.append(
                Row(
                    date=_date(cell(r, "Datum")),
                    package=package,
                    task=str(cell(r, "Aufgabe") or ""),
                    person=str(cell(r, "Benutzer") or ""),
                    hours=hours,
                    billed=_decimal(bps.billed_hours(hours)),
                    comment=str(cell(r, "Beschreibung") or ""),
                )
            )
    out = sorted(by_key.values(), key=lambda s: rank[s.key])
    for s in out:
        s.rows.sort(key=lambda r: (r.date or dt.date.min, r.person.casefold()))
    return out


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def _sheet_name(title: str, taken: set[str]) -> str:
    """An Excel-safe, unique sheet name of at most 31 characters."""
    base = _SHEET_BAD.sub(" ", title).strip()[:31] or "Sheet"
    name, n = base, 2
    while name.casefold() in taken:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    taken.add(name.casefold())
    return name


def workbook(items: list[Sheet], labels: dict[str, Any], month_label: str) -> bytes:
    """One .xlsx: an overview sheet when there is more than one invoice, then one per invoice."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEF4")
    hours_fmt = "0.00"
    wb = Workbook()
    first = wb.active
    taken: set[str] = set()

    def title_rows(ws: Any, title: str) -> None:
        ws.append([labels["title"]])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([_safe_cell(title)])
        ws["A2"].font = Font(bold=True, size=12)
        ws.append([month_label])
        ws.append([])

    def header(ws: Any, values: list[str]) -> None:
        ws.append(values)
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = head_fill

    if len(items) > 1:
        ws = first
        ws.title = _sheet_name(labels["overview"], taken)
        title_rows(ws, labels["overview"])
        header(ws, [labels["invoice"], labels["bookings"], labels["billed"]])
        start = ws.max_row + 1
        for s in items:
            ws.append([_safe_cell(s.title), len(s.rows), float(s.billed)])
        end = ws.max_row
        ws.append([labels["total"], f"=SUM(B{start}:B{end})", f"=SUM(C{start}:C{end})"])
        for c in ws[ws.max_row]:
            c.font = bold
        for row in ws.iter_rows(min_row=start, max_row=ws.max_row, min_col=3, max_col=3):
            for c in row:
                c.number_format = hours_fmt
        for col, width in zip("ABC", (34, 12, 16), strict=True):
            ws.column_dimensions[col].width = width
        first = None

    # Only the billed (rounded) hours: what the invoice charges (#408).
    columns = [
        (labels["date"], 12),
        (labels["package"], 24),
        (labels["task"], 24),
        (labels["person"], 22),
        (labels["comment"], 64),
        (labels["billed"], 16),
    ]
    for s in items:
        if first is not None:
            ws, first = first, None
            ws.title = _sheet_name(s.title, taken)
        else:
            ws = wb.create_sheet(_sheet_name(s.title, taken))
        title_rows(ws, s.title)
        header(ws, [c for c, _ in columns])
        head_row = ws.max_row
        for r in s.rows:
            ws.append(
                [r.date]
                + [_safe_cell(v) for v in (r.package, r.task, r.person, r.comment)]
                + [float(r.billed)]
            )
            ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
            ws.cell(ws.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(ws.max_row, 6).number_format = hours_fmt
        end = ws.max_row
        total = f"=SUM(F{head_row + 1}:F{end})" if s.rows else 0
        ws.append([labels["total"], None, None, None, None, total])
        for c in ws[ws.max_row]:
            c.font = bold
        ws.cell(ws.max_row, 6).number_format = hours_fmt
        for i, (_, width) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = ws.cell(head_row + 1, 1)
        if s.rows:
            ws.auto_filter.ref = f"A{head_row}:F{end}"
    if first is not None:  # no invoice at all: one empty sheet that says so
        first.title = _sheet_name(labels["overview"], taken)
        title_rows(first, labels["empty"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def _font_path() -> Path | None:
    """DejaVu Sans as shipped with matplotlib: a Unicode TTF on every host,
    so comments with umlauts, dashes or euro signs print as written."""
    try:
        import matplotlib

        path = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans.ttf"
    except Exception:
        return None
    return path if path.is_file() else None


def _hours(value: Decimal, fmt: Any) -> str:
    return str(fmt(value))


def pdf(items: list[Sheet], labels: dict[str, Any], month_label: str, number: Any = None) -> bytes:
    """One PDF, landscape A4: a page (or more) per invoice, totals under each.

    `number(Decimal) -> str` formats hours in the reader's locale.
    """
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    fmt = number or (lambda v: f"{v:.2f}")
    doc = FPDF(orientation="L", unit="mm", format="A4")
    doc.set_auto_page_break(True, margin=14)
    doc.set_margins(12, 12, 12)
    font = _font_path()
    family = "Helvetica"
    if font is not None:
        doc.add_font("DejaVu", "", str(font))
        bold = font.with_name("DejaVuSans-Bold.ttf")
        doc.add_font("DejaVu", "B", str(bold if bold.is_file() else font))
        family = "DejaVu"

    def text(value: str) -> str:
        # The core Helvetica only has latin-1; never fail a download on a glyph.
        if family == "DejaVu":
            return value
        return value.encode("latin-1", "replace").decode("latin-1")

    def page_head(title: str) -> None:
        doc.add_page()
        doc.set_font(family, "B", 15)
        doc.cell(0, 8, text(labels["title"]), new_x="LMARGIN", new_y="NEXT")
        doc.set_font(family, "B", 12)
        doc.cell(0, 7, text(title), new_x="LMARGIN", new_y="NEXT")
        doc.set_font(family, "", 10)
        doc.cell(0, 6, text(month_label), new_x="LMARGIN", new_y="NEXT")
        doc.ln(3)

    head_style = FontFace(emphasis="BOLD", fill_color=(232, 238, 244))
    if len(items) > 1:
        page_head(labels["overview"])
        doc.set_font(family, "", 10)
        with doc.table(
            col_widths=(120, 30, 40),
            width=190,
            align="LEFT",
            text_align=("LEFT", "RIGHT", "RIGHT"),
            headings_style=head_style,
        ) as table:
            table.row([text(labels["invoice"]), text(labels["bookings"]), text(labels["billed"])])
            for s in items:
                table.row([text(s.title), str(len(s.rows)), _hours(s.billed, fmt)])
            table.row(
                [
                    text(labels["total"]),
                    str(sum(len(s.rows) for s in items)),
                    _hours(sum((s.billed for s in items), Decimal(0)), fmt),
                ],
                style=FontFace(emphasis="BOLD"),
            )
    if not items:
        page_head(labels["empty"])
    # Only the billed (rounded) hours: what the invoice charges (#408).
    for s in items:
        page_head(s.title)
        doc.set_font(family, "", 8.5)
        with doc.table(
            col_widths=(20, 40, 36, 34, 103, 40),
            text_align=("LEFT", "LEFT", "LEFT", "LEFT", "LEFT", "RIGHT"),
            headings_style=head_style,
            repeat_headings=1,
            line_height=4.6,
        ) as table:
            table.row(
                [
                    text(labels[k])
                    for k in ("date", "package", "task", "person", "comment", "billed")
                ]
            )
            for r in s.rows:
                table.row(
                    [
                        r.date.strftime("%d.%m.%Y") if r.date else "",
                        text(r.package),
                        text(r.task),
                        text(r.person),
                        text(r.comment),
                        _hours(r.billed, fmt),
                    ]
                )
            table.row(
                [
                    text(labels["total"]),
                    "",
                    "",
                    "",
                    text(labels["bookings_n"](len(s.rows))),
                    _hours(s.billed, fmt),
                ],
                style=FontFace(emphasis="BOLD"),
            )
    return bytes(doc.output())


# --------------------------------------------------------------------------
# Several files at once
# --------------------------------------------------------------------------


def zipped(files: list[tuple[str, bytes]]) -> bytes:
    """[(file name, bytes)] as one .zip -- "each invoice its own file"."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files:
            zf.writestr(name, data)
    return buf.getvalue()

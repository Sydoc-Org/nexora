"""Integration tests for nx_lib.views.reporting.

Seed test users have no reporting.* perms, so guarded routes return 403 (authed)
or redirect (anon). Endpoints whose first DB hit needs tables absent from the
TEST schema are asserted as (200, 500) to stay forward-compatible, mirroring the
dashboard route tests.
"""

import io
from datetime import date, datetime, time, timedelta
from unittest.mock import patch

from openpyxl import load_workbook

from nx_lib.db import engine_nexora_db


def test_reporting_anonymous_redirects_to_login(client):
    resp = client.get("/reporting", follow_redirects=False)
    assert resp.status_code == 302
    assert "/login" in resp.headers.get("Location", "")


def test_reporting_without_perm_returns_403(user_client):
    resp = user_client.get("/reporting")
    assert resp.status_code == 403


def test_reporting_page_renders_chat_panel_when_ai_enabled(admin_client):
    # TestAdmin holds reporting.ai.use (seed) -- the chat panel renders, and the
    # old Ask-AI mode button + inline panel it replaces are gone.
    resp = admin_client.get("/reporting")
    assert resp.status_code == 200
    assert b"rpChatPanel" in resp.data
    assert b"rpAiPanel" not in resp.data
    assert b"rpModeAi" not in resp.data


def test_reporting_page_caption_slots_need_only_explain_data(admin_client):
    """Regression for the Phase 4 review finding: the caption <div>s must be
    gated on reporting.ai.explain_data alone (D-CAPTION), not on the
    ai_explain_enabled AND-combo (which also requires reporting.sql.run --
    that extra requirement is for Surface C's live-SQL tool binding, an
    unrelated concern). A caller with explain_data but NOT sql.run must still
    see both #rpCaption (Advanced) and #rsCaption (Simple)."""

    def _perm(code):
        return code in ("reporting.view", "reporting.ai.explain_data")

    with (
        patch("nx_lib.security.has_permission", side_effect=_perm),
        patch("nx_lib.views.reporting.has_permission", side_effect=_perm),
    ):
        resp = admin_client.get("/reporting")
    assert resp.status_code == 200
    assert b'id="rpCaption"' in resp.data
    assert b'id="rsCaption"' in resp.data


def test_sources_without_perm_returns_403(user_client):
    resp = user_client.get("/api/reporting/sources")
    assert resp.status_code == 403


def test_run_invalid_json_returns_400_or_403(user_client):
    # no perm → 403 before body parsing
    resp = user_client.post("/api/reporting/run", data="not-json")
    assert resp.status_code in (400, 403)


def test_run_anonymous_redirects(client):
    resp = client.post("/api/reporting/run", json={}, follow_redirects=False)
    assert resp.status_code in (302, 401)


def test_reports_list_without_perm_403(user_client):
    resp = user_client.get("/api/reporting/reports")
    assert resp.status_code == 403


def test_sql_run_without_perm_403(user_client):
    resp = user_client.post(
        "/api/reporting/sql/run", json={"target": "statistics", "sql": "SELECT 1"}
    )
    assert resp.status_code in (400, 403)


def test_sql_run_anonymous_redirects(client):
    resp = client.post(
        "/api/reporting/sql/run",
        json={"target": "statistics", "sql": "SELECT 1"},
        follow_redirects=False,
    )
    assert resp.status_code in (302, 401)


def test_sql_ack_without_perm_403(user_client):
    resp = user_client.post("/api/reporting/sql/ack", json={})
    assert resp.status_code in (400, 403)


def test_sql_run_octopus_target_without_perm_403(user_client):
    # The base reporting.sql.run gate blocks before the Octopus target check.
    resp = user_client.post("/api/reporting/sql/run", json={"target": "octopus", "sql": "SELECT 1"})
    assert resp.status_code in (400, 403)


def test_sql_run_serializes_binary_and_time_cells(user_client):
    """A result row with bytes (varbinary/rowversion) or datetime.time must
    serialize as strings, not crash jsonify with a 500 (regression)."""
    columns = [{"field": "Data", "header": "Data"}, {"field": "T", "header": "T"}]
    rows = [[b"\x00\x01\x02", time(13, 45, 0)]]
    with (
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._has_acked", return_value=True),
        patch("nx_lib.views.reporting._authorize_sql_target"),
        patch("nx_lib.views.reporting._run_sql", return_value=(columns, rows)),
    ):
        resp = user_client.post(
            "/api/reporting/sql/run", json={"target": "statistics", "sql": "SELECT 1"}
        )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["rows"][0][0] == "0x000102"
    assert data["rows"][0][1] == "13:45:00"


def test_export_sql_target_coerces_bytes_and_control_char_cells(user_client):
    """A SQL-target export row with a raw bytes cell (varbinary) and a control
    character must download as xlsx, not 500 (Task 52 regression: openpyxl
    raises on both, and _safe_cell used to pass them through unchanged)."""
    columns = [{"field": "Data", "header": "Data"}, {"field": "Note", "header": "Note"}]
    rows = [[b"caf\xc3\xa9", "bell\x07ringer"]]
    with (
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._has_acked", return_value=True),
        patch("nx_lib.views.reporting._authorize_sql_target"),
        patch("nx_lib.views.reporting._run_sql", return_value=(columns, rows)),
    ):
        resp = user_client.post(
            "/api/reporting/export",
            json={
                "kind": "sql",
                "target": "statistics",
                "sql": "SELECT 1",
                "title": "Bytes Test",
                "format": "xlsx",
            },
        )
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert ws["A5"].value == "café"  # decoded from bytes, not the b'...' repr
    assert ws["B5"].value == "bellringer"  # control char stripped


# --- /api/reporting/export/grid (client-supplied grid; no DB access) ---


def test_export_grid_without_perm_403(user_client):
    resp = user_client.post(
        "/api/reporting/export/grid",
        json={"columns": [{"header": "A"}], "rows": [["x"]]},
    )
    assert resp.status_code == 403


def test_export_grid_anonymous_redirects(client):
    resp = client.post(
        "/api/reporting/export/grid",
        json={"columns": [{"header": "A"}], "rows": [["x"]]},
        follow_redirects=False,
    )
    assert resp.status_code in (302, 401)


def test_export_grid_xlsx_ok(admin_client):
    # TestAdmin holds reporting.export; the endpoint serializes without touching a DB.
    resp = admin_client.post(
        "/api/reporting/export/grid",
        json={
            "columns": [{"header": "Client"}, {"header": "Total"}],
            "rows": [["Acme", 1700.5], ["Globex", 5940.25]],
            "title": "Pivot",
            "format": "xlsx",
        },
    )
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["Content-Type"]
    assert "Pivot.xlsx" in resp.headers["Content-Disposition"]


def test_export_grid_csv_ok_has_bom(admin_client):
    resp = admin_client.post(
        "/api/reporting/export/grid",
        json={
            "columns": [{"header": "Client"}, {"header": "Total"}],
            "rows": [["Acme", 1700.5]],
            "title": "Pivot",
            "format": "csv",
        },
    )
    assert resp.status_code == 200
    assert resp.headers["Content-Type"].startswith("text/csv")
    assert "Pivot.csv" in resp.headers["Content-Disposition"]
    assert resp.data[:3] == b"\xef\xbb\xbf"  # UTF-8 BOM
    assert b"Client,Total" in resp.data


def test_export_grid_bad_body_400(admin_client):
    resp = admin_client.post("/api/reporting/export/grid", json={"columns": "nope"})
    assert resp.status_code == 400


def test_export_grid_csv_strips_control_char_cell(admin_client):
    """A client-supplied grid cell with a control character must download as
    csv, not leak the raw byte into the file (Task 52 regression)."""
    resp = admin_client.post(
        "/api/reporting/export/grid",
        json={
            "columns": [{"header": "Note"}],
            "rows": [["bell\x07ringer"]],
            "title": "Notes",
            "format": "csv",
        },
    )
    assert resp.status_code == 200
    assert b"\x07" not in resp.data
    assert b"bellringer" in resp.data


# --- Cross-user sharing (A2). TestAdmin owns the reports it creates. ---


def _create_report(admin_client, name="Share Test"):
    resp = admin_client.post(
        "/api/reporting/reports",
        json={
            "name": name,
            "definition": {
                "kind": "sql",
                "target": "statistics",
                "sql": "SELECT 1 AS one",
                "title": name,
            },
        },
    )
    assert resp.status_code == 200, resp.data
    return resp.get_json()["id"]


def test_shares_get_owner_default_private(admin_client):
    rid = _create_report(admin_client)
    try:
        resp = admin_client.get(f"/api/reporting/reports/{rid}/shares")
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["visibility"] == "private"
        assert data["shares"] == []
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_set_visibility_shared_reflects_in_list(admin_client):
    rid = _create_report(admin_client)
    try:
        resp = admin_client.post(
            f"/api/reporting/reports/{rid}/shares", json={"visibility": "shared"}
        )
        assert resp.status_code == 200
        assert resp.get_json()["visibility"] == "shared"
        lst = admin_client.get("/api/reporting/reports").get_json()
        row = next(r for r in lst if r["id"] == rid)
        assert row["visibility"] == "shared" and row["owned"] is True
        bad = admin_client.post(f"/api/reporting/reports/{rid}/shares", json={"visibility": "nope"})
        assert bad.status_code == 400
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_add_and_remove_user_share(admin_client):
    rid = _create_report(admin_client)
    try:
        resp = admin_client.post(
            f"/api/reporting/reports/{rid}/shares",
            json={"user": "user@test.local", "canEdit": True},
        )
        assert resp.status_code == 200
        shares = resp.get_json()["shares"]
        assert len(shares) == 1 and shares[0]["canEdit"] is True
        uid = shares[0]["userId"]
        rem = admin_client.delete(f"/api/reporting/reports/{rid}/shares/{uid}")
        assert rem.status_code == 200
        assert rem.get_json()["shares"] == []
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_share_unknown_user_404(admin_client):
    rid = _create_report(admin_client)
    try:
        resp = admin_client.post(
            f"/api/reporting/reports/{rid}/shares", json={"user": "ghost@nowhere.invalid"}
        )
        assert resp.status_code == 404
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_share_with_self_400(admin_client):
    rid = _create_report(admin_client)
    try:
        resp = admin_client.post(
            f"/api/reporting/reports/{rid}/shares", json={"user": "admin@test.local"}
        )
        assert resp.status_code == 400
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


# --- Task 57: one malformed saved definition must not 500 the whole library ---


def test_reports_list_survives_malformed_columns_row(admin_client):
    """A saved definition with a non-dict first column (cols[0]) must not
    500 the whole library listing for every user - _preview_kind falls back
    to a safe default kind for that row instead of raising.
    """
    good_rid = _create_report(admin_client, name="Good Def")
    bad_resp = admin_client.post(
        "/api/reporting/reports",
        json={
            "name": "Corrupt Def",
            "definition": {"kind": "table", "columns": ["not-a-dict"]},
        },
    )
    assert bad_resp.status_code == 200, bad_resp.data
    bad_rid = bad_resp.get_json()["id"]
    try:
        resp = admin_client.get("/api/reporting/reports")
        assert resp.status_code == 200
        by_id = {row["id"]: row for row in resp.get_json()}
        assert good_rid in by_id
        # The corrupt row survives with a safe fallback previewKind rather
        # than raising and taking the whole listing down with it.
        assert bad_rid in by_id
        assert by_id[bad_rid]["previewKind"] == "bar"
    finally:
        admin_client.delete(f"/api/reporting/reports/{good_rid}")
        admin_client.delete(f"/api/reporting/reports/{bad_rid}")


def test_reports_list_skips_row_when_preview_kind_still_raises(admin_client):
    """Defense-in-depth: even if some other unexpected shape slips past
    _preview_kind's internal guard, the per-row loop in api_reports_list
    must skip just that row (and log it) rather than 500 the whole library.
    """
    good_rid = _create_report(admin_client, name="Good Def 2")
    boom_resp = admin_client.post(
        "/api/reporting/reports",
        json={"name": "Boom Def", "definition": {"kind": "table", "marker": "boom"}},
    )
    assert boom_resp.status_code == 200, boom_resp.data
    boom_rid = boom_resp.get_json()["id"]

    from nx_lib.views.reporting import _preview_kind as real_preview_kind

    def _boom_preview_kind(defn):
        if isinstance(defn, dict) and defn.get("marker") == "boom":
            raise RuntimeError("simulated unexpected shape")
        return real_preview_kind(defn)

    try:
        with patch("nx_lib.views.reporting._preview_kind", side_effect=_boom_preview_kind):
            resp = admin_client.get("/api/reporting/reports")
        assert resp.status_code == 200
        ids = {row["id"] for row in resp.get_json()}
        assert good_rid in ids
        assert boom_rid not in ids
    finally:
        admin_client.delete(f"/api/reporting/reports/{good_rid}")
        admin_client.delete(f"/api/reporting/reports/{boom_rid}")


def test_shares_endpoints_without_perm_403(user_client):
    assert user_client.get("/api/reporting/reports/1/shares").status_code == 403
    assert (
        user_client.post(
            "/api/reporting/reports/1/shares", json={"visibility": "shared"}
        ).status_code
        == 403
    )


# --- Scheduled report delivery (A1) ---


def test_schedules_without_perm_403(user_client):
    assert user_client.get("/api/reporting/reports/1/schedules").status_code == 403
    assert (
        user_client.post(
            "/api/reporting/reports/1/schedules",
            json={"frequency": "daily", "hour": 6, "recipients": "a@x.com"},
        ).status_code
        == 403
    )


def test_schedule_crud(admin_client):
    rid = _create_report(admin_client)
    try:
        assert admin_client.get(f"/api/reporting/reports/{rid}/schedules").get_json() == []
        cr = admin_client.post(
            f"/api/reporting/reports/{rid}/schedules",
            json={
                "frequency": "daily",
                "hour": 6,
                "minute": 0,
                "format": "xlsx",
                "recipients": "a@x.com, b@y.com",
            },
        )
        assert cr.status_code == 200, cr.data
        sid = cr.get_json()["id"]
        lst = admin_client.get(f"/api/reporting/reports/{rid}/schedules").get_json()
        assert len(lst) == 1 and lst[0]["frequency"] == "daily" and lst[0]["nextRunAt"]

        up = admin_client.put(
            f"/api/reporting/reports/{rid}/schedules/{sid}",
            json={
                "frequency": "weekly",
                "hour": 7,
                "minute": 30,
                "weekday": 1,
                "format": "csv",
                "recipients": "c@z.com",
                "enabled": False,
            },
        )
        assert up.status_code == 200
        lst = admin_client.get(f"/api/reporting/reports/{rid}/schedules").get_json()
        assert lst[0]["frequency"] == "weekly" and lst[0]["enabled"] is False
        assert lst[0]["weekday"] == 1 and lst[0]["format"] == "csv"

        assert (
            admin_client.delete(f"/api/reporting/reports/{rid}/schedules/{sid}").status_code == 200
        )
        assert admin_client.get(f"/api/reporting/reports/{rid}/schedules").get_json() == []
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_schedules_overview_without_perm_403(user_client):
    assert user_client.get("/api/reporting/schedules").status_code == 403


def test_schedules_overview_lists_owned(admin_client):
    """Console Scheduled screen: GET /api/reporting/schedules returns every
    owned schedule with the joined report name."""
    rid = _create_report(admin_client)
    try:
        cr = admin_client.post(
            f"/api/reporting/reports/{rid}/schedules",
            json={
                "frequency": "daily",
                "hour": 6,
                "minute": 0,
                "format": "xlsx",
                "recipients": "a@x.com",
            },
        )
        assert cr.status_code == 200, cr.data
        rows = admin_client.get("/api/reporting/schedules").get_json()
        mine = [r for r in rows if r["reportId"] == rid]
        assert len(mine) == 1
        row = mine[0]
        assert row["reportName"]
        assert row["frequency"] == "daily" and row["enabled"] is True
        assert row["nextRunAt"]
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_share_targets_without_perm_403(user_client):
    assert user_client.get("/api/reporting/share_targets?q=ad").status_code == 403


def test_share_targets_typeahead(admin_client):
    """2+ chars returns matching users (username + display name only);
    shorter queries return an empty list."""
    assert admin_client.get("/api/reporting/share_targets?q=a").get_json() == []
    rows = admin_client.get("/api/reporting/share_targets?q=test").get_json()
    assert isinstance(rows, list) and len(rows) <= 8
    for row in rows:
        assert set(row) == {"username", "name"}


def test_sources_health_without_perm_403(user_client):
    assert user_client.get("/api/reporting/sources/health").status_code == 403


def test_sources_health_shape(admin_client):
    """One row per accessible source: id, ok flag, latencyMs and the real
    database name behind the source (None when the probe fails — the TEST
    env's engines may or may not be reachable)."""
    resp = admin_client.get("/api/reporting/sources/health")
    assert resp.status_code == 200
    data = resp.get_json()
    assert isinstance(data.get("sources"), list)
    for row in data["sources"]:
        assert set(row) == {"id", "ok", "latencyMs", "db"}
        assert isinstance(row["ok"], bool)
        if row["ok"]:
            assert row["db"]  # DB_NAME() rides along on a successful probe


def test_schedule_validation_400(admin_client):
    rid = _create_report(admin_client)
    try:
        assert (
            admin_client.post(
                f"/api/reporting/reports/{rid}/schedules",
                json={"frequency": "daily", "hour": 6, "recipients": "bad"},
            ).status_code
            == 400
        )
        assert (
            admin_client.post(
                f"/api/reporting/reports/{rid}/schedules",
                json={"frequency": "weekly", "hour": 6, "recipients": "a@x.com"},
            ).status_code
            == 400
        )
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_schedule_alert_fields_roundtrip(admin_client):
    rid = _create_report(admin_client)
    try:
        cr = admin_client.post(
            f"/api/reporting/reports/{rid}/schedules",
            json={
                "frequency": "daily",
                "hour": 6,
                "minute": 0,
                "format": "csv",
                "recipients": "a@x.com",
                "alertOp": "gt",
                "alertThreshold": 250,
            },
        )
        assert cr.status_code == 200, cr.data
        lst = admin_client.get(f"/api/reporting/reports/{rid}/schedules").get_json()
        assert lst[0]["alertOp"] == "gt" and lst[0]["alertThreshold"] == 250.0

        sid = lst[0]["id"]  # PUT without alert fields clears the condition (full-replace)
        up = admin_client.put(
            f"/api/reporting/reports/{rid}/schedules/{sid}",
            json={
                "frequency": "daily",
                "hour": 6,
                "minute": 0,
                "format": "csv",
                "recipients": "a@x.com",
                "enabled": True,
            },
        )
        assert up.status_code == 200
        lst = admin_client.get(f"/api/reporting/reports/{rid}/schedules").get_json()
        assert lst[0]["alertOp"] is None and lst[0]["alertThreshold"] is None
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_schedule_alert_validation_400(admin_client):
    rid = _create_report(admin_client)
    try:
        for bad in (
            {"alertOp": "eq", "alertThreshold": 1},
            {"alertOp": "gt"},
            {"alertOp": "gt", "alertThreshold": "soon"},
        ):
            r = admin_client.post(
                f"/api/reporting/reports/{rid}/schedules",
                json={"frequency": "daily", "hour": 6, "recipients": "a@x.com", **bad},
            )
            assert r.status_code == 400, r.data
    finally:
        admin_client.delete(f"/api/reporting/reports/{rid}")


def test_runner_dry_run_processes_due_table_report(admin_client):
    # End-to-end: register a 'table' source over Users, save a report on it, queue
    # a past-due schedule, then run the runner in dry-run (builds the report via
    # the session-independent runner; no mail, no NextRunAt advance).
    from ops import run_scheduled_reports

    src = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "sched_users",
            "kind": "curated",
            "label": "Sched Users",
            "permission": "reporting.source.docprocessing",
            "provider": "table",
            "engine": "nexora",
            "baseObject": "dbo.Users",
            "columns": [
                {
                    "field": "username",
                    "label": "Username",
                    "type": "string",
                    "filterable": True,
                    "sortable": True,
                }
            ],
            "enabled": True,
            "sortOrder": 15,
        },
    )
    src_id = src.get_json()["id"]
    rep = admin_client.post(
        "/api/reporting/reports",
        json={
            "name": "Sched Users Report",
            "definition": {
                "schemaVersion": 1,
                "source": "sched_users",
                "visualization": "table",
                "title": "Sched Users Report",
                "columns": [{"field": "username"}],
                "filters": [],
                "sort": [],
                "scope": {},
                "rowLimit": 10,
            },
        },
    )
    rid = rep.get_json()["id"]
    conn = engine_nexora_db.raw_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT userID FROM dbo.Users WHERE username = 'admin@test.local'")
        owner = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO dbo.ReportSchedules (ReportID, OwnerUserID, Recipients, Format, "
            "Frequency, Hour, Minute, Enabled, NextRunAt) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (rid, owner, "a@x.com", "csv", "daily", 6, 0, 1, datetime(2000, 1, 1)),
        )
        conn.commit()

        failed = run_scheduled_reports.run_once(dry_run=True)
        assert failed == 0
    finally:
        cur = conn.cursor()
        cur.execute("DELETE FROM dbo.ReportSchedules WHERE ReportID = ?", (rid,))
        conn.commit()
        conn.close()
        admin_client.delete(f"/api/reporting/reports/{rid}")
        admin_client.delete(f"/api/reporting/admin/sources/{src_id}")


def test_zero_dim_latest_metric_run_constrains_to_latest_bucket(admin_client):
    """#178 coverage gap: prove _prepare_run's decision logic (not just
    build_generic_query directly) computes and passes latest_of end-to-end.
    Mirrors backlog_history/backlog_total's shape (migrations 0053-0056):
    one metric, TotalMode='latest', one grainable field, zero-dim run. The
    admin metrics API has no TotalMode write path (Task 8 only wired the
    read path + the 0056 seed for backlog_total), so TotalMode is flipped
    via direct SQL after creating the metric, same idiom the runner test
    above uses for ReportSchedules."""
    src = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "latest_test_src",
            "kind": "curated",
            "label": "Latest Test Src",
            "permission": "reporting.source.docprocessing",
            "provider": "table",
            "engine": "nexora",
            "baseObject": "dbo.Users",
            "columns": [
                {
                    "field": "LastLoginAt",
                    "label": "Last Login",
                    "type": "datetime",
                    "filterable": True,
                    "sortable": True,
                    "grainable": True,
                }
            ],
            "enabled": True,
            "sortOrder": 17,
        },
    )
    src_id = src.get_json()["id"]
    met = admin_client.post(
        "/api/reporting/admin/metrics",
        json={
            "code": "latest_test_metric",
            "sourceId": "latest_test_src",
            "label": "Latest Test Metric",
            "aggregation": "count",
            "enabled": True,
            "sortOrder": 17,
        },
    )
    mid = met.get_json()["id"]
    conn = engine_nexora_db.raw_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            "UPDATE dbo.ReportingMetrics SET TotalMode = 'latest' WHERE Code = ?",
            ("latest_test_metric",),
        )
        conn.commit()
    finally:
        conn.close()

    try:
        resp = admin_client.post(
            "/api/reporting/run",
            json={
                "schemaVersion": 1,
                "source": "latest_test_src",
                "visualization": "table",
                "title": "Latest total",
                "columns": [],
                "filters": [],
                "sort": [],
                "scope": {},
                "rowLimit": 10,
                "metrics": [{"metric": "latest_test_metric"}],
            },
        )
        assert resp.status_code == 200, resp.data
        sql = resp.get_json()["sql"]
        # The decisive proof: _prepare_run computed latest_of="LastLoginAt"
        # (from the one grainable field + the all-'latest' modes set) and
        # passed it through -- build_generic_query's MAX() subquery fired.
        assert "SELECT MAX([LastLoginAt]) FROM [dbo].[Users]" in sql
        assert "[LastLoginAt] = (SELECT MAX([LastLoginAt])" in sql
    finally:
        admin_client.delete(f"/api/reporting/admin/metrics/{mid}")
        admin_client.delete(f"/api/reporting/admin/sources/{src_id}")


def test_runner_alert_skips_mail_and_advances(admin_client):
    # Non-metric definition -> the alert value is the row count (>=1 seeded user).
    # 'gt 1e9' never trips: no mail, NextRunAt advances. 'gte 1' trips: mail sent.
    from ops import run_scheduled_reports

    src = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "alert_users",
            "kind": "curated",
            "label": "Alert Users",
            "permission": "reporting.source.docprocessing",
            "provider": "table",
            "engine": "nexora",
            "baseObject": "dbo.Users",
            "columns": [
                {
                    "field": "username",
                    "label": "Username",
                    "type": "string",
                    "filterable": True,
                    "sortable": True,
                }
            ],
            "enabled": True,
            "sortOrder": 16,
        },
    )
    src_id = src.get_json()["id"]
    rep = admin_client.post(
        "/api/reporting/reports",
        json={
            "name": "Alert Users Report",
            "definition": {
                "schemaVersion": 1,
                "source": "alert_users",
                "visualization": "table",
                "title": "Alert Users Report",
                "columns": [{"field": "username"}],
                "filters": [],
                "sort": [],
                "scope": {},
                "rowLimit": 10,
            },
        },
    )
    rid = rep.get_json()["id"]
    conn = engine_nexora_db.raw_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT userID FROM dbo.Users WHERE username = 'admin@test.local'")
        owner = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO dbo.ReportSchedules (ReportID, OwnerUserID, Recipients, Format, "
            "Frequency, Hour, Minute, Enabled, NextRunAt, AlertOp, AlertThreshold) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (rid, owner, "a@x.com", "csv", "daily", 6, 0, 1, datetime(2000, 1, 1), "gt", 1e9),
        )
        conn.commit()

        with patch("ops.run_scheduled_reports.send_mail") as sm:
            assert run_scheduled_reports.run_once(dry_run=False) == 0
        sm.assert_not_called()
        cur.execute(
            "SELECT NextRunAt, LastRunAt FROM dbo.ReportSchedules WHERE ReportID = ?", (rid,)
        )
        nxt, last = cur.fetchone()
        if isinstance(nxt, str):
            nxt = datetime.fromisoformat(nxt)
        assert nxt > datetime(2020, 1, 1) and last is not None  # advanced, not re-fired

        cur.execute(
            "UPDATE dbo.ReportSchedules SET AlertOp='gte', AlertThreshold=1, NextRunAt=? "
            "WHERE ReportID = ?",
            (datetime(2000, 1, 1), rid),
        )
        conn.commit()
        with patch("ops.run_scheduled_reports.send_mail") as sm2:
            assert run_scheduled_reports.run_once(dry_run=False) == 0
        sm2.assert_called_once()
    finally:
        cur = conn.cursor()
        cur.execute("DELETE FROM dbo.ReportSchedules WHERE ReportID = ?", (rid,))
        conn.commit()
        conn.close()
        admin_client.delete(f"/api/reporting/reports/{rid}")
        admin_client.delete(f"/api/reporting/admin/sources/{src_id}")


# --- Seeded curated 'table' sources (A5 Generali, A6 Workitems) ---


def test_generali_and_workitems_sources_registered(admin_client):
    # TestAdmin holds both source permissions (seed); listing builds the field
    # catalog from ColumnsJSON with no Generali/Octopus DB access.
    by_id = {s["id"]: s for s in admin_client.get("/api/reporting/sources").get_json()}
    assert by_id["generali_pdqm"]["kind"] == "curated"
    assert any(f["field"] == "ParentCategory" for f in by_id["generali_pdqm"]["fields"])
    assert by_id["workitems"]["kind"] == "curated"
    assert any(f["field"] == "WorkItemIdentifier" for f in by_id["workitems"]["fields"])


# --- Source registry admin (A3) ---


def test_admin_sources_page_without_perm_403(user_client):
    assert user_client.get("/reporting/sources").status_code == 403
    assert user_client.get("/api/reporting/admin/sources").status_code == 403


def test_admin_sources_page_renders(admin_client):
    resp = admin_client.get("/reporting/sources")
    assert resp.status_code == 200
    assert b"reporting-sources-admin" in resp.data


def test_admin_sources_list_has_defaults(admin_client):
    data = admin_client.get("/api/reporting/admin/sources").get_json()
    assert any(s["id"] == "docprocessing" for s in data["defaults"])
    assert isinstance(data["rows"], list)


def test_admin_sources_crud(admin_client):
    create = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "crud_src",
            "kind": "sql",
            "label": "CRUD Source",
            "permission": "reporting.sql.run",
            "target": "statistics",
            "enabled": True,
            "sortOrder": 50,
        },
    )
    assert create.status_code == 200, create.data
    sid = create.get_json()["id"]
    try:
        rows = admin_client.get("/api/reporting/admin/sources").get_json()["rows"]
        assert any(r["id"] == sid and r["label"] == "CRUD Source" for r in rows)
        upd = admin_client.put(
            f"/api/reporting/admin/sources/{sid}",
            json={
                "code": "crud_src",
                "kind": "sql",
                "label": "Renamed",
                "permission": "reporting.sql.run",
                "target": "statistics",
                "enabled": False,
                "sortOrder": 50,
            },
        )
        assert upd.status_code == 200
        rows = admin_client.get("/api/reporting/admin/sources").get_json()["rows"]
        assert next(r for r in rows if r["id"] == sid)["label"] == "Renamed"
    finally:
        dele = admin_client.delete(f"/api/reporting/admin/sources/{sid}")
        assert dele.status_code == 200
    assert admin_client.delete(f"/api/reporting/admin/sources/{sid}").status_code == 404


def test_admin_sources_validation(admin_client):
    base = {"kind": "curated", "label": "L", "permission": "reporting.view"}
    assert (
        admin_client.post(
            "/api/reporting/admin/sources", json={**base, "code": "bad code!"}
        ).status_code
        == 400
    )
    assert (
        admin_client.post(
            "/api/reporting/admin/sources",
            json={"code": "ok1", "label": "L", "permission": "p", "kind": "weird"},
        ).status_code
        == 400
    )
    assert (
        admin_client.post(
            "/api/reporting/admin/sources",
            json={**base, "code": "ok2", "columns": "{not json"},
        ).status_code
        == 400
    )


def test_table_source_end_to_end(admin_client):
    # Register a generic 'table' source over the TEST Users table and run it.
    create = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "e2e_users",
            "kind": "curated",
            "label": "E2E Users",
            "permission": "reporting.source.docprocessing",
            "provider": "table",
            "engine": "nexora",
            "baseObject": "dbo.Users",
            "columns": [
                {
                    "field": "username",
                    "label": "Username",
                    "type": "string",
                    "filterable": True,
                    "sortable": True,
                },
                {
                    "field": "Email",
                    "label": "Email",
                    "type": "string",
                    "filterable": True,
                    "sortable": True,
                },
            ],
            "enabled": True,
            "sortOrder": 10,
        },
    )
    assert create.status_code == 200, create.data
    sid = create.get_json()["id"]
    try:
        srcs = admin_client.get("/api/reporting/sources").get_json()
        entry = next((s for s in srcs if s["id"] == "e2e_users"), None)
        assert entry is not None and any(f["field"] == "username" for f in entry["fields"])

        run = admin_client.post(
            "/api/reporting/run",
            json={
                "schemaVersion": 1,
                "source": "e2e_users",
                "visualization": "table",
                "title": "Users",
                "columns": [{"field": "username"}, {"field": "Email"}],
                "filters": [{"field": "username", "op": "contains", "value": "admin"}],
                "sort": [{"field": "username", "dir": "asc"}],
                "scope": {},
                "rowLimit": 100,
            },
        )
        assert run.status_code == 200, run.data
        data = run.get_json()
        assert data["rowCount"] >= 1
        flat = [c for row in data["rows"] for c in row]
        assert "admin@test.local" in flat
    finally:
        admin_client.delete(f"/api/reporting/admin/sources/{sid}")


def test_run_response_includes_sql_and_params(admin_client):
    """The run endpoint echoes the executed SQL and its bind parameters."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = "SELECT TOP (100) COUNT(*) AS [n] FROM [dbo].[MyView]"
    fake_params = [42, "hello"]
    fake_rows = [[99]]
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, fake_params, None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json={"source": "x"})
    assert resp.status_code == 200
    body = resp.get_json()
    assert "sql" in body and body["sql"].lstrip().upper().startswith("SELECT")
    assert "params" in body and isinstance(body["params"], list)
    assert body["params"] == [42, "hello"]


def test_run_response_includes_pretty_sql(admin_client):
    """/run also echoes a display-formatted copy (sqlPretty) of the SQL.

    The raw `sql` stays byte-exact (Copy and exports read it); sqlPretty is
    cosmetic and must be multi-line with placeholders preserved.
    """
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = (
        "SELECT TOP (100) [a] AS [a], COUNT(*) AS [n] FROM "
        "(SELECT [A] AS [a] FROM [dbo].[T] WHERE [D] >= ? AND [D] < ?) t "
        "GROUP BY [a] ORDER BY [n] DESC"
    )
    fake_params = ["2026-01-01", "2026-02-01"]
    fake_rows = [["x", 1]]
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, fake_params, None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json={"source": "x"})
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["sql"] == fake_sql  # raw untouched
    assert "sqlPretty" in body
    assert "\n" in body["sqlPretty"]  # actually formatted
    assert body["sqlPretty"].count("?") == 2  # placeholders preserved


def test_run_response_includes_inlined_display_sql(admin_client):
    """/run echoes sqlDisplay — the pretty SQL with parameter literals inlined
    (display/copy only; sql+params stay authoritative for execution)."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = (
        "SELECT TOP (100) [a] AS [a], COUNT(*) AS [n] FROM "
        "(SELECT [A] AS [a] FROM [dbo].[T] WHERE [D] >= ? AND [D] < ?) t "
        "GROUP BY [a] ORDER BY [n] DESC"
    )
    fake_params = ["2026-07-01", "2026-08-01"]
    fake_rows = [["x", 1]]
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, fake_params, None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json={"source": "x"})
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["sql"] == fake_sql  # raw untouched
    assert body["params"] == fake_params  # still echoed
    assert "?" not in body["sqlDisplay"]
    assert "'2026-07-01'" in body["sqlDisplay"]
    assert "'2026-08-01'" in body["sqlDisplay"]
    assert "\n" in body["sqlDisplay"]  # pretty-printed


# --- compare: true (Phase 3 comparison deltas) ---


def test_run_compare_true_with_token_filter_returns_comparison(admin_client):
    """compare: true + a single token filter runs the definition a second time
    over the prior window and surfaces it as a `comparison` key."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = "SELECT COUNT(*) AS [n] FROM [dbo].[T]"
    main_rows = [[10]]
    comparison_rows = [[7]]
    body = {
        "source": "x",
        "filters": [{"field": "CreatedDate", "op": "between", "value": {"token": "last_month"}}],
        "compare": True,
    }
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, [], None),
        ),
        patch("nx_lib.views.reporting._execute", side_effect=[main_rows, comparison_rows]),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json=body)
    assert resp.status_code == 200
    data = resp.get_json()
    assert "comparison" in data
    comp = data["comparison"]
    assert comp["columns"] == [{"field": "n", "header": "N"}]
    assert comp["rows"] == comparison_rows
    assert isinstance(comp["priorStart"], str)
    assert isinstance(comp["priorEnd"], str)


def test_run_compare_true_without_token_filter_omits_comparison(admin_client):
    """compare: true with no (or an ambiguous) token filter can't be shifted —
    the response stays 200 with no `comparison` key."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = "SELECT COUNT(*) AS [n] FROM [dbo].[T]"
    fake_rows = [[10]]
    body = {"source": "x", "filters": [], "compare": True}
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, [], None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows) as mock_execute,
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json=body)
    assert resp.status_code == 200
    data = resp.get_json()
    assert "comparison" not in data
    assert mock_execute.call_count == 1


def test_run_compare_execute_error_degrades_without_failing_main_run(admin_client):
    """A failure while running the comparison query must never fail the main
    run — it just logs and omits the `comparison` key."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = "SELECT COUNT(*) AS [n] FROM [dbo].[T]"
    main_rows = [[10]]
    body = {
        "source": "x",
        "filters": [{"field": "CreatedDate", "op": "between", "value": {"token": "last_month"}}],
        "compare": True,
    }
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, [], None),
        ),
        patch(
            "nx_lib.views.reporting._execute",
            side_effect=[main_rows, RuntimeError("comparison boom")],
        ),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json=body)
    assert resp.status_code == 200
    data = resp.get_json()
    assert "comparison" not in data
    assert data["rowCount"] == 1
    assert data["rows"] == main_rows


def test_run_without_compare_key_calls_execute_exactly_once(admin_client):
    """No `compare` key at all: single query, no comparison branch touched."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = "SELECT COUNT(*) AS [n] FROM [dbo].[T]"
    fake_rows = [[10]]
    body = {"source": "x"}
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(fake_cols, fake_sql, [], None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows) as mock_execute,
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._resolved_dates_meta", return_value=None),
    ):
        resp = admin_client.post("/api/reporting/run", json=body)
    assert resp.status_code == 200
    data = resp.get_json()
    assert "comparison" not in data
    assert mock_execute.call_count == 1


_TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4"
    "2mP8/5+hHgAHggJ/PchI7wAAAABJRU5ErkJggg=="
)


def test_export_xlsx_embeds_chart_image(admin_client):
    """When chartImage is sent with the export request, the XLSX contains xl/media/."""
    import io
    import zipfile

    fake_cols = [{"field": "n", "header": "N"}]
    fake_sql = "SELECT COUNT(*) AS [n] FROM [dbo].[T]"
    fake_rows = [[99]]
    body = {
        "source": "x",
        "format": "xlsx",
        "chartImage": "data:image/png;base64," + _TINY_PNG_B64,
    }
    with (
        patch("nx_lib.views.reporting._prepare_run", return_value=(fake_cols, fake_sql, [], None)),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
    ):
        resp = admin_client.post("/api/reporting/export", json=body)
    assert resp.status_code == 200
    assert resp.data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(resp.data)) as z:
        assert [n for n in z.namelist() if n.startswith("xl/media/")]


def test_export_ignores_garbage_chart_image(admin_client):
    """Garbage chartImage must not break the export — it degrades to chartless."""
    fake_cols = [{"field": "n", "header": "N"}]
    fake_rows = [[99]]
    body = {
        "source": "x",
        "format": "xlsx",
        "chartImage": "data:image/png;base64,@@@not-b64@@@",
    }
    with (
        patch(
            "nx_lib.views.reporting._prepare_run", return_value=(fake_cols, "SELECT 1", [], None)
        ),
        patch("nx_lib.views.reporting._execute", return_value=fake_rows),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
    ):
        resp = admin_client.post("/api/reporting/export", json=body)
    assert resp.status_code == 200
    assert resp.data[:2] == b"PK"


def test_shared_report_visible_to_non_owner(admin_client):
    # A report owned by user@test.local (a different user), marked 'shared', must
    # appear in admin's list (owned=false) and be loadable, but not manageable.
    conn = engine_nexora_db.raw_connection()
    rid = None
    try:
        cur = conn.cursor()
        cur.execute("SELECT userID FROM dbo.Users WHERE username = 'user@test.local'")
        other = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO dbo.Reports (OwnerUserID, Name, DefinitionJSON, Visibility) "
            "OUTPUT INSERTED.ReportID VALUES (?, ?, ?, 'shared')",
            (
                other,
                "Shared By Other",
                '{"kind":"sql","target":"statistics","sql":"SELECT 1 AS one",'
                '"title":"Shared By Other"}',
            ),
        )
        rid = cur.fetchone()[0]
        conn.commit()

        lst = admin_client.get("/api/reporting/reports").get_json()
        row = next((r for r in lst if r["id"] == rid), None)
        assert row is not None
        assert row["owned"] is False and row["canEdit"] is False
        assert row["ownerName"] == "user@test.local"
        got = admin_client.get(f"/api/reporting/reports/{rid}")
        assert got.status_code == 200 and got.get_json()["owned"] is False
        # Not the owner => cannot manage its shares.
        assert admin_client.get(f"/api/reporting/reports/{rid}/shares").status_code == 404
    finally:
        if rid is not None:
            cur = conn.cursor()
            cur.execute("DELETE FROM dbo.ReportShares WHERE ReportID = ?", (rid,))
            cur.execute("DELETE FROM dbo.Reports WHERE ReportID = ?", (rid,))
            conn.commit()
        conn.close()


# --- Error boundary: translated `error`, raw-English `detail` ---


def test_run_validation_error_is_translated_with_detail(admin_client):
    """A 400 from /run carries a gettext boundary message; the raw builder
    text moves to `detail` (devtools-only, never rendered on the page)."""
    from nx_lib.reporting.schema import ReportDefinitionError

    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            side_effect=ReportDefinitionError("at least one column is required"),
        ),
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
    ):
        resp = admin_client.post("/api/reporting/run", json={"source": "x"})
    assert resp.status_code == 400
    body = resp.get_json()
    assert body["detail"] == "at least one column is required"
    assert body["error"] != body["detail"]
    assert "column" not in body["error"]  # raw builder text no longer leaks


def test_sql_run_sandbox_error_is_translated_with_rule_and_detail(admin_client):
    from nx_lib.reporting.sandbox import SqlSandboxError

    with (
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._has_acked", return_value=True),
        patch("nx_lib.views.reporting._authorize_sql_target"),
        patch(
            "nx_lib.views.reporting._run_sql",
            side_effect=SqlSandboxError(
                "not_select", "only SELECT / WITH / set-operations are allowed"
            ),
        ),
    ):
        resp = admin_client.post(
            "/api/reporting/sql/run", json={"target": "statistics", "sql": "SELECT 1"}
        )
    assert resp.status_code == 400
    body = resp.get_json()
    assert body["rule"] == "not_select"
    assert body["detail"] == "only SELECT / WITH / set-operations are allowed"
    assert body["error"] != body["detail"]


def test_sql_run_generic_500_detail_is_humanized(admin_client):
    # Task 6: the generic 500 (a raw driver exception, not a SqlSandboxError)
    # carries a humanized `detail` — no ODBC/pyodbc noise reaching the client.
    odbc_text = (
        "('42000', '[42000] [Microsoft][ODBC SQL Server Driver][SQL Server]"
        "The ORDER BY clause is invalid in views, inline functions, derived "
        "tables, subqueries, and common table expressions, unless TOP, OFFSET "
        "or FOR XML is also specified. (1033) (SQLExecDirectW)')"
    )
    with (
        patch("nx_lib.security.has_permission", return_value=True),
        patch("nx_lib.views.reporting.has_permission", return_value=True),
        patch("nx_lib.views.reporting._has_acked", return_value=True),
        patch("nx_lib.views.reporting._authorize_sql_target"),
        patch("nx_lib.views.reporting._run_sql", side_effect=Exception(odbc_text)),
    ):
        resp = admin_client.post(
            "/api/reporting/sql/run", json={"target": "statistics", "sql": "SELECT 1"}
        )
    assert resp.status_code == 500
    body = resp.get_json()
    assert "SQLExecDirectW" not in body["detail"]
    assert "[Microsoft]" not in body["detail"]
    assert "Hint:" in body["detail"]


# --- forecast: definition toggle (issue #168) ---

_FC_DEF = {
    "schemaVersion": 1,
    "visualization": "table",
    "source": "docprocessing",
    "title": "Backlog over time",
    "columns": [{"field": "export_date", "grain": "month"}],
    "metrics": [{"metric": "doc_count"}],
    "filters": [],
    "sort": [],
    "scope": {"clients": [], "processes": []},
    "rowLimit": 5000,
    "forecast": {"enabled": True, "horizon": 3},
}

_FC_COLS = [{"field": "export_date"}, {"field": "doc_count"}]
_FC_ROWS = [[f"2025-{m:02d}-01", 10 + 2 * (m - 1)] for m in range(1, 9)]


def test_run_forecast_enabled_returns_block(admin_client):
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(_FC_COLS, "SELECT 1", [], None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=_FC_ROWS),
    ):
        resp = admin_client.post("/api/reporting/run", json=_FC_DEF)
    assert resp.status_code == 200
    fc = resp.get_json().get("forecast")
    assert fc and "unavailable" not in fc
    assert fc["horizon"] == 3 and len(fc["buckets"]) == 3
    assert fc["series"][0]["field"] == "doc_count"
    assert len(fc["series"][0]["values"]) == 3
    assert (
        fc["series"][0]["lower"][0] <= fc["series"][0]["values"][0] <= fc["series"][0]["upper"][0]
    )


def test_run_forecast_disabled_or_absent_omits_block(admin_client):
    quiet = dict(_FC_DEF, forecast={"enabled": False, "horizon": 3})
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(_FC_COLS, "SELECT 1", [], None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=_FC_ROWS),
    ):
        resp = admin_client.post("/api/reporting/run", json=quiet)
    assert resp.status_code == 200
    assert "forecast" not in resp.get_json()


def test_run_forecast_short_history_reports_unavailable(admin_client):
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(_FC_COLS, "SELECT 1", [], None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=_FC_ROWS[:3]),
    ):
        resp = admin_client.post("/api/reporting/run", json=_FC_DEF)
    assert resp.status_code == 200
    assert resp.get_json()["forecast"]["unavailable"] == "insufficient_history"


def test_export_forecast_appends_marker_rows(admin_client):
    body = dict(_FC_DEF, format="csv")
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            return_value=(_FC_COLS, "SELECT 1", [], None),
        ),
        patch("nx_lib.views.reporting._execute", return_value=_FC_ROWS),
    ):
        resp = admin_client.post("/api/reporting/export", json=body)
    assert resp.status_code == 200
    text = resp.data.decode("utf-8-sig")
    assert text.splitlines()[0].endswith("Forecast")
    assert text.count("forecast") == 3  # horizon 3 marker rows


# --- forecast: grain-dependent lookback widens the fit window (#178) ---

_FC_WIDE_DEF = {
    "schemaVersion": 1,
    "visualization": "table",
    "source": "docprocessing",
    "title": "Imports by day",
    "columns": [{"field": "import_date", "grain": "day"}],
    "metrics": [{"metric": "doc_count"}],
    "filters": [{"field": "import_date", "op": "between", "value": {"token": "this_month"}}],
    "sort": [],
    "scope": {"clients": [], "processes": []},
    "rowLimit": 5000,
    "forecast": {"enabled": True, "horizon": 3},
}

_FC_WIDE_COLS = [{"field": "import_date"}, {"field": "doc_count"}]
# Visible chart window: only 6 daily buckets (a real "this month, day grain"
# window) — nowhere near enough for a seasonal fit on its own.
_FC_WIDE_VISIBLE_ROWS = [[f"2026-08-{d:02d}", 10 + d] for d in range(1, 7)]
# Widened fit window: 60 contiguous daily buckets with a weekend dip, so
# trend_seasonal can only have come from the widened rerun.
_FC_WIDE_WIDE_START = date(2026, 6, 6)
_FC_WIDE_WIDE_ROWS = [
    [
        (_FC_WIDE_WIDE_START + timedelta(days=d)).isoformat(),
        4 if (_FC_WIDE_WIDE_START + timedelta(days=d)).weekday() >= 5 else 10,
    ]
    for d in range(60)
]


def test_run_forecast_fits_on_widened_history_not_visible_window(admin_client):
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            side_effect=[
                (_FC_WIDE_COLS, "SELECT 1", [], None),
                (_FC_WIDE_COLS, "SELECT 2", [], None),
            ],
        ) as mock_prepare,
        patch(
            "nx_lib.views.reporting._execute",
            side_effect=[_FC_WIDE_VISIBLE_ROWS, _FC_WIDE_WIDE_ROWS],
        ),
    ):
        resp = admin_client.post("/api/reporting/run", json=_FC_WIDE_DEF)
    assert resp.status_code == 200
    body = resp.get_json()
    # The visible chart response is unaffected by the wider fit window.
    assert body["rows"] == _FC_WIDE_VISIBLE_ROWS
    assert mock_prepare.call_count == 2
    widened_rd = mock_prepare.call_args_list[1].args[0]
    widened_filter = widened_rd["filters"][0]
    assert widened_filter["op"] == "between"
    assert isinstance(widened_filter["value"], list) and len(widened_filter["value"]) == 2
    assert "compare" not in widened_rd
    fc = body.get("forecast")
    assert fc and fc.get("method") == "trend_seasonal"


def test_export_forecast_fits_on_widened_history_not_visible_window(admin_client):
    # Same widen-refit-with-fallback as api_run (via the shared _forecast_for
    # helper) — the export's forecast must not silently regress to a
    # trend-only fit on the 6-row visible window (#178 finding 2).
    body = dict(_FC_WIDE_DEF, format="csv")
    with (
        patch(
            "nx_lib.views.reporting._prepare_run",
            side_effect=[
                (_FC_WIDE_COLS, "SELECT 1", [], None),
                (_FC_WIDE_COLS, "SELECT 2", [], None),
            ],
        ) as mock_prepare,
        patch(
            "nx_lib.views.reporting._execute",
            side_effect=[_FC_WIDE_VISIBLE_ROWS, _FC_WIDE_WIDE_ROWS],
        ),
    ):
        resp = admin_client.post("/api/reporting/export", json=body)
    assert resp.status_code == 200
    assert mock_prepare.call_count == 2
    widened_rd = mock_prepare.call_args_list[1].args[0]
    assert widened_rd["filters"][0]["op"] == "between"
    assert "compare" not in widened_rd
    text = resp.data.decode("utf-8-sig")
    # horizon 3 (explicit in _FC_WIDE_DEF) -> 3 marker rows; only reachable if
    # the widened rerun actually produced a usable (non-unavailable) fit.
    assert text.count("forecast") == 3


def test_runner_forecast_export_rows_failure_still_sends_mail(admin_client):
    # A forecast_export_rows failure must degrade to the unmarked attachment,
    # not skip the mail or stall NextRunAt (#168).
    from ops import run_scheduled_reports

    src = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "sched_fc_users",
            "kind": "curated",
            "label": "Sched Forecast Users",
            "permission": "reporting.source.docprocessing",
            "provider": "table",
            "engine": "nexora",
            "baseObject": "dbo.Users",
            "columns": [
                {
                    "field": "username",
                    "label": "Username",
                    "type": "string",
                    "filterable": True,
                    "sortable": True,
                }
            ],
            "enabled": True,
            "sortOrder": 17,
        },
    )
    src_id = src.get_json()["id"]
    rep = admin_client.post(
        "/api/reporting/reports",
        json={
            "name": "Sched Forecast Users Report",
            "definition": {
                "schemaVersion": 1,
                "source": "sched_fc_users",
                "visualization": "table",
                "title": "Sched Forecast Users Report",
                "columns": [{"field": "username"}],
                "filters": [],
                "sort": [],
                "scope": {},
                "rowLimit": 10,
                "forecast": {"enabled": True, "horizon": 3},
            },
        },
    )
    rid = rep.get_json()["id"]
    conn = engine_nexora_db.raw_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT userID FROM dbo.Users WHERE username = 'admin@test.local'")
        owner = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO dbo.ReportSchedules (ReportID, OwnerUserID, Recipients, Format, "
            "Frequency, Hour, Minute, Enabled, NextRunAt) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (rid, owner, "a@x.com", "csv", "daily", 6, 0, 1, datetime(2000, 1, 1)),
        )
        conn.commit()

        with (
            patch(
                "ops.run_scheduled_reports.compute_forecast",
                return_value={"buckets": [1, 2, 3], "series": [{"field": "username"}]},
            ),
            patch(
                "ops.run_scheduled_reports.forecast_export_rows",
                side_effect=Exception("boom"),
            ),
            patch("ops.run_scheduled_reports.send_mail") as sm,
        ):
            assert run_scheduled_reports.run_once(dry_run=False) == 0
        sm.assert_called_once()

        cur.execute(
            "SELECT NextRunAt, LastRunAt FROM dbo.ReportSchedules WHERE ReportID = ?", (rid,)
        )
        nxt, last = cur.fetchone()
        if isinstance(nxt, str):
            nxt = datetime.fromisoformat(nxt)
        assert nxt > datetime(2020, 1, 1) and last is not None  # advanced, not stalled
    finally:
        cur = conn.cursor()
        cur.execute("DELETE FROM dbo.ReportSchedules WHERE ReportID = ?", (rid,))
        conn.commit()
        conn.close()
        admin_client.delete(f"/api/reporting/reports/{rid}")
        admin_client.delete(f"/api/reporting/admin/sources/{src_id}")


def _create_field_values_source(admin_client):
    """A 'backlog_history'-shaped table source (#178), created dynamically
    since the TEST NexoraDB fixture doesn't seed the real migration-0053 row.
    Reuses the already-granted reporting.source.docprocessing permission,
    same idiom as test_zero_dim_latest_metric_run_constrains_to_latest_bucket."""
    src = admin_client.post(
        "/api/reporting/admin/sources",
        json={
            "code": "field_values_test_src",
            "kind": "curated",
            "label": "Field Values Test Src",
            "permission": "reporting.source.docprocessing",
            "provider": "table",
            "engine": "nexora",
            "baseObject": "dbo.Users",
            "columns": [
                {
                    "field": "username",
                    "label": "Username",
                    "type": "string",
                    "filterable": True,
                    "sortable": True,
                },
                {
                    "field": "locale",
                    "label": "Locale",
                    "type": "string",
                    "filterable": False,
                    "sortable": True,
                },
            ],
            "enabled": True,
            "sortOrder": 18,
        },
    )
    return src.get_json()["id"]


def test_field_values_returns_distinct_values(admin_client):
    """#178: /api/reporting/field_values returns the distinct values (from
    _execute) of a whitelisted filterable field of a table source, unwrapped
    from row tuples into a flat list."""
    src_id = _create_field_values_source(admin_client)
    fake_rows = [["01_EasyTax"], ["03_Invoice_New"]]
    try:
        with patch("nx_lib.views.reporting._execute", return_value=fake_rows) as mock_execute:
            resp = admin_client.post(
                "/api/reporting/field_values",
                json={"source": "field_values_test_src", "field": "username"},
            )
        assert resp.status_code == 200, resp.data
        assert resp.get_json() == {"values": ["01_EasyTax", "03_Invoice_New"]}
        mock_execute.assert_called_once()
        sql = mock_execute.call_args[0][1]
        assert "SELECT DISTINCT TOP (100) [username]" in sql
        assert "[dbo].[Users]" in sql
    finally:
        admin_client.delete(f"/api/reporting/admin/sources/{src_id}")


def test_field_values_unknown_field_returns_400(admin_client):
    """An unwhitelisted (and a non-filterable) field is rejected with 400
    before any query executes."""
    src_id = _create_field_values_source(admin_client)
    try:
        with patch("nx_lib.views.reporting._execute") as mock_execute:
            resp = admin_client.post(
                "/api/reporting/field_values",
                json={"source": "field_values_test_src", "field": "Nope"},
            )
        assert resp.status_code == 400
        assert "error" in resp.get_json()
        mock_execute.assert_not_called()

        with patch("nx_lib.views.reporting._execute") as mock_execute:
            resp = admin_client.post(
                "/api/reporting/field_values",
                json={"source": "field_values_test_src", "field": "locale"},
            )
        assert resp.status_code == 400
        mock_execute.assert_not_called()
    finally:
        admin_client.delete(f"/api/reporting/admin/sources/{src_id}")


def test_field_values_unknown_source_returns_400(admin_client):
    resp = admin_client.post(
        "/api/reporting/field_values",
        json={"source": "does_not_exist", "field": "username"},
    )
    assert resp.status_code == 400


def test_field_values_without_perm_403(user_client):
    resp = user_client.post(
        "/api/reporting/field_values",
        json={"source": "field_values_test_src", "field": "username"},
    )
    assert resp.status_code == 403

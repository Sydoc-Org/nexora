"""Integration tests for nx_lib.views.finance -- page, section API, CSV export (#408).

Routes covered:
- GET /finance                          page (finance.view-gated)
- GET /api/finance/section/<key>        one section's figures as JSON
- GET /api/finance/export.csv           the whole month as CSV
- GET /api/finance/bps-export           billable BPS hours as .xlsx / .pdf / .zip
- POST /api/finance/close|reopen        freeze / unfreeze a month (#415)

The test tier seeds finance.view (sql/test/seed.sql, TestAdmin holds every
code) but none of the billing sources, so against the seeded registry every
section answers with its error payload -- exactly what the page must render
for a source that is missing. The happy path stubs the three registry
touch-points the view uses, so no Statistics database is needed.
"""

import pytest

import nx_lib.views.finance as fv

FAKE_SOURCE = {
    "id": "em_invoice",
    "kind": "curated",
    "label": "Elektro-Material — Verrechnung",
    "permission": "reporting.source.em_invoice.use",
    "engine": "statistics",
    "provider": "table",
    "baseObject": "dbo.EM_Invoice",
    "columns": [
        {"field": "ExportEM_dt", "type": "date", "grainable": True},
        {"field": "Eingang", "type": "string"},
        {"field": "OrdItmPosCount", "type": "number"},
        {"field": "AnzImagesOut", "type": "number"},
    ],
}
FAKE_METRICS = {
    "em_documents": {
        "aggregation": "count",
        "base_field": None,
        "filter": None,
        "label": "Dokumente",
    },
    "em_images_out": {
        "aggregation": "sum",
        "base_field": "AnzImagesOut",
        "filter": None,
        "label": "Bilder",
    },
    "em_order_positions": {
        "aggregation": "sum",
        "base_field": "OrdItmPosCount",
        "filter": None,
        "label": "Positionen",
    },
}


@pytest.fixture()
def fake_registry(monkeypatch):
    """Route the EM section at an in-memory source; every other one stays unregistered."""
    calls = []

    def execute(engine, sql, params):
        calls.append((sql, params))
        if "GROUP BY" in sql:
            return [
                ("OPEX Scan Scanner", 120, 340, 88),
                ("E_MAIL", 30, 61, 12),
                ("Nexora", 0, 0, 0),
            ]
        if params and params[-1] == "2026-09-01":  # the previous-month query
            return [(140, 380, 90)]
        return [(150, 401, 100)]

    monkeypatch.setattr(
        fv, "_get_effective_source", lambda code: FAKE_SOURCE if code == "em_invoice" else None
    )
    monkeypatch.setattr(fv, "_metrics_for_source", lambda sid, locale=None: dict(FAKE_METRICS))
    monkeypatch.setattr(fv, "_CURATED_ENGINES", {"statistics": object()})
    monkeypatch.setattr(fv, "_execute", execute)
    return calls


# ---- gates -----------------------------------------------------------------


def test_finance_anonymous_redirects_to_login(client):
    resp = client.get("/finance", follow_redirects=False)
    assert resp.status_code == 302
    assert "/login" in resp.headers.get("Location", "")


@pytest.mark.parametrize(
    "path",
    [
        "/finance",
        "/api/finance/section/compass",
        "/api/finance/export.csv",
        "/api/finance/bps-export?month=2026-09",
    ],
)
def test_finance_without_perm_returns_403(noperm_client, path):
    assert noperm_client.get(path).status_code == 403


def test_finance_user_without_the_code_returns_403(user_client):
    """dashboard.view alone (TestUser) does not open the accounting page."""
    assert user_client.get("/finance").status_code == 403


# ---- the page --------------------------------------------------------------


def test_finance_page_renders_every_section_shell(admin_client):
    resp = admin_client.get("/finance?month=2026-08")
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert 'data-testid="finance-title"' in html
    for key in (
        "elektro_material",
        "compass",
        "privera_posteingang",
        "privera_invoice",
        "privera_nachsendungen",
        "privera_neuzugaenge",
        "frigemo",
        "xpert",
        "bucherer",
        "mediamarkt",
        "bps",
    ):
        assert f'data-testid="finance-section-{key}"' in html
    assert 'data-month="2026-08"' in html  # the picker's selected month
    assert 'data-testid="finance-period-button"' in html
    assert 'class="nx-app nx-sydoc"' in html
    assert ">August<" in html  # the headline's month word
    assert "month=2026-07" in html  # the previous-month link
    assert "export.csv?month=2026-08" in html


def test_finance_page_ignores_a_garbage_month(admin_client):
    resp = admin_client.get("/finance?month=nope")
    assert resp.status_code == 200
    assert 'data-testid="finance-period-button"' in resp.get_data(as_text=True)


def test_finance_page_marks_closed_months_for_the_picker(admin_client, monkeypatch):
    monkeypatch.setattr(fv, "_closed_months_safe", lambda: {"2026-07"})
    html = admin_client.get("/finance?month=2026-08").get_data(as_text=True)
    assert '"value": "2026-07"' in html
    assert '"state": "closed"' in html


def test_finance_page_is_reachable_from_the_sidebar(admin_client):
    resp = admin_client.get("/dashboard")
    assert resp.status_code == 200
    assert "/finance" in resp.get_data(as_text=True)


# ---- the section API -------------------------------------------------------


def test_section_api_unknown_key_is_404(admin_client):
    assert admin_client.get("/api/finance/section/nope").status_code == 404


def test_section_api_reports_an_unregistered_source_in_place(admin_client):
    """The seeded registry has no billing sources: the section answers with its
    error, status 200, so the page shows it where the figures would be and the
    other sections are unaffected."""
    resp = admin_client.get("/api/finance/section/compass?month=2026-08")
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["key"] == "compass"
    assert body["client"] == "Compass Group"
    assert body["month"] == "2026-08"
    assert body["blocks"] == []
    assert "compass_invoice" in body["error"]


def test_section_api_returns_figures_and_breakdowns(admin_client, fake_registry):
    resp = admin_client.get("/api/finance/section/elektro_material?month=2026-09")
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["error"] is None
    assert body["source"] == {"code": "em_invoice", "label": "Elektro-Material — Verrechnung"}
    block = body["blocks"][0]
    assert block["figures"] == [
        {"code": "em_documents", "label": "Dokumente", "value": 150, "prev": 140},
        {"code": "em_images_out", "label": "Bilder", "value": 401, "prev": 380},
        {"code": "em_order_positions", "label": "Positionen", "value": 100, "prev": 90},
    ]
    br = block["breakdowns"][0]
    assert br["dim"] == "Eingang"
    assert [r["key"] for r in br["rows"]] == ["OPEX Scan Scanner", "E_MAIL"]  # the 0-row is gone
    assert br["totals"] == [150, 401, 100]
    # Three statements: this month, the month before, the channel breakdown.
    assert len(fake_registry) == 3
    assert fake_registry[0][1][-2:] == ["2026-09-01", "2026-10-01"]


def test_section_api_reports_a_query_failure_in_place(admin_client, fake_registry, monkeypatch):
    def boom(engine, sql, params):
        raise RuntimeError("Invalid object name 'dbo.EM_Invoice'.\nsecond line")

    monkeypatch.setattr(fv, "_execute", boom)
    body = admin_client.get("/api/finance/section/elektro_material?month=2026-09").get_json()
    assert body["error"]
    assert body["detail"] == "Invalid object name 'dbo.EM_Invoice'."
    assert body["blocks"] == []


# ---- the export ------------------------------------------------------------


def test_export_is_a_csv_download_with_one_line_per_figure(admin_client, fake_registry):
    resp = admin_client.get("/api/finance/export.csv?month=2026-09")
    assert resp.status_code == 200
    assert resp.mimetype == "text/csv"
    assert 'filename="sydoc-finance-2026-09.csv"' in resp.headers["Content-Disposition"]
    assert resp.headers["Cache-Control"] == "no-store"
    text = resp.get_data(as_text=True)
    assert text.startswith("﻿")
    lines = text.strip().splitlines()
    assert lines[0].endswith(",2026-09,Previous month,Detail")
    assert "Elektro-Material,,by export date,figure,,,Dokumente,150,140," in lines
    assert "Elektro-Material,,by export date,breakdown,Channel,E_MAIL,Bilder,61,," in lines
    # Every unregistered section still appears, as its error line.
    assert any(line.startswith("Compass Group,,,error,") for line in lines)


# ---- month close (#415) ----------------------------------------------------


@pytest.fixture()
def all_sections_readable(monkeypatch):
    """Every section answers with one figure, so a close can go through."""
    value = {"n": 7}

    def payload(section, year, month):
        return {
            "key": section.key,
            "client": section.client,
            "title": section.title,
            "group": section.group,
            "source": {"code": section.source, "label": None},
            "note": None,
            "blocks": [
                {
                    "basis": "b",
                    "figures": [{"code": "x", "label": "Docs", "value": value["n"], "prev": 1}],
                    "breakdowns": [],
                }
            ],
            "bookings": None,
            "error": None,
        }

    monkeypatch.setattr(fv, "_section_payload", payload)
    return value


@pytest.fixture()
def reopen_after(admin_client):
    yield
    admin_client.post("/api/finance/reopen?month=2026-08")


@pytest.mark.parametrize(
    "path", ["/api/finance/close?month=2026-08", "/api/finance/reopen?month=2026-08"]
)
def test_close_and_reopen_need_finance_close(user_client, path):
    assert user_client.post(path).status_code == 403


@pytest.mark.parametrize("month", ["", "nope", "2026-13", "2999-01"])
def test_close_refuses_an_unnamed_running_or_future_month(admin_client, month):
    """A write never falls back to a default month, and a running month cannot close."""
    resp = admin_client.post(f"/api/finance/close?month={month}")
    assert resp.status_code == 400
    assert resp.get_json()["error"]


def test_close_refuses_while_a_section_cannot_be_read(admin_client):
    """The test tier registers no billing sources: every section errors, so nothing is frozen."""
    resp = admin_client.post("/api/finance/close?month=2026-08")
    assert resp.status_code == 409
    assert "Compass Group" in resp.get_json()["error"]
    body = admin_client.get("/api/finance/section/compass?month=2026-08").get_json()
    assert body["closed"] is None


def test_a_closed_month_serves_its_snapshot_and_reports_live_drift(
    admin_client, all_sections_readable, reopen_after
):
    resp = admin_client.post("/api/finance/close?month=2026-08")
    assert resp.status_code == 200, resp.get_json()
    assert admin_client.post("/api/finance/close?month=2026-08").status_code == 409

    # Nothing moved: the snapshot is served, with an empty drift list.
    body = admin_client.get("/api/finance/section/compass?month=2026-08").get_json()
    assert body["closed"]["by"]
    assert body["blocks"][0]["figures"][0]["value"] == 7
    assert body["live_diff"] == []

    # The source is edited after the close: the snapshot stays, the drift shows.
    all_sections_readable["n"] = 9
    body = admin_client.get("/api/finance/section/compass?month=2026-08").get_json()
    assert body["blocks"][0]["figures"][0]["value"] == 7
    assert body["live_diff"] == [{"label": "Docs", "closed": 7, "live": 9}]
    live = admin_client.get("/api/finance/section/compass?month=2026-08&live=1").get_json()
    assert live["closed"] is None and live["blocks"][0]["figures"][0]["value"] == 9

    page = admin_client.get("/finance?month=2026-08").get_data(as_text=True)
    assert 'data-testid="finance-closed-badge"' in page
    assert 'data-action="reopen"' in page

    csv = admin_client.get("/api/finance/export.csv?month=2026-08")
    assert 'filename="sydoc-finance-2026-08-closed.csv"' in csv.headers["Content-Disposition"]
    assert "Compass Group,,b,figure,,,Docs,7,1," in csv.get_data(as_text=True).splitlines()

    assert admin_client.post("/api/finance/reopen?month=2026-08").status_code == 200
    assert admin_client.post("/api/finance/reopen?month=2026-08").status_code == 404
    body = admin_client.get("/api/finance/section/compass?month=2026-08").get_json()
    assert body["closed"] is None and body["blocks"][0]["figures"][0]["value"] == 9


# ---- the BPS hours export -------------------------------------------------


def _bps_payload():
    from nx_lib.finance import SECTIONS_BY_KEY

    spec = SECTIONS_BY_KEY["bps"].bookings
    cols = [{"field": f} for f in spec.columns] + [{"field": "billed"}]

    def row(customer, package, hours):
        return ["2026-09-02", customer, package, "Change", "Ben", hours, "Kommentar ä", 0]

    return {
        "key": "bps",
        "error": None,
        "bookings": {
            "columns": cols,
            "groups": [
                {"key": "Aveniq", "rows": [row("Aveniq", "BFH", 0.3333)]},
                {
                    "key": "Privera",
                    "rows": [
                        row("Privera", "Tagesgeschäft Neuzugänge", 1),
                        row("Privera", "Posteingang", 0.5),
                    ],
                },
            ],
        },
    }


@pytest.fixture()
def bps_payload(monkeypatch):
    monkeypatch.setattr(fv, "_section_payload", lambda section, y, m: _bps_payload())


def test_bps_export_is_one_workbook_with_a_sheet_per_invoice(admin_client, bps_payload):
    import io

    from openpyxl import load_workbook

    resp = admin_client.get("/api/finance/bps-export?month=2026-09&format=xlsx")
    assert resp.status_code == 200
    assert 'filename="sydoc-bps-2026-09.xlsx"' in resp.headers["Content-Disposition"]
    assert resp.headers["Cache-Control"] == "no-store"
    wb = load_workbook(io.BytesIO(resp.data))
    # The overview, then the page's order: Privera before Aveniq, its streams split.
    assert wb.sheetnames[1:] == ["Privera Posteingang", "Privera Neuzugänge", "Aveniq"]
    aveniq = wb["Aveniq"]
    values = [c.value for c in aveniq[6]]
    assert values[5:] == [0.3333, 0.5]  # booked, billed (rounded up to the quarter)


def test_bps_export_one_invoice_as_pdf(admin_client, bps_payload):
    resp = admin_client.get(
        "/api/finance/bps-export?month=2026-09&format=pdf&sheet=privera-neuzugaenge"
    )
    assert resp.status_code == 200
    assert resp.mimetype == "application/pdf" and resp.data.startswith(b"%PDF")
    assert "sydoc-bps-2026-09-privera-neuzugaenge.pdf" in resp.headers["Content-Disposition"]


def test_bps_export_separate_files_come_as_a_zip(admin_client, bps_payload):
    import io
    import zipfile

    resp = admin_client.get("/api/finance/bps-export?month=2026-09&format=pdf&files=separate")
    assert resp.status_code == 200 and resp.mimetype == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(resp.data)).namelist()
    assert names == [
        "sydoc-bps-2026-09-privera-posteingang.pdf",
        "sydoc-bps-2026-09-privera-neuzugaenge.pdf",
        "sydoc-bps-2026-09-aveniq.pdf",
    ]


def test_bps_export_refuses_an_unknown_format_or_invoice(admin_client, bps_payload):
    assert admin_client.get("/api/finance/bps-export?format=docx").status_code == 400
    assert admin_client.get("/api/finance/bps-export?sheet=nope").status_code == 404


def test_bps_section_lists_its_export_sheets(admin_client, bps_payload):
    body = admin_client.get("/api/finance/section/bps?month=2026-09").get_json()
    assert [e["key"] for e in body["exports"]] == [
        "privera-posteingang",
        "privera-neuzugaenge",
        "aveniq",
    ]

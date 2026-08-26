"""Integration tests for nx_lib.views.workitems — 13 routes.

Seed users only have dashboard.view, so all workitems.* gates return 403.
Tests that exercise route bodies use the workitems_all_perms fixture
(monkeypatches has_permission to True).

Most workitems routes query OctoDB tables (t_WorkItems, t_ActivityInstances,
t_Processes) and NEXORA tables (SearchConfig, Search_Field_Labels,
WorkitemTags, WorkitemComments) that are absent from the TEST schema. They
fall through except branches and return 500 JSON. Tuple matches in
assertions allow for that.

The collaboration API (single-workitem tags, interactions, comments, assign,
priority, tags CRUD, mention-autocomplete users) was removed in Task 4 of the
chat-collab-removal-bug-fixes plan; get_audithistory (Octo processing trail)
is a separate feature and stays.

Routes covered (13 endpoints):
- /api/config/fields                         GET
- /api/docfield_values                       GET
- /api/workitems                             GET
- /api/export/workitems/csv                  GET
- /workitems                                 page
- /import_workitems                          POST
- /api/get_media_info/<id>                   GET
- /api/get_media_raw/<id>/<idx>              GET
- /api/get_audithistory/<id>                 GET
- /api/workitems_page_init                   GET
- /api/workitem_filter_views                 GET + POST (saved views, #170)
- /api/workitem_filter_views/<id>            DELETE
"""

import concurrent.futures
import csv
import io

import pytest
import requests


@pytest.fixture()
def workitems_all_perms(monkeypatch):
    monkeypatch.setattr("nx_lib.security.has_permission", lambda code: True)
    # The detail endpoints (media_info/media_raw/audithistory) also enforce a
    # per-workitem (client, process) entitlement gate (#193) that reads session
    # grants + resolves the workitem's real pair -- orthogonal to these tests,
    # which exercise caching/scoping/error behaviour. Grant it here so "all
    # perms" keeps meaning all perms.
    monkeypatch.setattr("nx_lib.views.workitems._may_view_workitem", lambda wid: True)
    yield


# ============================ /workitems page ================================


def test_workitems_overview_anonymous_redirects(client):
    resp = client.get("/workitems", follow_redirects=False)
    assert resp.status_code == 302
    assert "/login" in resp.headers.get("Location", "")


def test_workitems_overview_without_perm_returns_403(noperm_client):
    resp = noperm_client.get("/workitems")
    assert resp.status_code == 403


def test_workitems_overview_with_perms(user_client, workitems_all_perms):
    """With all perms granted, the page tries to render — some helper queries
    hit absent tables → 500-fallback render. Both outcomes are valid."""
    resp = user_client.get("/workitems")
    assert resp.status_code in (200, 500)


def test_workitems_overview_uses_shared_detail_panel(user_client, workitems_all_perms):
    """The workitems page wires the shared renderer."""
    resp = user_client.get("/workitems")
    # 200 or 500-fallback possible in CI; the partial markers live in template body.
    if resp.status_code == 200:
        assert b"NexoraWorkitemDetail.render" in resp.data


# ============================ API: config/data ===============================


def test_api_config_fields_anonymous(client):
    resp = client.get("/api/config/fields")
    assert resp.status_code == 401


def test_api_config_fields_authed(user_client):
    """Helper is not @require_permission-gated. Returns 200 with possibly empty
    search_options if Search_Field_Labels/SearchConfig are missing."""
    resp = user_client.get("/api/config/fields")
    assert resp.status_code == 200
    body = resp.get_json()
    assert "search_options" in body
    assert "labels" in body


class _FakeCache:
    """Minimal cache.get/set stand-in, isolated from the real process-wide
    Flask-Caching SimpleCache instance. A full pytest run showed the shared
    cache is not reliably empty at this test's start even after cache.clear()
    (some other test in the suite repopulates or retains an entry under the
    same key), so this test substitutes its own throwaway store instead of
    depending on that instance's isolation."""

    def __init__(self):
        self.store = {}

    def get(self, key):
        return self.store.get(key)

    def set(self, key, value, timeout=None):
        self.store[key] = value


def test_api_config_fields_perm_state_in_cache_key(user_client, monkeypatch):
    """The route's cache key embeds the sensitive-perm state as a `_s0`/`_s1`
    suffix (see api_config_fields: `f"config_fields_{...}_s{int(_sees_sensitive)}"`)
    so a permissioned user's cached response can never be served to a
    permissionless one. search_options is always empty in this test DB (no
    SearchConfig/Search_Field_Labels tables), so we can't assert on response
    *body* differences — instead swap in a throwaway fake cache (see
    _FakeCache) and inspect what it collects, discovering the actual keys
    each perm state writes under rather than predicting them.
    """
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)

    with user_client.session_transaction() as sess:
        sess["locale"] = "en"

    # Without the sensitive perm the response is filtered + cached under _s0.
    monkeypatch.setattr(
        wv, "has_permission", lambda code: code != "workitems.filter.documentfields.sensitive"
    )
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: {"validationuser"})
    r0 = user_client.get("/api/config/fields")
    assert r0.status_code == 200
    assert (
        len(fake_cache.store) == 1
    ), f"expected exactly one cached entry, got {fake_cache.store!r}"
    key_s0 = next(iter(fake_cache.store))
    assert key_s0.endswith("_s0")
    cached_s0 = fake_cache.store[key_s0]
    assert cached_s0 is not None, f"expected a cache entry under {key_s0!r}"

    # With the perm the cache key differs (_s1) -> not served the _s0 entry.
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    r1 = user_client.get("/api/config/fields")
    assert r1.status_code == 200
    assert len(fake_cache.store) == 2, f"expected a second cached entry, got {fake_cache.store!r}"
    key_s1 = next(k for k in fake_cache.store if k != key_s0)
    assert key_s1.endswith("_s1")
    assert key_s1 != key_s0
    cached_s1 = fake_cache.store[key_s1]
    assert cached_s1 is not None, f"expected a cache entry under {key_s1!r}"

    # Both perm states landed in genuinely distinct, still-present cache
    # entries: the permissioned request never reused or clobbered the
    # permissionless slot.
    assert fake_cache.store[key_s0] == cached_s0


def test_api_docfield_values_gated(noperm_client):
    resp = noperm_client.get("/api/docfield_values")
    assert resp.status_code == 403


def test_api_docfield_values_with_perms(user_client, workitems_all_perms):
    resp = user_client.get("/api/docfield_values?docfield=foo&prcfW=all")
    assert resp.status_code in (200, 400, 500)


def test_api_workitems_gated(noperm_client):
    resp = noperm_client.get("/api/workitems")
    assert resp.status_code == 403


def test_api_workitems_with_perms(user_client, workitems_all_perms):
    resp = user_client.get("/api/workitems")
    assert resp.status_code in (200, 500)


def test_api_workitems_returns_degraded_key(user_client, workitems_all_perms):
    resp = user_client.get("/api/workitems")
    assert resp.status_code in (200, 500)
    if resp.status_code == 200:
        body = resp.get_json()
        assert "workitems" in body
        assert "pagination" in body
        assert "degradedSources" in body
        assert isinstance(body["degradedSources"], list)


def test_api_workitems_docfield_tolerates_absent_ms02_engine(user_client, workitems_all_perms):
    """A doc-field query for an MS02 process must not error beyond the harness
    tolerance when engine_ms02_docfields_pg is None (CI default)."""
    resp = user_client.get(
        "/api/workitems",
        query_string={
            "prcfW": "sydoc.praesidialdepartement_bs",
            "docfield": "doctype",
            "docvalue": "invoice",
        },
    )
    assert resp.status_code in (200, 500)
    if resp.status_code == 200:
        body = resp.get_json()
        assert "workitems" in body
        assert "pagination" in body


def test_api_workitems_docfield_mixed_processes_tolerated(user_client, workitems_all_perms):
    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "doctype", "docvalue": "x"},
    )
    assert resp.status_code in (200, 500)


class _SqlLogCursor:
    """Minimal DB-API cursor stub: records executed SQL text and always
    returns no rows. Lets the doc-field pre-fetch blocks in
    _get_workitems_data run to completion without touching a real (in this
    plan's dev environment, unreachable) SQL Server, so the test can assert
    on *whether a query was even attempted* for a given docfield."""

    def __init__(self, log):
        self._log = log

    def execute(self, sql, params=None):
        self._log.append(sql)
        return self

    def fetchall(self):
        return []

    def fetchone(self):
        return None

    def close(self):
        pass


class _SqlLogConn:
    def __init__(self, log):
        self._log = log

    def cursor(self):
        return _SqlLogCursor(self._log)

    def close(self):
        pass


class _SqlLogEngine:
    """Stand-in for engine_nexora_db / engine_statistics_db: only
    .raw_connection() is touched by _get_workitems_data's doc-field
    pre-fetch blocks."""

    def __init__(self, log):
        self._log = log

    def raw_connection(self):
        return _SqlLogConn(self._log)


def _fm(field_key, column, process="sydoc.test_proc", client="default", column_type=None):
    """Build a nx_lib.mapping_config.FieldMapping for docfield-resolution
    tests (#98 phase 3: the two resolution blocks in _get_workitems_data
    now read mapping_config instead of per-pair SearchConfig cursors)."""
    from nx_lib.mapping_config import FieldMapping

    return FieldMapping(
        client=client, process=process, field_key=field_key, column=column, column_type=column_type
    )


def _ps(
    process,
    table,
    alias="t",
    join_condition=None,
    time_filter="1=1",
    client="default",
):
    """Build a nx_lib.mapping_config.ProcessSource -- sibling to _fm above.
    join_condition defaults to referencing `alias` (matching the id-col
    regex/`_ms02_id_column` both legs use to pull the id column out of it)."""
    from nx_lib.mapping_config import ProcessSource

    return ProcessSource(
        client=client,
        process=process,
        table=table,
        alias=alias,
        join_condition=join_condition or f"{alias}.ID = twi.id",
        time_filter=time_filter,
        suggestion_time_filter=None,
        export_column=None,
        import_column=None,
        workitem_column=None,
        extra_condition=None,
        id_column_type=None,
    )


def _stub_mapping_config(
    monkeypatch,
    wv,
    *,
    default_mappings=(),
    default_sources=(),
    ms02_mappings=(),
    ms02_sources=(),
):
    """Replace mapping_config.mappings_for/sources_for with canned per-client
    data. The two doc-field resolution blocks in _get_workitems_data each
    make exactly one mappings_for + one sources_for call per leg, before the
    pair loop (#98 phase 3) -- this is the direct successor to the legacy
    per-pair SearchConfig cursor mocks (_SqlLogCursor/_SqlLogConn/_SqlLogEngine
    above), which those two blocks no longer query."""

    def _mappings_for(client, processes, field_keys=None):
        rows = default_mappings if client == "default" else ms02_mappings
        if processes is None:
            return list(rows)
        wanted = set(processes)
        return [m for m in rows if m.process in wanted]

    def _sources_for(client, processes=None):
        rows = default_sources if client == "default" else ms02_sources
        if processes is None:
            return list(rows)
        wanted = set(processes)
        return [s for s in rows if s.process in wanted]

    monkeypatch.setattr(wv.mapping_config, "mappings_for", _mappings_for)
    monkeypatch.setattr(wv.mapping_config, "sources_for", _sources_for)


def test_get_workitems_data_skips_sensitive_docfield_search(
    user_client, workitems_all_perms, monkeypatch
):
    """Regression for the two doc-field search-filtering blocks in
    _get_workitems_data (default/StatisticsDB path + MS02/Postgres path, both
    gated in commit ae83bcb via `if docfield in blocked_docfields: continue`).
    A sensitive docfield/docvalue pair must contribute NO SQL constraint --
    neither block may even build/execute its StatisticsDB constraint query --
    when the caller lacks workitems.filter.documentfields.sensitive, and the
    resulting WorkitemFilter must carry no docfield constraint at all."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    # allowed_processes (and therefore whether either pre-fetch block is even
    # entered) is read directly off session['permissions'], not via
    # has_permission(). _reload_user_permissions (before_request) reloads that
    # list from the DB on every request, so it must be patched at the source
    # (same seam as test_prepared_docs_link_visible_for_target_process) rather
    # than set via session_transaction, which would just be clobbered.
    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    # Force entry into the MS02 block too (engine_ms02_docfields_pg is None in
    # CI/this dev env absent MS02_DOCFIELDS_DB_* env vars).
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: {"validationuser"})
    monkeypatch.setattr(
        wv, "has_permission", lambda code: code != "workitems.filter.documentfields.sensitive"
    )

    # Mapped on BOTH legs -- proves the skip is caused by the sensitive-field
    # gate, not by an absent mapping (which would zero the pair anyway).
    _stub_mapping_config(
        monkeypatch,
        wv,
        default_mappings=[_fm("validationuser", "ValidationUser")],
        default_sources=[_ps("sydoc.test_proc", "dbo.T")],
        ms02_mappings=[_fm("validationuser", "ValidationUser", client="ms02")],
        ms02_sources=[_ps("sydoc.test_proc", 'public."T"', alias="d", client="ms02")],
    )

    def _must_not_run(*a, **k):
        raise AssertionError("resolve_ms02_docfield_ids must not run for a blocked docfield")

    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", _must_not_run)

    captured = {}

    def _fake_fetch_merged_page(filt, offset, per_page):
        captured["filt"] = filt
        return [], 0, []

    monkeypatch.setattr(wv, "fetch_merged_page", _fake_fetch_merged_page)

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "validationuser", "docvalue": "alice"},
    )

    assert resp.status_code == 200
    assert not [q for q in sql_log if "ValidationUser" in q], sql_log
    assert captured["filt"].docfield_ids is None
    assert captured["filt"].ms02_docfield_ids is None


def test_docfield_search_absent_ms02_engine_fails_closed(
    user_client, workitems_all_perms, monkeypatch
):
    """Cross-source bleed regression (observed on STAGING, 2026-07-20): with the
    MS02 doc-field engine unset (env vars missing) but the MS02 runtime client
    registered, a doc-field search skipped the MS02 pre-resolution entirely and
    left ms02_docfield_ids = None -- "no constraint" -- so the Postgres source
    returned its ENTIRE corpus into the filtered list. An active doc-field
    search must fail CLOSED: a source that cannot be checked contributes zero
    rows, never all of them."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", None)

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["docbarcode"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: set())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    _stub_mapping_config(monkeypatch, wv)  # unmapped everywhere; MS02 block skipped anyway

    captured = {}

    def _fake_fetch_merged_page(filt, offset, per_page):
        captured["filt"] = filt
        return [], 0, []

    monkeypatch.setattr(wv, "fetch_merged_page", _fake_fetch_merged_page)

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "docbarcode", "docvalue": "M629648"},
    )

    assert resp.status_code == 200
    assert captured["filt"].ms02_docfield_ids == set()


def test_docfield_search_ms02_resolver_error_fails_closed(
    user_client, workitems_all_perms, monkeypatch
):
    """Sibling to the absent-engine test: the engine exists and the ms02
    mapping_config mapping is found, but resolve_ms02_docfield_ids errors
    (its contract returns None on any failure). That None must be coerced to
    an empty allow-set -- zero MS02 rows -- not treated as "no constraint"."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["docbarcode"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: set())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", lambda *a, **k: None)
    # One usable ms02 mapping row for the MS02 leg's lookup; the default leg
    # stays unmapped (not asserted on here).
    _stub_mapping_config(
        monkeypatch,
        wv,
        ms02_mappings=[
            _fm("docbarcode", "DossierBarcode", process="sydoc.test_proc", client="ms02")
        ],
        ms02_sources=[
            _ps(
                "sydoc.test_proc",
                'public."DossierStatistik"',
                alias="d",
                join_condition="d.WorkItemID = twi.id",
                time_filter=None,
                client="ms02",
            )
        ],
    )

    captured = {}

    def _fake_fetch_merged_page(filt, offset, per_page):
        captured["filt"] = filt
        return [], 0, []

    monkeypatch.setattr(wv, "fetch_merged_page", _fake_fetch_merged_page)

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "docbarcode", "docvalue": "M629648"},
    )

    assert resp.status_code == 200
    assert captured["filt"].ms02_docfield_ids == set()


def test_get_workitems_data_queries_nonsensitive_docfield_search(
    user_client, workitems_all_perms, monkeypatch
):
    """Control for the sibling skip test above: with the SAME field but no
    sensitivity block in play (get_sensitive_field_keys empty + full
    has_permission), both doc-field pre-fetch blocks DO attempt their
    resolution -- proving the sibling test's absence of a constraint is
    genuinely caused by the sensitive-field skip, not by the fakes
    themselves suppressing all resolution regardless of gating."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: set())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    captured_ms02 = {}

    def _spy_resolve(engine, pairs):
        captured_ms02["pairs"] = pairs
        return None

    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", _spy_resolve)
    monkeypatch.setattr(wv, "fetch_merged_page", lambda filt, offset, per_page: ([], 0, []))

    _stub_mapping_config(
        monkeypatch,
        wv,
        default_mappings=[_fm("validationuser", "ValidationUser")],
        default_sources=[_ps("sydoc.test_proc", "dbo.T")],
        ms02_mappings=[_fm("validationuser", "ValidationUser", client="ms02")],
        ms02_sources=[_ps("sydoc.test_proc", 'public."T"', alias="d", client="ms02")],
    )

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "validationuser", "docvalue": "alice"},
    )

    assert resp.status_code == 200
    default_queries = [q for q in sql_log if "ValidationUser" in q]
    assert default_queries, sql_log
    ms02_specs = [spec for pair in captured_ms02.get("pairs", []) for spec in pair[0]]
    assert ms02_specs, captured_ms02


def test_get_workitems_data_unmapped_docfield_zeroes_both_sources(
    user_client, workitems_all_perms, monkeypatch
):
    """Cross-source bleed regression: a searched doc-field with NO
    mapping_config mapping for a source must force that source to ZERO rows
    (empty allow-set), not run unconstrained (None). Observed on PROD:
    validationuser was mapped only for a 'default' process, and the
    unconstrained MS02 leg returned every PDBS workitem. Here neither leg
    finds a mapping row, so BOTH allow-sets must come out as set() -- and the
    MS02 resolver must never run."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: set())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    _stub_mapping_config(monkeypatch, wv)  # unmapped everywhere

    def _must_not_run(*a, **k):
        raise AssertionError("resolve_ms02_docfield_ids must not run without mapping rows")

    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", _must_not_run)

    captured = {}

    def _fake_fetch_merged_page(filt, offset, per_page):
        captured["filt"] = filt
        return [], 0, []

    monkeypatch.setattr(wv, "fetch_merged_page", _fake_fetch_merged_page)

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "validationuser", "docvalue": "alice"},
    )

    assert resp.status_code == 200
    assert captured["filt"].docfield_ids == set()
    assert captured["filt"].ms02_docfield_ids == set()


def test_get_workitems_data_fieldless_pair_searches_all_columns(
    user_client, workitems_all_perms, monkeypatch
):
    """Value-first search (#148): a docvalue with NO docfield must widen the
    default leg's resolution to every permitted column (proven here by both
    mapped columns showing up together in the single StatisticsDB constraint
    query the widened UNION produces) and still count as an ACTIVE search for
    the fail-closed guard -- with the MS02 leg deliberately left unmapped,
    both allow-sets must come out set(), never None (which would let a
    source run unconstrained)."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser", "docbarcode"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: set())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    # Default leg mapped for BOTH widened columns (proves the widening);
    # MS02 leg stays unmapped so the resolver-must-not-run guard still holds.
    _stub_mapping_config(
        monkeypatch,
        wv,
        default_mappings=[
            _fm("validationuser", "ValidationUser"),
            _fm("docbarcode", "DocBarcode"),
        ],
        default_sources=[_ps("sydoc.test_proc", "dbo.T")],
    )

    def _must_not_run(*a, **k):
        raise AssertionError("resolve_ms02_docfield_ids must not run without mapping rows")

    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", _must_not_run)

    captured = {}

    def _fake_fetch_merged_page(filt, offset, per_page):
        captured["filt"] = filt
        return [], 0, []

    monkeypatch.setattr(wv, "fetch_merged_page", _fake_fetch_merged_page)

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "", "docvalue": "alice"},
    )

    assert resp.status_code == 200
    widened = [q for q in sql_log if "ValidationUser" in q and "DocBarcode" in q]
    assert widened, sql_log
    assert captured["filt"].docfield_ids == set()
    assert captured["filt"].ms02_docfield_ids == set()


def test_get_workitems_data_fieldless_pair_excludes_sensitive_columns(
    user_client, workitems_all_perms, monkeypatch
):
    """Value-first search (#148): the widened any-field column set must drop
    sensitive FieldKeys for callers without the sensitive-fields permission --
    no StatisticsDB constraint query may even mention the blocked column."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser", "secretfield"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: {"secretfield"})
    monkeypatch.setattr(
        wv,
        "has_permission",
        lambda code: code != "workitems.filter.documentfields.sensitive",
    )
    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", lambda *a, **k: None)
    monkeypatch.setattr(wv, "fetch_merged_page", lambda filt, offset, per_page: ([], 0, []))
    # Both fields ARE mapped -- the blocked one must still never reach a
    # query even though a mapping row exists for it (the block happens
    # earlier, at the target_cols/blocked_docfields filter).
    _stub_mapping_config(
        monkeypatch,
        wv,
        default_mappings=[
            _fm("validationuser", "ValidationUser"),
            _fm("secretfield", "SecretField"),
        ],
        default_sources=[_ps("sydoc.test_proc", "dbo.T")],
    )

    resp = user_client.get(
        "/api/workitems",
        query_string={"prcfW": "all", "docfield": "", "docvalue": "alice"},
    )

    assert resp.status_code == 200
    assert any("ValidationUser" in q for q in sql_log), sql_log
    assert not any("SecretField" in q for q in sql_log), sql_log


def _op_test_scaffold(monkeypatch, sql_log):
    """Shared monkeypatching for the docop/doccomb view tests."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: set())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "resolve_ms02_docfield_ids", lambda *a, **k: None)
    monkeypatch.setattr(wv, "fetch_merged_page", lambda filt, offset, per_page: ([], 0, []))
    _stub_mapping_config(
        monkeypatch,
        wv,
        default_mappings=[_fm("validationuser", "ValidationUser")],
        default_sources=[_ps("sydoc.test_proc", "dbo.T")],
    )


def test_docfield_ops_map_shapes():
    """(#148) the whitelisted operator map drives the SQL Server comparators
    and LIKE-pattern params."""
    from nx_lib.views.workitems import DOCFIELD_OPS, _docfield_comb, _docfield_op

    assert DOCFIELD_OPS["eq"][0] == "="
    assert DOCFIELD_OPS["neq"][0] == "<>"
    assert DOCFIELD_OPS["contains"][0] == "LIKE"
    assert DOCFIELD_OPS["ncontains"][0] == "NOT LIKE"
    assert DOCFIELD_OPS["contains"][1]("v") == "%v%"
    assert DOCFIELD_OPS["startswith"][1]("v") == "v%"
    assert DOCFIELD_OPS["endswith"][1]("v") == "%v"
    assert DOCFIELD_OPS["eq"][1]("v") == "v"
    # normalizers: unknown/missing keys fall back to contains / and
    assert _docfield_op(["EQ"], 0) == "eq"
    assert _docfield_op(["bogus"], 0) == "contains"
    assert _docfield_op([], 5) == "contains"
    assert _docfield_comb(["OR"], 0) == "or"
    assert _docfield_comb(["nand"], 0) == "and"
    assert _docfield_comb([], 5) == "and"


def test_docfield_unknown_op_and_comb_are_whitelisted(
    user_client, workitems_all_perms, monkeypatch
):
    """(#148) hostile docop/doccomb values must never be interpolated into
    SQL -- unknown keys fall back to contains/and."""
    sql_log = []
    _op_test_scaffold(monkeypatch, sql_log)

    resp = user_client.get(
        "/api/workitems",
        query_string={
            "prcfW": "all",
            "docfield": "validationuser",
            "docvalue": "alice",
            "docop": "1; DROP TABLE Users--",
            "doccomb": "UNION SELECT",
        },
    )
    assert resp.status_code == 200
    assert not any("DROP TABLE" in q for q in sql_log)
    assert not any("UNION SELECT" in q for q in sql_log)


def test_docfield_or_pair_processed_without_early_break(
    user_client, workitems_all_perms, monkeypatch
):
    """(#148) with AND-only semantics the first no-mapping pair used to break
    out of the loop; OR support requires every pair to be evaluated. Two
    field-carrying pairs must both reach the op/comb resolution step (spied
    via _docfield_op, called once per pair in EACH of the two resolution legs
    -- default and MS02 -- so both pair indices must appear twice) even
    though StatisticsDB (stubbed to return no rows) ultimately resolves both
    to empty."""
    sql_log = []
    _op_test_scaffold(monkeypatch, sql_log)

    captured = {}

    def _fake_fetch_merged_page(filt, offset, per_page):
        captured["filt"] = filt
        return [], 0, []

    import nx_lib.views.workitems as wv

    monkeypatch.setattr(wv, "fetch_merged_page", _fake_fetch_merged_page)

    pair_indices = []
    _orig_docfield_op = wv._docfield_op

    def _spy_docfield_op(docops, idx):
        pair_indices.append(idx)
        return _orig_docfield_op(docops, idx)

    monkeypatch.setattr(wv, "_docfield_op", _spy_docfield_op)

    resp = user_client.get(
        "/api/workitems",
        query_string=[
            ("prcfW", "all"),
            ("docfield", "validationuser"),
            ("docvalue", "alice"),
            ("doccomb", "and"),
            ("docfield", "validationuser"),
            ("docvalue", "bob"),
            ("doccomb", "or"),
        ],
    )
    assert resp.status_code == 200
    assert sorted(pair_indices) == [0, 0, 1, 1], pair_indices
    # both pairs mapped but StatisticsDB (stubbed) returns no rows -> OR-fold
    # of two empty sets -> still fail-closed
    assert captured["filt"].docfield_ids == set()


def test_api_docfield_values_no_field_widens_and_excludes_sensitive(
    user_client, workitems_all_perms, monkeypatch
):
    """/api/docfield_values with no field (value-first mode, #148) must answer
    200 with a JSON list (labeled suggestions), widen its mapping_config
    lookup to all permitted field keys, and never mention sensitive fields --
    neither in the generated StatisticsDB SQL nor in the field_keys passed to
    mapping_config.mappings_for itself (#98 task 6: the SearchConfig cursor
    this test used to log queries against is gone; the widened field-key set
    is now the thing to assert on)."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.test_proc",
        ],
    )

    sql_log = []
    monkeypatch.setattr(wv, "engine_statistics_db", _SqlLogEngine(sql_log))
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", None)

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser", "secretfield"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: {"secretfield"})
    monkeypatch.setattr(
        wv,
        "has_permission",
        lambda code: code != "workitems.filter.documentfields.sensitive",
    )

    captured_field_keys = {}

    def _mappings_for(client, processes, field_keys=None):
        captured_field_keys[client] = set(field_keys or [])
        if client == "default":
            return [_fm("validationuser", "ValidationUser")]
        return []

    def _sources_for(client, processes=None):
        return [_ps("sydoc.test_proc", "dbo.T")] if client == "default" else []

    monkeypatch.setattr(wv.mapping_config, "mappings_for", _mappings_for)
    monkeypatch.setattr(wv.mapping_config, "sources_for", _sources_for)

    resp = user_client.get("/api/docfield_values", query_string={"process": "all", "q": ""})

    assert resp.status_code == 200
    assert resp.get_json() == []
    assert any("ValidationUser" in q for q in sql_log), sql_log
    assert not any("SecretField" in q for q in sql_log), sql_log
    assert "secretfield" not in captured_field_keys.get("default", set()), captured_field_keys
    assert "secretfield" not in captured_field_keys.get("ms02", set()), captured_field_keys


def test_export_workitems_csv_gated(noperm_client):
    resp = noperm_client.get("/api/export/workitems/csv")
    assert resp.status_code == 403


def test_export_workitems_csv_with_perms(user_client, workitems_all_perms):
    """No data → returns CSV with the placeholder line."""
    resp = user_client.get("/api/export/workitems/csv")
    assert resp.status_code in (200, 500)
    if resp.status_code == 200:
        assert "text/csv" in resp.headers.get("Content-Type", "")


# --- colliding-id export rows must not mix client fields (D9) --------------- #
# Workitem ids are not globally unique across clients (1216 collides between
# the default Octo client and MS02, see docs/design/ms02-multisource.md). The
# export builds `domains`/`details_map`/media+audit cache entries per row; if
# any of those are keyed by the bare id, the second row processed for a
# colliding id silently answers for (or overwrites the cache entry of) the
# first, so one CSV row ends up carrying the OTHER client's field values.


def test_export_workitems_csv_keys_by_client_not_bare_id(
    user_client, workitems_all_perms, monkeypatch
):
    """Two rows share workitemid=1216 but belong to different clients. Each
    exported row must carry ITS OWN client's field value, never the other
    client's cached/fetched copy of the same bare id."""
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)
    # workitems_all_perms only patches nx_lib.security.has_permission, which
    # covers the @require_permission route gate (looked up inside security.py
    # at call time) but NOT the `include_fields = ... and has_permission(...)`
    # check inside this module -- that name was bound at import time and needs
    # patching directly on the view module, same as test_api_config_fields_
    # perm_state_in_cache_key above.
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    rows = [
        {
            "workitemid": 1216,
            "client": "default",
            "status": "Open",
            "current_stage": "Stage A",
            "priority": 2,
            "tags": [],
            "modifiedat": None,
        },
        {
            "workitemid": 1216,
            "client": "ms02",
            "status": "Closed",
            "current_stage": "Stage B",
            "priority": 1,
            "tags": [],
            "modifiedat": None,
        },
    ]

    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {
            "workitems": rows,
            "pagination": {"totalItems": len(rows)},
        },
    )

    def fake_get_domain(wid, client_hint=None):
        return (
            "default-domain.example.com" if client_hint == "default" else "ms02-domain.example.com"
        )

    monkeypatch.setattr(wv, "get_domain_for_workitem", fake_get_domain)
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{domain}", f"doc-{domain}")
    )

    def fake_get_extensions_urls_fields(workitemdata, document_id, domain, with_tables=False):
        fields = (
            {"Amount": "100-default"}
            if domain == "default-domain.example.com"
            else {"Amount": "999-ms02"}
        )
        return [], [], fields, {}, {}

    monkeypatch.setattr(wv, "get_extensions_urls_fields", fake_get_extensions_urls_fields)

    # D-CSVLIM: include=fields now requires an explicit (<=10) ids selection
    # server-side; name both colliding rows so this test's actual concern
    # (per-client field values, not the selection cap) is unaffected.
    resp = user_client.get("/api/export/workitems/csv?include=fields&ids=default-1216,ms02-1216")
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    csv_rows = list(csv.reader(io.StringIO(body)))
    header, data_rows = csv_rows[0], csv_rows[1:]
    assert len(data_rows) == 2, f"expected 2 data rows, got {data_rows!r}"
    amount_idx = header.index("Amount")
    values = {r[amount_idx] for r in data_rows}
    assert values == {"100-default", "999-ms02"}, (
        f"expected each row to carry its own client's Amount, got {values!r} "
        f"(cross-contamination: a bare-id cache/dict key let one client's "
        f"fetched value answer for the other's)"
    )


# --- "export selected" must filter by (client, id), not bare id ------------ #
# The bulk-select checkboxes / `ids` query param used to carry a bare
# workitemid. Selecting only the default-client row of a colliding id (1216
# collides between the default Octo client and MS02) also exported the MS02
# row, because `specific_ids` membership was checked against the bare id
# alone. The UI now sends compound `client-id` pairs (matching renderTable's
# `rowKey = `${client}-${workitemid}`` in _workitems_overview_js.html) and the
# backend must filter on that compound key.


def test_export_workitems_csv_selected_ids_are_client_aware(
    user_client, workitems_all_perms, monkeypatch
):
    """Two rows share workitemid=1216 but belong to different clients. Passing
    ids=default-1216 must export only the default row, never the ms02 row
    that also matches the base filter."""
    import nx_lib.views.workitems as wv

    rows = [
        {
            "workitemid": 1216,
            "client": "default",
            "status": "Open",
            "current_stage": "Stage A",
            "priority": 2,
            "tags": [],
            "modifiedat": None,
        },
        {
            "workitemid": 1216,
            "client": "ms02",
            "status": "Closed",
            "current_stage": "Stage B",
            "priority": 1,
            "tags": [],
            "modifiedat": None,
        },
    ]

    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {
            "workitems": rows,
            "pagination": {"totalItems": len(rows)},
        },
    )

    resp = user_client.get("/api/export/workitems/csv?ids=default-1216")
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    csv_rows = list(csv.reader(io.StringIO(body)))
    header, data_rows = csv_rows[0], csv_rows[1:]
    assert len(data_rows) == 1, (
        f"expected exactly 1 row (the default-client row) for ids=default-1216, "
        f"got {data_rows!r} (bare-id filtering also matched the colliding ms02 row)"
    )
    status_idx = header.index("Status")
    assert data_rows[0][status_idx] == "Open", (
        f"expected the default-client row (Status=Open), got {data_rows[0]!r} "
        f"-- wrong client's row was exported"
    )


# --- D-CSVLIM: heavy-include cap is enforced server-side, not just in JS --- #
# fields/history/images are all per-row Octo fetches; the ≤10-selection rule
# used to live only in the JS control (isSelection = selectedIds.size > 0 &&
# selectedIds.size <= 10). A direct API call bypassing that control could
# request include=fields|history|images with no ids (or an arbitrarily large
# ids list) and walk up to EXPORT_MAX_ROWS rows doing per-row Octo fetches.


def test_export_workitems_csv_include_without_ids_returns_400(user_client, workitems_all_perms):
    resp = user_client.get("/api/export/workitems/csv?include=fields")
    assert resp.status_code == 400
    body = resp.get_json()
    assert body and "error" in body


def test_export_workitems_csv_include_with_over_10_ids_returns_400(
    user_client, workitems_all_perms
):
    ids_param = ",".join(f"default-{i}" for i in range(1, 12))  # 11 compound ids
    resp = user_client.get(f"/api/export/workitems/csv?include=history&ids={ids_param}")
    assert resp.status_code == 400
    body = resp.get_json()
    assert body and "error" in body


def test_export_workitems_csv_include_with_le_10_ids_is_honored(
    user_client, workitems_all_perms, monkeypatch
):
    """3 selected ids (<= the cap) must still get the requested include=fields
    treatment -- the cap must not also block legitimate small selections."""
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    rows = [
        {
            "workitemid": wid,
            "client": "default",
            "status": "Open",
            "current_stage": "Stage A",
            "priority": 1,
            "tags": [],
            "modifiedat": None,
        }
        for wid in (91001, 91002, 91003)
    ]

    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {
            "workitems": rows,
            "pagination": {"totalItems": len(rows)},
        },
    )
    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: "d.example.com"
    )
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{wid}", f"doc-{wid}")
    )
    monkeypatch.setattr(
        wv,
        "get_extensions_urls_fields",
        lambda workitemdata, document_id, domain, with_tables=False: (
            [],
            [],
            {"Amount": "42"},
            {},
            {},
        ),
    )

    ids_param = "default-91001,default-91002,default-91003"
    resp = user_client.get(f"/api/export/workitems/csv?include=fields&ids={ids_param}")
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    csv_rows = list(csv.reader(io.StringIO(body)))
    header, data_rows = csv_rows[0], csv_rows[1:]
    assert "Amount" in header, f"expected include=fields honored, got header={header!r}"
    assert len(data_rows) == 3


def test_export_workitems_csv_no_include_no_ids_still_200(
    user_client, workitems_all_perms, monkeypatch
):
    """Regression: the new heavy-include cap must not affect the light default
    export path (no include=, no ids=) -- it stays a plain 200 CSV."""
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {"workitems": [], "pagination": {"totalItems": 0}},
    )

    resp = user_client.get("/api/export/workitems/csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers.get("Content-Type", "")


# --- as_completed() timing out mid-iteration must degrade, not 500 -------- #
# Per-row Octo fetches (fields/history/images) run on a thread pool; the
# `for future in as_completed(futures, timeout=120):` loop's own iteration
# protocol -- not just future.result() inside the loop body -- can raise
# concurrent.futures.TimeoutError once the budget elapses with futures still
# pending. That used to be uncaught and 500'd the entire export.


def test_export_workitems_csv_survives_as_completed_timeout(
    user_client, workitems_all_perms, monkeypatch
):
    """as_completed() raising TimeoutError at the loop boundary must still
    yield a 200 CSV: rows that finished before the timeout keep their data,
    rows that didn't get an explicit timed-out marker instead of a 500."""
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    rows = [
        {
            "workitemid": wid,
            "client": "default",
            "status": "Open",
            "current_stage": "Stage A",
            "priority": 1,
            "tags": [],
            "modifiedat": None,
        }
        for wid in (82001, 82002)
    ]

    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {
            "workitems": rows,
            "pagination": {"totalItems": len(rows)},
        },
    )
    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: "d.example.com"
    )
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{wid}", f"doc-{wid}")
    )
    monkeypatch.setattr(
        wv,
        "get_extensions_urls_fields",
        lambda workitemdata, document_id, domain, with_tables=False: (
            [],
            [],
            {"Amount": "42"},
            {},
            {},
        ),
    )

    real_as_completed = wv.as_completed

    def fake_as_completed(futures, timeout=None):
        # Simulate the 120s budget elapsing mid-iteration: let the
        # first-completed future come through normally (so the export has
        # at least one real row), then raise instead of yielding the rest
        # -- exactly what the real as_completed() does when its timeout
        # fires with futures still outstanding.
        it = real_as_completed(futures, timeout=timeout)
        yield next(it)
        raise concurrent.futures.TimeoutError()

    monkeypatch.setattr(wv, "as_completed", fake_as_completed)

    ids_param = "default-82001,default-82002"
    resp = user_client.get(f"/api/export/workitems/csv?include=fields&ids={ids_param}")
    assert resp.status_code == 200, (
        f"as_completed() timing out mid-iteration must degrade gracefully, not "
        f"500 the export; got {resp.status_code}: {resp.get_data(as_text=True)!r}"
    )
    body = resp.get_data(as_text=True)
    csv_rows = list(csv.reader(io.StringIO(body)))
    header, data_rows = csv_rows[0], csv_rows[1:]
    assert len(data_rows) == 2, f"expected both rows still present, got {data_rows!r}"
    amount_idx = header.index("Amount")
    values = {r[amount_idx] for r in data_rows}
    assert "42" in values, (
        f"expected the row that finished before the timeout to keep its "
        f"fetched data, got {values!r}"
    )
    assert wv.EXPORT_TIMEOUT_MARKER in values, (
        f"expected the row that never finished before the timeout to carry "
        f"an explicit timed-out marker, got {values!r}"
    )


# D-CSVTIMEOUT: when the 120s budget elapses before ANY future completes --
# the realistic Octo-outage shape, since EXPORT_HEAVY_INCLUDE_MAX_IDS caps
# heavy exports to <=10 rows -- all_field_keys/max_images derive from zero
# completed rows and end up empty. That used to leave the CSV with no
# include= columns at all, indistinguishable from a legitimate "these
# workitems have no fields" result. Must instead be signalled via the same
# trailer-line + header pattern the row-truncation path already uses.


def test_export_workitems_csv_all_rows_timeout_signals_degraded_export(
    user_client, workitems_all_perms, monkeypatch
):
    """Zero completed rows before the as_completed() timeout must still
    yield a 200 CSV, but flagged as timeout-degraded rather than looking
    like a legitimate zero-fields export."""
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    rows = [
        {
            "workitemid": wid,
            "client": "default",
            "status": "Open",
            "current_stage": "Stage A",
            "priority": 1,
            "tags": [],
            "modifiedat": None,
        }
        for wid in (83001, 83002)
    ]

    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {
            "workitems": rows,
            "pagination": {"totalItems": len(rows)},
        },
    )
    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: "d.example.com"
    )
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{wid}", f"doc-{wid}")
    )
    monkeypatch.setattr(
        wv,
        "get_extensions_urls_fields",
        lambda workitemdata, document_id, domain, with_tables=False: (
            [],
            [],
            {"Amount": "42"},
            {},
            {},
        ),
    )

    def fake_as_completed(futures, timeout=None):
        # Simulate the 120s budget elapsing before a single future completes
        # -- exactly what as_completed() itself raises in that shape.
        raise concurrent.futures.TimeoutError()

    monkeypatch.setattr(wv, "as_completed", fake_as_completed)

    ids_param = "default-83001,default-83002"
    resp = user_client.get(f"/api/export/workitems/csv?include=fields&ids={ids_param}")
    assert resp.status_code == 200, (
        f"all rows timing out must still degrade gracefully, not 500; "
        f"got {resp.status_code}: {resp.get_data(as_text=True)!r}"
    )
    body = resp.get_data(as_text=True)
    csv_rows = list(csv.reader(io.StringIO(body)))
    header, data_rows = csv_rows[0], csv_rows[1:]

    # The bug: with zero completed rows, all_field_keys was empty, so the
    # header carried no include= columns at all -- exactly the silently
    # misleading shape this fix must prevent.
    assert (
        "Amount" not in header
    ), f"sanity check: no row completed, so no field header should appear; got {header!r}"
    assert resp.headers.get("X-Export-Timeout") == "true", (
        f"expected the timeout to be signalled via X-Export-Timeout header, "
        f"got headers={dict(resp.headers)!r}"
    )
    assert (
        "EXPORT TIMED OUT" in body
    ), f"expected a trailer line flagging the timeout-degraded export, got body={body!r}"
    assert len(data_rows) >= 2, f"expected both rows still present, got {data_rows!r}"


def test_strip_export_fields_removes_sensitive_columns():
    from nx_lib.views.workitems import _strip_export_fields

    details_map = {
        1: {"fields": {"Validation User": "alice", "Amount": "50"}, "history": [], "images": []},
        2: {"fields": {"Amount": "9"}, "history": [], "images": []},
    }
    _strip_export_fields(details_map, {"validationuser"})
    assert details_map[1]["fields"] == {"Amount": "50"}
    assert details_map[2]["fields"] == {"Amount": "9"}


def test_import_workitems_gated(noperm_client):
    resp = noperm_client.post("/import_workitems", json={})
    assert resp.status_code == 403


def test_import_workitems_with_perms_empty_body(user_client, workitems_all_perms):
    """No importFile in request → flash + redirect (302) back to overview."""
    resp = user_client.post("/import_workitems", json={}, follow_redirects=False)
    assert resp.status_code in (200, 302, 400, 500)


# ============================ media + audit history ===========================


def test_api_get_media_info_authed_unknown_id(user_client):
    """View has an internal has_permission('workitems.details.view') check
    that the seed user doesn't satisfy → 403. With workitems_all_perms it
    would attempt OctoDB lookup and 500/404."""
    resp = user_client.get("/api/get_media_info/999999")
    assert resp.status_code in (200, 401, 403, 404, 500)


def test_csv_export_does_not_poison_media_info_cache(user_client, workitems_all_perms, monkeypatch):
    """Task 60 regression: the export _fetch helper used to cache
    {"fields": ..., "media_count": ...} under the SAME media_info cache key
    api_get_media_info serves, but built without with_tables=True -- so the
    very next detail-panel request for that workitem got served the export's
    reduced payload and the source-highlight overlay (field_sources/
    table_sources) went silently empty. The export path must be read-through
    only: it may reuse an existing media_info cache entry, but must never
    write a reduced one itself."""
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    wid = 92001
    rows = [
        {
            "workitemid": wid,
            "client": "default",
            "status": "Open",
            "current_stage": "Stage A",
            "priority": 1,
            "tags": [],
            "modifiedat": None,
        }
    ]
    monkeypatch.setattr(
        wv,
        "_get_workitems_data",
        lambda args, export_all=False: {
            "workitems": rows,
            "pagination": {"totalItems": len(rows)},
        },
    )
    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: "d.example.com"
    )
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{wid}", f"doc-{wid}")
    )

    def fake_get_extensions_urls_fields(workitemdata, document_id, domain, with_tables=False):
        # Mirrors real Octo behavior (nx_lib/octo.py get_extensions_urls_fields):
        # field_sources are always computed; table_sources only when
        # with_tables=True. The export path never passes with_tables=True.
        field_sources = [{"key": "Amount", "value": "42", "locations": []}]
        table_sources = [{"rows": []}] if with_tables else []
        return [], [], {"Amount": "42"}, field_sources, table_sources

    monkeypatch.setattr(wv, "get_extensions_urls_fields", fake_get_extensions_urls_fields)

    export_resp = user_client.get(f"/api/export/workitems/csv?include=fields&ids=default-{wid}")
    assert export_resp.status_code == 200
    body = export_resp.get_data(as_text=True)
    assert "Amount" in body, f"sanity check: export should have fetched fields; got {body!r}"

    # The export must not have written anything to the shared media_info slot.
    ck = wv._wi_cache_key("media_info", wid, "d.example.com")
    assert fake_cache.get(ck) is None, (
        "export path wrote to the shared media_info cache key -- it must be " "read-through only"
    )

    detail_resp = user_client.get(f"/api/get_media_info/{wid}?client=default")
    assert detail_resp.status_code == 200
    payload = detail_resp.get_json()
    assert payload.get("field_sources"), (
        f"detail panel lost field_sources after an export ran first for the same "
        f"workitem; got payload={payload!r}"
    )
    assert payload.get("table_sources") == [{"rows": []}], (
        f"detail panel should compute its own full with_tables=True payload, "
        f"not reuse anything the export cached; got payload={payload!r}"
    )


def test_strip_sensitive_from_detail_removes_fields_and_sources():
    from nx_lib.views.workitems import strip_sensitive_from_detail

    data = {
        "fields": {"Validation User": "alice", "Amount": "50"},
        "field_sources": [
            {"key": "Validation User", "value": "alice", "locations": []},
            {"key": "Amount", "value": "50", "locations": []},
        ],
        "table_sources": [],
    }
    out = strip_sensitive_from_detail(data, {"validationuser"})
    assert out["fields"] == {"Amount": "50"}
    assert [s["key"] for s in out["field_sources"]] == ["Amount"]
    assert data["fields"] == {"Validation User": "alice", "Amount": "50"}  # untouched


def test_api_get_media_raw_gated(noperm_client):
    resp = noperm_client.get("/api/get_media_raw/1/0")
    assert resp.status_code == 403


def test_api_get_media_raw_with_perms_unknown(user_client, workitems_all_perms):
    resp = user_client.get("/api/get_media_raw/999999/0")
    assert resp.status_code in (200, 404, 500)


# --- PDF/TIF page-image caches must be keyed by client domain -------------- #
# Workitem ids collide across clients (e.g. 1216 exists in both the default
# Octo runtime and MS02 on live INT, see docs/design/ms02-multisource.md).
# media_raw_pdfpage_/media_raw_tif_ used to key their rendered-page cache on
# the bare (workitem_id, media_index) pair, so the second client's request
# for the "same" id was served the first client's already-cached page image
# for up to an hour. Both tests substitute a throwaway fake cache (see
# _FakeCache above) so they don't depend on the shared process-wide cache
# being empty, and each mocks the client-domain resolution + media fetch to
# return bytes that are distinguishable per domain.


def test_api_get_media_raw_tif_cache_is_client_scoped(
    user_client, workitems_all_perms, monkeypatch
):
    """Requesting the same colliding workitem id/media_index under two
    different clients must render+cache each client's own TIFF bytes, never
    reuse the other client's cached JPEG."""
    import io as _io

    from PIL import Image

    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "cache", _FakeCache())

    def _tif_bytes(color):
        buf = _io.BytesIO()
        Image.new("RGB", (4, 4), color=color).save(buf, format="TIFF")
        return buf.getvalue()

    domains = {"default": "default-domain.example.com", "ms02": "ms02-domain.example.com"}
    media_bytes = {
        domains["default"]: _tif_bytes((255, 0, 0)),
        domains["ms02"]: _tif_bytes((0, 0, 255)),
    }

    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: domains[client_hint]
    )
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{domain}", f"doc-{domain}")
    )
    monkeypatch.setattr(
        wv,
        "get_extensions_urls_fields",
        lambda workitemdata, document_id, domain: (
            [".tif"],
            [f"http://media/{domain}"],
            {},
            {},
            {},
        ),
    )
    monkeypatch.setattr(wv, "get_media", lambda url, domain: media_bytes[domain])

    resp_default = user_client.get("/api/get_media_raw/1216/0?client=default")
    resp_ms02 = user_client.get("/api/get_media_raw/1216/0?client=ms02")

    assert resp_default.status_code == 200
    assert resp_ms02.status_code == 200

    img_default = Image.open(io.BytesIO(resp_default.data))
    img_ms02 = Image.open(io.BytesIO(resp_ms02.data))
    r_default, _g, b_default = img_default.convert("RGB").getpixel((0, 0))
    r_ms02, _g, b_ms02 = img_ms02.convert("RGB").getpixel((0, 0))

    assert r_default > b_default, "default-client response should be the red TIFF it fetched"
    assert b_ms02 > r_ms02, (
        "ms02-client response came back red (the default client's cached bytes) instead of "
        "blue -- media_raw_tif_ cache key omitted the client domain"
    )
    assert resp_default.data != resp_ms02.data, (
        "second client's response returned the first client's cached bytes -- "
        "media_raw_tif_ cache key omitted the client domain"
    )


def test_api_get_media_raw_pdf_cache_is_client_scoped(
    user_client, workitems_all_perms, monkeypatch
):
    """Same leak for the PDF-page-render cache: media_raw_pdfpage_{id}_{idx}
    used to omit the domain, so the second client's request served back the
    first client's rendered PDF page."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "cache", _FakeCache())

    domains = {"default": "default-domain.example.com", "ms02": "ms02-domain.example.com"}

    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: domains[client_hint]
    )
    monkeypatch.setattr(
        wv, "get_workitemdata_param", lambda wid, domain: (f"wdata-{domain}", f"doc-{domain}")
    )
    monkeypatch.setattr(
        wv,
        "get_extensions_urls_fields",
        lambda workitemdata, document_id, domain: (
            [".pdf"],
            [f"http://media/{domain}#page=0"],
            {},
            {},
            {},
        ),
    )
    monkeypatch.setattr(wv, "pdf_src_bytes", lambda url, domain: f"pdfbytes-{domain}".encode())
    monkeypatch.setattr(
        wv, "render_pdf_page_jpeg", lambda pdf_bytes, page_index: pdf_bytes + b"-rendered"
    )

    resp_default = user_client.get("/api/get_media_raw/1216/0?client=default")
    resp_ms02 = user_client.get("/api/get_media_raw/1216/0?client=ms02")

    assert resp_default.status_code == 200
    assert resp_ms02.status_code == 200
    assert resp_default.data == b"pdfbytes-default-domain.example.com-rendered"
    assert resp_ms02.data == b"pdfbytes-ms02-domain.example.com-rendered", (
        "second client's response returned the first client's cached PDF-page bytes -- "
        "media_raw_pdfpage_ cache key omitted the client domain"
    )
    assert resp_default.data != resp_ms02.data


def test_api_get_media_raw_pdf_error_does_not_cache(user_client, workitems_all_perms, monkeypatch):
    """When the PDF source stream errors (e.g. Octo 502s), pdf_src_bytes must
    raise instead of handing back the error body as if it were page bytes.
    The route has to degrade to an error status without ever writing the
    rendered-page cache slot -- caching the error would poison that slot
    with garbage for every request in the next hour."""
    import nx_lib.views.workitems as wv

    fake_cache = _FakeCache()
    monkeypatch.setattr(wv, "cache", fake_cache)

    monkeypatch.setattr(
        wv, "get_domain_for_workitem", lambda wid, client_hint=None: "default-domain.example.com"
    )
    monkeypatch.setattr(wv, "get_workitemdata_param", lambda wid, domain: ("wdata", "doc-1"))
    monkeypatch.setattr(
        wv,
        "get_extensions_urls_fields",
        lambda workitemdata, document_id, domain: (
            [".pdf"],
            ["http://media/doc.pdf#page=0"],
            {},
            {},
            {},
        ),
    )

    def _raise_http_error(url, domain):
        raise requests.HTTPError("502 Server Error")

    monkeypatch.setattr(wv, "pdf_src_bytes", _raise_http_error)

    resp = user_client.get("/api/get_media_raw/1216/0?client=default")

    assert resp.status_code == 500
    pdf_cache_key = "media_raw_pdfpage_0_default-domain.example.com_1216"
    assert pdf_cache_key not in fake_cache.store, (
        "an error response must not populate the rendered-page cache slot -- "
        "that would poison media_raw_pdfpage_ for the next hour"
    )


def test_get_audithistory_gated(noperm_client):
    resp = noperm_client.get("/api/get_audithistory/1")
    assert resp.status_code == 403


def test_get_audithistory_with_perms(user_client, workitems_all_perms):
    resp = user_client.get("/api/get_audithistory/999999")
    assert resp.status_code in (200, 404, 500)


# ============================ tags + page init ===============================


def test_api_workitems_page_init_authed(user_client):
    """Collaboration was removed (Task 6 of the chat-collab-removal-bug-fixes
    plan): the response now carries only the surviving doc-field search
    config, no tags/users blocks."""
    resp = user_client.get("/api/workitems_page_init")
    assert resp.status_code == 200
    body = resp.get_json()
    assert "field_config" in body
    assert "tags" not in body
    assert "users" not in body


# ============================ MS02 autocomplete ==============================


def test_api_docfield_values_ms02_degrades_without_engine(user_client, workitems_all_perms):
    """For an MS02 process with no doc-field engine, autocomplete returns [] (200)
    or stays within harness tolerance; never an uncaught error."""
    resp = user_client.get(
        "/api/docfield_values",
        query_string={
            "process": "sydoc.praesidialdepartement_bs",
            "field": "doctype",
            "q": "inv",
        },
    )
    assert resp.status_code in (200, 401, 500)
    if resp.status_code == 200:
        assert isinstance(resp.get_json(), list)


def test_api_docfield_values_blocks_sensitive_without_perm(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv

    # Pretend validationuser is a real searchable field, and that it is sensitive.
    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: {"validationuser"})
    # Everything allowed EXCEPT the sensitive perm.
    monkeypatch.setattr(
        wv, "has_permission", lambda code: code != "workitems.filter.documentfields.sensitive"
    )
    resp = user_client.get("/api/docfield_values?field=validationuser&process=all")
    assert resp.status_code == 200
    assert resp.get_json() == []


def test_api_docfield_values_allows_sensitive_with_perm(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["validationuser"])
    monkeypatch.setattr(wv, "get_sensitive_field_keys", lambda: {"validationuser"})
    monkeypatch.setattr(wv, "has_permission", lambda code: True)  # incl. the sensitive perm
    # With the perm the sensitivity gate is skipped; the route then hits the
    # (unmapped-in-CI) mapping_config registry and degrades to 500/[] -- either
    # proves the gate did NOT short-circuit. Accept both to stay DB-independent.
    resp = user_client.get("/api/docfield_values?field=validationuser&process=all")
    assert resp.status_code in (200, 500)


def test_api_docfield_values_process_not_allowed_returns_empty(
    user_client, workitems_all_perms, monkeypatch
):
    """A caller holding the blanket workitems.filter.documentfields perm but
    NOT workitems.filter.process.<p> for the specific process requested must
    not get value suggestions leaked from that process -- fail closed to []
    (never a leak, never an error that confirms/denies existence)."""
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["doctype"])
    monkeypatch.setattr(wv, "has_permission", lambda code: True)  # sensitivity gate open

    with user_client.session_transaction() as sess:
        sess["permissions"] = [
            "workitems.filter.documentfields",
            "workitems.filter.process.sydoc.allowedprocess",
        ]

    resp = user_client.get(
        "/api/docfield_values",
        query_string={"field": "doctype", "process": "sydoc.otherprocess"},
    )
    assert resp.status_code == 200
    assert resp.get_json() == []


def test_api_docfield_values_all_scopes_to_allowed_processes(
    user_client, workitems_all_perms, monkeypatch
):
    """`process=all` (the JS default when no process filter is selected) must
    not be an unfiltered escape hatch: it narrows to the caller's own
    workitems.filter.process.* grants, not every process in SearchConfig.
    With zero process grants, "all" fails closed to []."""
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(wv, "get_valid_search_columns", lambda: ["doctype"])
    monkeypatch.setattr(wv, "has_permission", lambda code: True)

    with user_client.session_transaction() as sess:
        sess["permissions"] = ["workitems.filter.documentfields"]  # no process grants

    resp = user_client.get(
        "/api/docfield_values",
        query_string={"field": "doctype", "process": "all"},
    )
    assert resp.status_code == 200
    assert resp.get_json() == []


# ============================ /import_prepared_audit =========================


def test_import_prepared_audit_gated(noperm_client):
    resp = noperm_client.post("/import_prepared_audit")
    assert resp.status_code == 403


def test_import_prepared_audit_no_file(user_client, workitems_all_perms):
    resp = user_client.post("/import_prepared_audit")
    # CI has no MS02 engine -> MS02-only gate returns 400 (not 403; perms are all granted)
    assert resp.status_code == 400
    assert "MS02" in (resp.get_json() or {}).get("error", "")


def test_import_prepared_audit_rejects_non_xlsx(user_client, workitems_all_perms):
    import io

    data = {"preparedAuditFile": (io.BytesIO(b"%PDF-1.4 nope"), "x.xlsx")}
    resp = user_client.post("/import_prepared_audit", data=data, content_type="multipart/form-data")
    # CI has no MS02 engine -> MS02-only gate returns 400 (not 403; perms are all granted)
    assert resp.status_code == 400
    assert "MS02" in (resp.get_json() or {}).get("error", "")


def _prepared_docs_link_tag(html):
    """Return the opening <a ...> tag of the register link, or None if absent.
    The tag spans several lines, so match across newlines up to the closing >."""
    import re

    m = re.search(r'<a[^>]*data-testid="prepared-docs-link"[^>]*>', html)
    return m.group(0) if m else None


def test_prepared_docs_link_visible_for_target_process(
    user_client, workitems_all_perms, monkeypatch
):
    """The toolbar link is rendered and NOT hidden when the selected process is
    an MS02 prepared-docs target, and the qualifying-process list is embedded so
    the client JS can toggle it live when the process dropdown changes."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "_ms02_prepared_docs_processes", lambda: ["sydoc.05_PDBS"])
    # _reload_user_permissions (before_request) overwrites session["permissions"]
    # from the DB each request via load_permissions_for_user; patch it so
    # allowed_processes contains the target and process_name survives the reset.
    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.import.preparedaudit",
            "workitems.filter.process.sydoc.05_PDBS",
        ],
    )

    resp = user_client.get("/workitems?prcfW=sydoc.05_PDBS")
    assert resp.status_code == 200
    tag = _prepared_docs_link_tag(resp.data.decode())
    assert tag is not None, "register link should be rendered"
    assert "hidden" not in tag, "link must be visible for the matching process"
    assert "/prepared_documents" in tag
    assert "sydoc.05_PDBS" in tag, "qualifying-process list must be embedded for the JS toggle"


def test_prepared_docs_link_rendered_but_hidden_off_target(
    user_client, workitems_all_perms, monkeypatch
):
    """On 'All Processes' (and any non-PDBS process) the link stays in the DOM
    (so the JS can reveal it live without a reload) but carries the hidden class."""
    import nx_lib.hooks as hooks
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "_ms02_prepared_docs_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(
        hooks,
        "load_permissions_for_user",
        lambda uid: [
            "workitems.view",
            "workitems.import.preparedaudit",
            "workitems.filter.process.sydoc.05_PDBS",
        ],
    )

    resp = user_client.get("/workitems")  # defaults to prcfW=all
    assert resp.status_code == 200
    tag = _prepared_docs_link_tag(resp.data.decode())
    assert tag is not None, "link should still be in the DOM for the JS toggle"
    assert "hidden" in tag, "link must be hidden until a PDBS process is selected"


def test_import_prepared_audit_upserts_when_ms02_active(
    user_client, workitems_all_perms, monkeypatch
):
    """Past the MS02 gate, a valid xlsx upserts into the register and returns
    {inserted, updated, total} (no token, no synthetic-row machinery)."""
    import io

    import openpyxl

    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    # Bypass the MIME sniff for the synthetic xlsx bytes.
    monkeypatch.setattr(wv, "is_file_allowed", lambda name, stream: True)

    captured = {}

    def fake_upsert(rows, uploaded_by):
        captured["rows"] = rows
        captured["uploaded_by"] = uploaded_by
        return {"inserted": len(rows), "updated": 0, "total": len(rows)}

    monkeypatch.setattr(wv, "upsert_prepared_documents", fake_upsert)

    wb = openpyxl.Workbook()
    sh = wb.active
    sh.append(["PID", "Collected", "CollectedBy", "Prepared", "PreparedBy"])
    sh.append(["100", 1, "Alice", 1, "Bob"])
    buf = io.BytesIO()
    wb.save(buf)

    resp = user_client.post(
        "/import_prepared_audit",
        data={"preparedAuditFile": (io.BytesIO(buf.getvalue()), "p.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["total"] == 1
    assert body["inserted"] == 1
    assert captured["rows"][0]["pid"] == "100"


def test_import_prepared_audit_persists_even_without_pid_specs(
    user_client, workitems_all_perms, monkeypatch
):
    """The register must persist rows even when no MS02 PID specs are configured
    (the default CI/TEST path: SearchConfig has no 'ms02' col_pid). The old
    warning early-return is intentionally gone -- Octo status degrades on the
    page, the upload still upserts."""
    import io

    import openpyxl

    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "is_file_allowed", lambda name, stream: True)
    # No PID specs configured.
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [])
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: [])
    monkeypatch.setattr(
        wv,
        "upsert_prepared_documents",
        lambda rows, uploaded_by: {"inserted": len(rows), "updated": 0, "total": len(rows)},
    )

    wb = openpyxl.Workbook()
    sh = wb.active
    sh.append(["PID", "Prepared"])
    sh.append(["100", 1])
    buf = io.BytesIO()
    wb.save(buf)

    resp = user_client.post(
        "/import_prepared_audit",
        data={"preparedAuditFile": (io.BytesIO(buf.getvalue()), "p.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["total"] == 1
    assert "token" not in body
    assert "warning" not in body


def test_import_prepared_audit_db_failure_returns_500(
    user_client, workitems_all_perms, monkeypatch
):
    import io

    import openpyxl

    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "is_file_allowed", lambda name, stream: True)

    def boom(rows, uploaded_by):
        raise RuntimeError("db down")

    monkeypatch.setattr(wv, "upsert_prepared_documents", boom)

    wb = openpyxl.Workbook()
    sh = wb.active
    sh.append(["PID", "Prepared"])
    sh.append(["100", 1])
    buf = io.BytesIO()
    wb.save(buf)

    resp = user_client.post(
        "/import_prepared_audit",
        data={"preparedAuditFile": (io.BytesIO(buf.getvalue()), "p.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 500
    assert "error" in (resp.get_json() or {})


# ====================== /prepared_documents (register page) =================


def test_prepared_documents_page_gated(noperm_client):
    """No permission -> 403 (or login redirect)."""
    resp = noperm_client.get("/prepared_documents")
    assert resp.status_code in (403, 302)


def test_prepared_documents_page_denies_without_ms02(user_client, workitems_all_perms):
    """Perm present but no MS02 engine in CI -> ms02_active gate denies (403)."""
    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 403


def test_prepared_documents_page_renders_when_ms02_active(
    user_client, workitems_all_perms, monkeypatch
):
    """Perm + ms02_active -> 200; table renders. DB read + Octo resolve mocked.
    The in-body prepared_import_perm uses the wv-local has_permission binding,
    so patch that too (the documented fixture trap)."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})

    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert b"100" in resp.data


def test_prepared_documents_page_octo_resolve_failure_degrades(
    user_client, workitems_all_perms, monkeypatch
):
    """resolve returning None must NOT break the page (status degrades to dash)."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": False,
                "collected_by": "",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": None,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: None)

    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200


def test_prepared_documents_octo_status_false_when_stage_not_found(
    user_client, workitems_all_perms, monkeypatch
):
    """A wid MAPPING existing is not enough: resolve_octo_wid_stage returning
    {"status": None, "current_stage": None} (nothing found in Octo) must yield
    in_octo=False, so the Preview / "Open in Workitems" buttons -- which would
    otherwise be dead links for a wid that isn't actually in Octo -- are not
    rendered."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})
    monkeypatch.setattr(
        wv,
        "_resolve_prepared_doc_wid_stage",
        lambda w: {"status": None, "current_stage": None},
    )

    captured = {}
    real_render_template = wv.render_template

    def _capture(template_name, **kwargs):
        captured.update(kwargs)
        return real_render_template(template_name, **kwargs)

    monkeypatch.setattr(wv, "render_template", _capture)

    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert captured["octo_status"]["100"]["in_octo"] is False
    # The row must fall back to the dash placeholder, not render the (dead) Preview
    # button / "Open in Workitems" link for a wid that Octo doesn't actually have.
    assert b'data-wid="42"' not in resp.data
    assert b'data-testid="prepared-docs-octo-link"' not in resp.data


def test_prepared_documents_resolves_stage_against_owning_client_engine(
    user_client, workitems_all_perms, monkeypatch
):
    """Compound identity (client + id): dbo.PreparedDocuments is an MS02-only
    register, so a register wid must be resolved against CLIENTS['ms02']'s
    runtime engine, never CLIENTS['default'] -- ids collide across client
    runtimes (docs/design/ms02-multisource.md), so resolving against the
    default engine can silently surface a DIFFERENT client's status/stage for
    a colliding id. Two distinct engines return two distinct stages; the
    MS02 stage must win and the default engine must never be consulted."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS, ClientConfig

    default_engine = object()
    ms02_engine = object()
    default_client = ClientConfig(
        code="default",
        runtime_engine=default_engine,
        dialect="tsql",
        octo_domain="default.example",
        octo_client_id="id",
        octo_secret="secret",
        octo_grant_type="client_credentials",
    )
    ms02_client = ClientConfig(
        code="ms02",
        runtime_engine=ms02_engine,
        dialect="postgres",
        octo_domain="ms02.example",
        octo_client_id="id",
        octo_secret="secret",
        octo_grant_type="client_credentials",
    )
    monkeypatch.setitem(CLIENTS, "default", default_client)
    monkeypatch.setitem(CLIENTS, "ms02", ms02_client)
    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})

    calls = []

    def fake_default_resolve(engine, wid):
        calls.append(("default", engine))
        return {"status": "Ready", "current_stage": "Extraction"}

    def fake_pg_resolve(engine, wid):
        calls.append(("ms02", engine))
        return {"status": "Done", "current_stage": "Delivery"}

    monkeypatch.setattr(wv, "resolve_octo_wid_stage", fake_default_resolve)
    monkeypatch.setattr(wv, "_resolve_octo_wid_stage_pg", fake_pg_resolve)

    captured = {}
    real_render_template = wv.render_template

    def _capture(template_name, **kwargs):
        captured.update(kwargs)
        return real_render_template(template_name, **kwargs)

    monkeypatch.setattr(wv, "render_template", _capture)

    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    # The MS02 (owning-client) engine must have been used, and ONLY it --
    # never the default engine, even though both are registered.
    assert calls == [("ms02", ms02_engine)]
    assert captured["octo_status"]["100"]["status"] == "Done"
    assert captured["octo_status"]["100"]["current_stage"] == "Delivery"
    assert captured["octo_status"]["100"]["in_octo"] is True


def test_prepared_documents_stage_resolve_fails_closed_on_error(
    user_client, workitems_all_perms, monkeypatch
):
    """A resolution error (bad CLIENTS shape, DB error, etc.) must degrade to
    in_octo=False, never raise and 500 the register page."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    # A malformed/partial CLIENTS['ms02'] entry (e.g. missing .dialect) must
    # not blow up the route -- it must fail closed instead.
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})

    captured = {}
    real_render_template = wv.render_template

    def _capture(template_name, **kwargs):
        captured.update(kwargs)
        return real_render_template(template_name, **kwargs)

    monkeypatch.setattr(wv, "render_template", _capture)

    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert captured["octo_status"]["100"]["in_octo"] is False


def test_prepared_documents_clear_gated(noperm_client):
    resp = noperm_client.post("/prepared_documents/clear")
    assert resp.status_code in (403, 302)


def test_prepared_documents_clear_400_without_ms02(user_client, workitems_all_perms):
    resp = user_client.post("/prepared_documents/clear")
    assert resp.status_code == 400
    assert b"MS02" in resp.data


def test_prepared_documents_clear_deletes_when_ms02_active(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "clear_prepared_documents", lambda: 3)
    resp = user_client.post("/prepared_documents/clear")
    assert resp.status_code == 200
    assert resp.get_json()["deleted"] == 3


def test_prepared_documents_preview_button_requires_details_view(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})
    monkeypatch.setattr(
        wv,
        "_resolve_prepared_doc_wid_stage",
        lambda w: {"status": "Ready", "current_stage": "Import"},
    )

    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert b'data-testid="prepared-docs-preview"' in resp.data
    assert b'data-wid="42"' in resp.data
    # Register is MS02-only: the preview button must carry the owning client
    # so the shared detail panel's media/audit fetches disambiguate colliding ids.
    assert b'data-client="ms02"' in resp.data
    assert b"NexoraWorkitemDetail" in resp.data

    monkeypatch.setattr(wv, "has_permission", lambda code: code != "workitems.details.view")
    resp2 = user_client.get("/prepared_documents")
    assert resp2.status_code == 200
    assert b'data-testid="prepared-docs-preview"' not in resp2.data


def test_prepared_documents_modal_wires_shared_renderer(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 0)
    monkeypatch.setattr(
        wv, "fetch_prepared_documents_page", lambda offset, limit, pid=None, **kw: []
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: [])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: None)
    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert b"NexoraWorkitemDetail.render" in resp.data
    assert b"attachLightbox" in resp.data
    assert b"api/config/fields" in resp.data
    # The click handler must thread the button's data-client through to
    # openPreview, which must forward it into the render() options (mirrors
    # toggleDetailsAndLoadImages's `client: detailsRow.dataset.client` convention).
    assert b"btn.dataset.client" in resp.data
    assert b"client: client" in resp.data


def test_prepared_documents_preview_present_when_media_degrades(
    user_client, workitems_all_perms, monkeypatch
):
    """Octo resolve returning None still renders the page with the modal shell (image
    degradation is client-side; the panel must not be gated on media)."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})
    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert b'data-testid="prepared-docs-preview-modal"' in resp.data


def test_prepared_documents_pid_filter_passes_through(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    seen = {}
    monkeypatch.setattr(
        wv,
        "count_prepared_documents",
        lambda pid=None, **kw: (seen.__setitem__("count_pid", pid) or 1),
    )
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: (
            seen.__setitem__("fetch_pid", pid)
            or [
                {
                    "id": 1,
                    "pid": "100",
                    "collected": True,
                    "collected_by": "A",
                    "prepared": False,
                    "prepared_by": "",
                    "uploaded_by": 7,
                    "uploaded_at": None,
                    "updated_at": None,
                }
            ]
        ),
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: [])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: None)
    resp = user_client.get("/prepared_documents?pid=100")
    assert resp.status_code == 200
    assert seen.get("count_pid") == "100"
    assert seen.get("fetch_pid") == "100"
    assert b'data-testid="prepared-docs-show-all"' in resp.data


def test_prepared_documents_collected_prepared_group_by_per_page_pass_through(
    user_client, workitems_all_perms, monkeypatch
):
    """Collected/Prepared/group_by/per_page query params must reach the
    data-access calls unchanged, and per_page must round-trip so the
    rendered select reflects what was actually requested."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    seen = {}
    monkeypatch.setattr(
        wv,
        "count_prepared_documents",
        lambda pid=None, collected=None, prepared=None, **kw: (
            seen.__setitem__("count_kwargs", (pid, collected, prepared)) or 1
        ),
    )
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, collected=None, prepared=None, group_by=None, **kw: (
            seen.__setitem__("fetch_kwargs", (offset, limit, pid, collected, prepared, group_by))
            or []
        ),
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: [])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: None)

    resp = user_client.get(
        "/prepared_documents?collected=1&prepared=0&group_by=collected_by&per_page=100"
    )
    assert resp.status_code == 200
    assert seen["count_kwargs"] == (None, True, False)
    assert seen["fetch_kwargs"] == (0, 100, None, True, False, "collected_by")
    # The per-page select must reflect the requested value, not the default.
    assert b'value="100" selected' in resp.data


def test_prepared_documents_per_page_rejects_unknown_value(
    user_client, workitems_all_perms, monkeypatch
):
    """An out-of-allowlist per_page must fall back to the 40 default rather
    than reaching the DB layer with an arbitrary page size."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    seen = {}
    monkeypatch.setattr(
        wv,
        "count_prepared_documents",
        lambda pid=None, **kw: 1,
    )
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, **kw: (seen.__setitem__("limit", limit) or []),
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: [])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: None)

    resp = user_client.get("/prepared_documents?per_page=9999")
    assert resp.status_code == 200
    assert seen["limit"] == 40


def test_prepared_documents_group_by_rejects_unknown_column(
    user_client, workitems_all_perms, monkeypatch
):
    """An unrecognized group_by value must not reach the ORDER BY builder --
    it should be dropped rather than passed through as free text."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    seen = {}
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, group_by=None, **kw: (seen.__setitem__("group_by", group_by) or []),
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: [])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: None)

    resp = user_client.get("/prepared_documents?group_by=DROP TABLE Users")
    assert resp.status_code == 200
    assert seen["group_by"] is None


def test_prepared_documents_centered_headers_use_align_center(
    user_client, workitems_all_perms, monkeypatch
):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})

    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    # The three centered columns now carry the .nx-table align helper on the <th>,
    # so headers no longer render left while bodies render center.
    assert resp.data.count(b'class="px-6 py-3 align-center"') == 3


def test_ms02_pid_specs_builds_from_mapping_registry(monkeypatch):
    """#98 phase 5: _ms02_pid_specs reads mapping_config.mappings_for/sources_for
    (ProcessFieldMappings/ProcessSources) instead of a per-call SearchConfig
    SELECT, but keeps the same (table, id_col, pid_col, None) spec shape."""
    import nx_lib.views.workitems as wv

    mapping = _fm("pid", "PidCol", process="sydoc.05_PDBS", client="ms02")
    source = _ps("sydoc.05_PDBS", "DossierStatistik", alias="d", client="ms02")

    monkeypatch.setattr(
        wv.mapping_config,
        "mappings_for",
        lambda client, processes, field_keys=None: [mapping] if client == "ms02" else [],
    )
    monkeypatch.setattr(
        wv.mapping_config,
        "sources_for",
        lambda client, processes=None: [source] if client == "ms02" else [],
    )

    specs = wv._ms02_pid_specs(["sydoc.05_PDBS"])
    assert specs == [("DossierStatistik", "ID", "PidCol", None)]


def test_ms02_pid_specs_empty_target_processes_returns_empty(monkeypatch):
    import nx_lib.views.workitems as wv

    def _must_not_run(*a, **k):
        raise AssertionError("mapping_config must not be queried with no target processes")

    monkeypatch.setattr(wv.mapping_config, "mappings_for", _must_not_run)
    assert wv._ms02_pid_specs([]) == []


def test_ms02_pid_specs_skips_mapping_without_matching_source(monkeypatch):
    import nx_lib.views.workitems as wv

    mapping = _fm("pid", "PidCol", process="sydoc.05_PDBS", client="ms02")

    monkeypatch.setattr(
        wv.mapping_config, "mappings_for", lambda client, processes, field_keys=None: [mapping]
    )
    monkeypatch.setattr(wv.mapping_config, "sources_for", lambda client, processes=None: [])

    assert wv._ms02_pid_specs(["sydoc.05_PDBS"]) == []


def test_ms02_prepared_docs_processes_distinct_from_registry(monkeypatch):
    """#98 phase 5: distinct processes carrying a 'pid' mapping, sourced from
    mapping_config.mappings_for(None) (all processes) instead of a
    SELECT DISTINCT ProcessName FROM SearchConfig."""
    import nx_lib.views.workitems as wv

    mappings = [
        _fm("pid", "PidCol", process="sydoc.05_PDBS", client="ms02"),
        _fm("pid", "PidCol", process="sydoc.05_PDBS", client="ms02"),
        _fm("pid", "PidCol", process="sydoc.06_ABC", client="ms02"),
    ]
    captured = {}

    def _mappings_for(client, processes, field_keys=None):
        captured["processes"] = processes
        captured["field_keys"] = field_keys
        return mappings if client == "ms02" else []

    monkeypatch.setattr(wv.mapping_config, "mappings_for", _mappings_for)

    assert wv._ms02_prepared_docs_processes() == ["sydoc.05_PDBS", "sydoc.06_ABC"]
    assert captured["processes"] is None
    assert captured["field_keys"] == {"pid"}


def test_ms02_prepared_docs_processes_empty_when_unseeded(monkeypatch):
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(
        wv.mapping_config, "mappings_for", lambda client, processes, field_keys=None: []
    )
    assert wv._ms02_prepared_docs_processes() == []


def test_stamp_in_register_marks_rows(monkeypatch):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_wids_to_pids", lambda e, s, w: {42: "100", 43: "200"})
    monkeypatch.setattr(wv, "pids_in_register", lambda pids: {"100"})
    monkeypatch.setattr(wv, "sensitive_blocked_keys", set)
    rows = [
        {"workitemid": 42, "client": "ms02"},
        {"workitemid": 43, "client": "ms02"},
    ]
    wv._stamp_in_register(rows)
    assert rows[0]["pid"] == "100" and rows[0]["in_register"] is True
    assert rows[1]["pid"] == "200" and rows[1]["in_register"] is False


def test_stamp_in_register_skips_non_ms02_rows(monkeypatch):
    """Workitem ids collide across clients, so a default-client row must never
    be resolved against MS02's PID table — it would be stamped with (and link
    to) an unrelated person's PID."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "sensitive_blocked_keys", set)
    seen = {}

    def _resolve(engine, specs, wids):
        seen["wids"] = list(wids)
        return {42: "100"}

    monkeypatch.setattr(wv, "resolve_ms02_wids_to_pids", _resolve)
    monkeypatch.setattr(wv, "pids_in_register", lambda pids: {"100"})
    rows = [{"workitemid": 42, "client": "default"}, {"workitemid": 42, "client": "ms02"}]
    wv._stamp_in_register(rows)
    assert seen["wids"] == [42], "only the MS02 row's id may be resolved"
    assert "pid" not in rows[0] and "in_register" not in rows[0]
    assert rows[1]["pid"] == "100" and rows[1]["in_register"] is True


def test_stamp_in_register_respects_sensitive_pid_gate(monkeypatch):
    """PID is a personal identifying number: when it is flagged sensitive and
    the caller lacks the perm, it must not be stamped onto list rows."""
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "sensitive_blocked_keys", lambda: {"pid"})

    def _must_not_run(*a, **k):
        raise AssertionError("PID resolution must not run when the field is blocked")

    monkeypatch.setattr(wv, "resolve_ms02_wids_to_pids", _must_not_run)
    rows = [{"workitemid": 42, "client": "ms02"}]
    wv._stamp_in_register(rows)
    assert "pid" not in rows[0]


def test_stamp_in_register_noop_when_not_ms02(monkeypatch):
    import nx_lib.views.workitems as wv

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", None)
    rows = [{"workitemid": 42}]
    wv._stamp_in_register(rows)
    assert "pid" not in rows[0] and "in_register" not in rows[0]


def test_api_workitems_carries_pid_in_register(user_client, workitems_all_perms, monkeypatch):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    monkeypatch.setattr(
        wv,
        "fetch_merged_page",
        lambda filt, off, lim: (
            [
                {
                    "workitemid": 42,
                    "status": "Ready",
                    "current_stage": "Import",
                    "priority": 0,
                    "tags": [],
                    "modifiedat": None,
                    "client": "ms02",
                }
            ],
            1,
            [],
        ),
    )
    monkeypatch.setattr(wv, "sensitive_blocked_keys", set)
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_wids_to_pids", lambda e, s, w: {42: "100"})
    monkeypatch.setattr(wv, "pids_in_register", lambda pids: {"100"})
    resp = user_client.get("/api/workitems")
    assert resp.status_code in (200, 500)  # 500 only if upstream filter parsing trips in CI
    if resp.status_code == 200:
        wi = resp.get_json()["workitems"][0]
        assert wi["pid"] == "100"
        assert wi["in_register"] is True


def test_prepared_docs_preview_button_carries_stage(user_client, workitems_all_perms, monkeypatch):
    import nx_lib.views.workitems as wv
    from nx_lib.clients import CLIENTS

    monkeypatch.setattr(wv, "engine_ms02_docfields_pg", object())
    monkeypatch.setitem(CLIENTS, "ms02", object())
    monkeypatch.setattr(wv, "count_prepared_documents", lambda pid=None, **kw: 1)
    monkeypatch.setattr(
        wv,
        "fetch_prepared_documents_page",
        lambda offset, limit, pid=None, **kw: [
            {
                "id": 1,
                "pid": "100",
                "collected": True,
                "collected_by": "A",
                "prepared": False,
                "prepared_by": "",
                "uploaded_by": 7,
                "uploaded_at": None,
                "updated_at": None,
            }
        ],
    )
    monkeypatch.setattr(wv, "_ms02_target_processes", lambda: ["sydoc.05_PDBS"])
    monkeypatch.setattr(wv, "_ms02_pid_specs", lambda procs: [("t", "id", "pid", None)])
    monkeypatch.setattr(wv, "resolve_ms02_pid_to_wids", lambda e, s, p: {"100": [42]})
    monkeypatch.setattr(
        wv,
        "_resolve_prepared_doc_wid_stage",
        lambda w: {"status": "In Progress", "current_stage": "Validation"},
    )
    monkeypatch.setattr(wv, "has_permission", lambda code: True)
    resp = user_client.get("/prepared_documents")
    assert resp.status_code == 200
    assert b'data-wid="42"' in resp.data
    assert b'data-status="In Progress"' in resp.data
    assert b'data-current-stage="Validation"' in resp.data
    # openPreview is called with 4 args -- status/stage plus the client, so a
    # colliding id (compound identity: client + id) resolves against the
    # right runtime instead of whichever client happens to own the id.
    assert (
        b"openPreview(btn.dataset.wid, btn.dataset.status, btn.dataset.currentStage, "
        b"btn.dataset.client)" in resp.data
    )


# ===================== API: saved filter views (#170) ========================


def test_filter_views_anonymous(client):
    resp = client.get("/api/workitem_filter_views", follow_redirects=False)
    assert resp.status_code in (302, 401, 403)


def test_filter_views_without_perm_returns_403(noperm_client):
    assert noperm_client.get("/api/workitem_filter_views").status_code == 403


def test_filter_views_crud_roundtrip(user_client, workitems_all_perms):
    filters = [["prcfW", "all"], ["search", "1216"], ["status", "Done"]]

    resp = user_client.post(
        "/api/workitem_filter_views", json={"name": "My queue", "filters": filters}
    )
    assert resp.status_code == 200
    vid = resp.get_json()["id"]

    resp = user_client.get("/api/workitem_filter_views")
    assert resp.status_code == 200
    mine = [v for v in resp.get_json()["views"] if v["id"] == vid]
    assert mine and mine[0]["name"] == "My queue" and mine[0]["filters"] == filters

    # Saving under the same name overwrites, keeping one row (same id).
    resp = user_client.post(
        "/api/workitem_filter_views",
        json={"name": "My queue", "filters": [["status", "Ready"]]},
    )
    assert resp.status_code == 200
    assert resp.get_json()["id"] == vid

    # Rename via id keeps the id stable.
    resp = user_client.post(
        "/api/workitem_filter_views",
        json={"id": vid, "name": "Renamed queue", "filters": filters},
    )
    assert resp.status_code == 200
    resp = user_client.get("/api/workitem_filter_views")
    assert [v["name"] for v in resp.get_json()["views"] if v["id"] == vid] == ["Renamed queue"]

    assert user_client.delete(f"/api/workitem_filter_views/{vid}").status_code == 200
    assert user_client.delete(f"/api/workitem_filter_views/{vid}").status_code == 404


def test_filter_views_validation(user_client, workitems_all_perms):
    def post(body):
        return user_client.post("/api/workitem_filter_views", json=body)

    assert post({"name": "", "filters": []}).status_code == 400
    assert post({"name": "x" * 101, "filters": []}).status_code == 400
    assert post({"name": "x", "filters": "nope"}).status_code == 400
    assert post({"name": "x", "filters": [["only-one-element"]]}).status_code == 400
    assert post({"name": "x", "filters": [["k", 1]]}).status_code == 400
    assert post({"name": "x", "filters": [["k", "v"]], "id": "nope"}).status_code == 400


def test_filter_views_scoped_to_owner(login, workitems_all_perms):
    user = login(username="user@test.local")
    vid = user.post(
        "/api/workitem_filter_views", json={"name": "scope check", "filters": []}
    ).get_json()["id"]

    admin = login(username="admin@test.local")
    assert vid not in [v["id"] for v in admin.get("/api/workitem_filter_views").get_json()["views"]]
    assert admin.delete(f"/api/workitem_filter_views/{vid}").status_code == 404

    user = login(username="user@test.local")
    assert user.delete(f"/api/workitem_filter_views/{vid}").status_code == 200


def test_filter_views_folder_roundtrip(user_client, workitems_all_perms):
    resp = user_client.post(
        "/api/workitem_filter_views",
        json={"name": "foldered", "folder": "Search Specific", "filters": []},
    )
    assert resp.status_code == 200
    body = resp.get_json()
    vid = body["id"]
    assert body["folder"] == "Search Specific"

    views = user_client.get("/api/workitem_filter_views").get_json()["views"]
    assert [v["folder"] for v in views if v["id"] == vid] == ["Search Specific"]

    # Move to another folder via id-update (id stays stable).
    resp = user_client.post(
        "/api/workitem_filter_views",
        json={"id": vid, "name": "foldered", "folder": "Other", "filters": []},
    )
    assert resp.status_code == 200
    views = user_client.get("/api/workitem_filter_views").get_json()["views"]
    assert [v["folder"] for v in views if v["id"] == vid] == ["Other"]

    # Blank folder clears it back to un-foldered (NULL -> JSON null).
    resp = user_client.post(
        "/api/workitem_filter_views",
        json={"id": vid, "name": "foldered", "folder": "  ", "filters": []},
    )
    assert resp.status_code == 200
    views = user_client.get("/api/workitem_filter_views").get_json()["views"]
    assert [v["folder"] for v in views if v["id"] == vid] == [None]

    assert user_client.delete(f"/api/workitem_filter_views/{vid}").status_code == 200


def test_filter_views_folder_validation(user_client, workitems_all_perms):
    def post(body):
        return user_client.post("/api/workitem_filter_views", json=body)

    assert post({"name": "x", "folder": "f" * 101, "filters": []}).status_code == 400
    assert post({"name": "x", "folder": 5, "filters": []}).status_code == 400

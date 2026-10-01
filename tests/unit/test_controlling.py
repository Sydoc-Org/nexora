"""Sydoc Controlling: stream mapping, rates, invoice assignment and the month payload (#433).

nx_lib/controlling.py is pure; the inputs are built here by hand. The
document specs are checked against the Finance sources as the migrations
seeded them (test_finance.REGISTRY), so a stream that names a measure its
source does not carry fails here, not on the page.
"""

import datetime as dt
import io
import re
from decimal import Decimal
from pathlib import Path

import pytest

from nx_lib import bexio, controlling, controlling_export
from nx_lib.controlling import (
    DRAFT,
    FOREIGN,
    INVOICED,
    MISSING,
    STREAMS,
    STREAMS_BY_KEY,
    UNLINKED,
    Hours,
    Inputs,
)
from nx_lib.reporting.table_query import table_source_catalog
from tests.unit.test_finance import REGISTRY

REPO = Path(__file__).resolve().parents[2]
MIGRATION = REPO / "sql" / "_migrations" / "NexoraDB" / "0145_controlling_page.sql"

BPS_CATALOG = table_source_catalog(
    [
        {"field": "Datum", "type": "date", "grainable": True},
        {"field": "Kunde", "type": "string"},
        {"field": "Projektpaket", "type": "string"},
        {"field": "Aufgabe", "type": "string"},
        {"field": "Benutzer", "type": "string"},
        {"field": "Stunden", "type": "number"},
        {"field": "Beschreibung", "type": "string"},
    ]
)


def _rate(id_, value, start, end=None, kind=controlling.HOURLY, stream=None):
    return {
        "id": id_,
        "kind": kind,
        "stream": stream,
        "value": Decimal(str(value)),
        "from": dt.date.fromisoformat(start),
        "to": dt.date.fromisoformat(end) if end else None,
        "changedAt": None,
        "changedBy": "test",
    }


RATES = [_rate(1, 85, "2025-01-01"), _rate(2, 8.4, "2025-01-01", kind=controlling.FTE_DAY_HOURS)]


def _inv(id_, date, project, excl, total=None, status="paid", currency="CHF", contact=294):
    return {
        "id": id_,
        "nr": f"RE-{id_}",
        "title": "September 2025",
        "contactId": contact,
        "projectId": project,
        "date": date,
        "due": None,
        "status": status,
        "total": total if total is not None else round(excl * 1.081, 2),
        "excl": excl,
        "currency": currency,
    }


# --------------------------------------------------------------------------
# Stream mapping
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "customer, package, task, expected",
    [
        ("Elektro Material", "Tagesgeschäft", "Validierung", "elektro_material"),
        ("Elektro Material", "Invoice", "Support", "elektro_material"),
        ("Privera", "Tagesgeschäft Posteingang", "Validierung", "privera_posteingang"),
        ("Privera", "Posteingang", "Support-verrechenbar", "privera_posteingang"),
        ("Privera", "Tagesgeschäft Invoice", "Validierung", "privera_invoice"),
        ("Privera", "Invoice", "Change Request", "privera_invoice"),
        ("Privera", "Tagesgeschäft Neuzugänge", "Scanning", "privera_neuzugaenge"),
        # BPS books Zupfen on the Invoice package; the workbook counts it under Posteingang.
        ("Privera", "Tagesgeschäft Invoice", "Zupfen", "privera_posteingang"),
        ("Privera", "Tagesgeschäft Invoice", "Zupfen ", "privera_posteingang"),
        ("ZHAW", "Tagesgeschäft", "Validierung", "zhaw"),
        ("Aveniq", "ZHAW", "Support", None),  # DocProStar, as in the workbook
        ("BFH", "Tagesgeschäft", "Validierung", "bfh"),
        ("Aveniq", "BFH", "Support", None),
        ("Bucherer", "D365", "Validierung", "bucherer"),
        ("MediaMarkt", "Tagesgeschäft", "Scanning", "mediamarkt"),
        ("frigemo", "fAPA", "Validierung", "frigemo"),  # case-insensitive, like SQL Server
        # Not on any stream: reported as unmapped, never dropped.
        ("Aveniq", "HAP", "Support", None),
        ("Privera", "Allgemein", "Support", None),
        ("sydoc intern", "sydoc", "Diverses", None),
    ],
)
def test_stream_of(customer, package, task, expected):
    assert controlling.stream_of(customer, package, task) == expected


def test_every_stream_is_unique_and_has_documents():
    assert len(STREAMS_BY_KEY) == len(STREAMS) == 10
    assert all(s.documents is not None for s in STREAMS)


def test_customers_are_the_streams_bps_customers():
    assert controlling.customers() == sorted(
        {
            "Elektro Material",
            "CompassGroup",
            "Privera",
            "Frigemo",
            "ZHAW",
            "BFH",
            "Bucherer",
            "MediaMarkt",
        }
    )


def test_fold_hours_merges_names_and_keeps_unmapped():
    rows = [
        [dt.date(2025, 9, 1), "Privera", "Tagesgeschäft Invoice", "Zupfen", Decimal("74.81")],
        [
            dt.date(2025, 9, 1),
            "Privera",
            "Tagesgeschäft Posteingang",
            "Validierung",
            Decimal("90.62"),
        ],
        [dt.date(2025, 9, 1), "Privera", "Posteingang", "Support ", Decimal("3.83")],
        [dt.date(2025, 9, 1), "Privera", "Tagesgeschäft Posteingang", "Support", Decimal("15.18")],
        [dt.date(2025, 9, 1), "Aveniq", "HAP", "Support", Decimal("4")],
        [dt.date(2025, 9, 1), "Privera", "Tagesgeschäft Invoice", "Validierung", None],
    ]
    h = controlling.fold_hours(rows)
    assert h.tasks("2025-09", "privera_posteingang") == {
        "Zupfen": Decimal("74.81"),
        "Validierung": Decimal("90.62"),
        "Support": Decimal("19.01"),
    }
    assert h.tasks("2025-09", "privera_invoice") == {}
    assert h.unmapped["2025-09"] == {("Aveniq", "HAP", "Support"): Decimal("4")}


def test_hours_query_reads_every_month_of_the_streams_customers():
    sql, params = controlling.hours_query(
        "dbo.BPS_ProjectReportAll", BPS_CATALOG, (2025, 1), (2025, 9)
    )
    assert "2025-01-01" in params and "2025-10-01" in params
    assert "Privera" in params and "Aveniq" not in params
    assert re.search(r"GROUP BY", sql)


# --------------------------------------------------------------------------
# Documents
# --------------------------------------------------------------------------


@pytest.mark.parametrize("stream", STREAMS, ids=lambda s: s.key)
def test_documents_queries_build_against_the_seeded_sources(stream):
    doc = stream.documents
    base_object, catalog, metrics = REGISTRY[doc.finance_section.source]
    assert doc.metric in metrics, f"{doc.metric} is not a measure of {doc.finance_section.source}"
    months = [(2025, 8), (2025, 9)]
    queries = controlling.documents_queries(stream, base_object, catalog, metrics, months)
    assert queries
    # A text date cannot be bucketed: one query per month.
    if doc.period.kind == "text_month":
        assert [q.mode for q in queries] == ["month", "month"]
    else:
        assert len(queries) == 1


def test_fold_documents_reads_every_mode():
    q = [
        controlling.DocQuery("compass", "grain", None, "", []),
        controlling.DocQuery("privera_neuzugaenge", "pair", None, "", []),
        controlling.DocQuery("privera_posteingang", "month", (2025, 9), "", []),
    ]
    rows = [
        [[dt.datetime(2025, 9, 1), 120], [dt.datetime(2025, 7, 1), 99]],
        [[2025, 9, Decimal("12")]],
        [[17344]],
    ]
    out = controlling.fold_documents(q, rows, [(2025, 8), (2025, 9)])
    assert out[("2025-09", "compass")] == 120
    assert out[("2025-08", "compass")] == 0
    assert ("2025-07", "compass") not in out
    assert out[("2025-09", "privera_neuzugaenge")] == 12
    assert out[("2025-09", "privera_posteingang")] == 17344


def test_snapshot_documents_reads_figures_and_breakdown_rows():
    payload = {
        "blocks": [
            {
                "figures": [{"code": "xpert_stats_documents", "value": 9000}],
                "breakdowns": [
                    {
                        "kind": "table",
                        "dim": "Client",
                        "columns": [{"code": "xpert_stats_documents"}],
                        "rows": [
                            {"key": "ZHAW", "values": [4200]},
                            {"key": "BFH", "values": [1800]},
                        ],
                    }
                ],
            }
        ]
    }
    assert controlling.snapshot_documents(STREAMS_BY_KEY["zhaw"].documents, payload) == 4200
    assert controlling.snapshot_documents(STREAMS_BY_KEY["bfh"].documents, payload) == 1800
    em = controlling.Documents("elektro_material", "em_documents")
    assert (
        controlling.snapshot_documents(
            em, {"blocks": [{"figures": [{"code": "em_documents", "value": 5}]}]}
        )
        == 5
    )
    assert controlling.snapshot_documents(em, {"error": "down"}) is None


# --------------------------------------------------------------------------
# Rates
# --------------------------------------------------------------------------


def test_rate_for_prefers_a_stream_override_and_the_latest_start():
    rates = [
        *RATES,
        _rate(3, 90, "2026-01-01"),
        _rate(4, 70, "2025-06-01", "2025-12-01", stream="mediamarkt"),
    ]
    assert controlling.rate_for(rates, controlling.HOURLY, "compass", 2025, 9) == Decimal(85)
    assert controlling.rate_for(rates, controlling.HOURLY, "compass", 2026, 2) == Decimal(90)
    assert controlling.rate_for(rates, controlling.HOURLY, "mediamarkt", 2025, 9) == Decimal(70)
    assert controlling.rate_for(rates, controlling.HOURLY, "mediamarkt", 2025, 12) == Decimal(70)
    assert controlling.rate_for(rates, controlling.HOURLY, "mediamarkt", 2026, 1) == Decimal(90)
    assert controlling.rate_for(rates, controlling.HOURLY, None, 2024, 12) is None


@pytest.mark.parametrize(
    "args, error",
    [
        (("weekly", None, 85, "2025-01", None), "Unknown rate kind."),
        (("hourly", "nope", 85, "2025-01", None), "Unknown stream."),
        (("hourly", None, "abc", "2025-01", None), "The rate must be a number."),
        (("hourly", None, 0, "2025-01", None), "The rate must be a positive number."),
        (("hourly", None, 85, "Jan", None), "A month must read YYYY-MM."),
        (("hourly", None, 85, "2025-05", "2025-01"), "'Valid to' is before 'valid from'."),
    ],
)
def test_check_rate_rejects(args, error):
    with pytest.raises(ValueError, match=re.escape(error)):
        controlling.check_rate(*args)


def test_check_rate_cleans_values():
    assert controlling.check_rate("hourly", "", "85.005", "2025-01", "") == (
        "hourly",
        None,
        Decimal("85.01"),
        dt.date(2025, 1, 1),
        None,
    )


# --------------------------------------------------------------------------
# Invoices
# --------------------------------------------------------------------------

PROJECTS = {
    10: "privera_posteingang",
    46: "privera_posteingang",
    4: "privera_invoice",
    17: "privera_neuzugaenge",
}


def test_assign_invoices_by_project_and_billed_month():
    invoices = [
        _inv(1, "2025-10-08", 10, 26000.0),
        _inv(2, "2025-10-08", 4, 7000.0),
        _inv(3, "2025-10-08", 46, 3000.0),
        _inv(4, "2025-10-08", 99, 500.0),  # linked contact, unmapped project
        _inv(5, "2025-10-08", None, 50.0, contact=1),  # not our customer
        _inv(6, "2025-10-08", 17, 700.0, status="cancelled"),
        _inv(7, "2025-11-04", 10, 25000.0),
    ]
    by, unassigned = controlling.assign_invoices(invoices, PROJECTS, {294})
    assert [i["id"] for i in by[("2025-09", "privera_posteingang")]] == [1, 3]
    assert [i["id"] for i in by[("2025-10", "privera_posteingang")]] == [7]
    assert [i["id"] for i in by[("2025-09", "privera_invoice")]] == [2]
    assert ("2025-09", "privera_neuzugaenge") not in by
    assert [i["id"] for i in unassigned["2025-09"]] == [4]


@pytest.mark.parametrize(
    "invoices, linked, expected",
    [
        ([], False, (UNLINKED, None, None)),
        ([], True, (MISSING, None, None)),
        ([_inv(1, "2025-10-01", 10, 100.0, status="draft")], True, (DRAFT, None, None)),
        # EUR is not converted: the month is FOREIGN, out of every total.
        ([_inv(1, "2025-10-01", 10, 100.0, currency="EUR")], True, (FOREIGN, None, None)),
        (
            [_inv(1, "2025-10-01", 10, 100.0, currency="EUR"), _inv(2, "2025-10-01", 10, 5.0)],
            True,
            (FOREIGN, None, None),
        ),
        (
            [
                _inv(1, "2025-10-01", 10, 100.0, 108.1),
                _inv(2, "2025-10-01", 10, 50.0, 54.05, status="draft"),
            ],
            True,
            (INVOICED, Decimal("100.0"), Decimal("108.1")),
        ),
        (
            [_inv(1, "2025-10-01", 10, 100.0, 108.1, currency="")],
            True,
            (INVOICED, Decimal("100.0"), Decimal("108.1")),
        ),
    ],
)
def test_invoice_state(invoices, linked, expected):
    assert controlling.invoice_state(invoices, linked) == expected


def test_billed_month_follows_the_invoice_offset():
    assert controlling.billed_month("2025-10-08") == (2025, 10 - bexio.INVOICE_MONTH_OFFSET)
    assert controlling.billed_month("2026-01-04") == (2025, 12)
    assert controlling.billed_month(None) is None


# --------------------------------------------------------------------------
# The month payload
# --------------------------------------------------------------------------


def _inputs(invoices=True):
    h = Hours()
    h.by[("2025-09", "privera_posteingang")] = {
        "Validierung": Decimal("100"),
        "Zupfen": Decimal("50"),
    }
    h.by[("2025-08", "privera_posteingang")] = {"Validierung": Decimal("80")}
    h.by[("2025-09", "compass")] = {"Validierung": Decimal("10")}
    h.unmapped["2025-09"] = {("Aveniq", "HAP", "Support"): Decimal("2.5")}
    by, unassigned = controlling.assign_invoices(
        [
            _inv(1, "2025-10-08", 10, 20000.0, 21620.0),
            _inv(2, "2025-09-05", 10, 9000.0, 9729.0),
            _inv(3, "2025-10-08", 99, 10.0),
        ],
        PROJECTS | {5: "compass"},
        {294},
    )
    return Inputs(
        hours=h,
        rates=RATES,
        documents={
            ("2025-09", "privera_posteingang"): 18000,
            ("2025-08", "privera_posteingang"): 15000,
        },
        documents_live={("2025-09", "privera_posteingang"): 18100},
        invoices=by if invoices else None,
        unassigned=unassigned,
        linked_streams=frozenset(PROJECTS.values()) | {"compass"},
        bexio_error=None if invoices else "Bexio down",
    )


def test_stream_month_computes_cost_margin_and_kpis():
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, _inputs())
    assert c["hours"] == 150.0
    assert c["rate"] == 85.0
    assert c["cost"] == 12750.0
    assert c["state"] == INVOICED
    assert c["invoiced"] == 20000.0 and c["invoicedIncl"] == 21620.0
    assert c["margin"] == 7250.0
    assert c["marginPct"] == pytest.approx(0.3625)
    assert c["documents"] == 18000 and c["documentsLive"] == 18100
    assert c["kpis"]["secondsPerDocument"] == 20.0  # 100 h * 3600 / 18000
    assert c["kpis"]["documentsPerHour"] == 120.0
    assert c["kpis"]["chfPerDocument"] == pytest.approx(1.1111)
    assert list(c["tasks"]) == ["Validierung", "Zupfen"]
    assert c["flags"] == ["incomplete"]  # Privera Jan-Oct 2025: BPS lacks hours


def test_stream_month_flags_incomplete_and_missing():
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 8, _inputs())
    assert "incomplete" in c["flags"] and c["incomplete"]
    assert c["state"] == INVOICED  # the 2025-09-05 invoice bills August
    compass = controlling.stream_month(STREAMS_BY_KEY["compass"], 2025, 9, _inputs())
    assert compass["state"] == MISSING
    assert compass["margin"] is None and compass["cost"] == 850.0
    mm = controlling.stream_month(STREAMS_BY_KEY["mediamarkt"], 2025, 9, _inputs())
    assert mm["state"] == UNLINKED


def test_stream_month_without_bexio_is_an_error_not_zero():
    c = controlling.stream_month(
        STREAMS_BY_KEY["privera_posteingang"], 2025, 9, _inputs(invoices=False)
    )
    assert c["state"] == controlling.ERROR
    assert c["invoiced"] is None and c["margin"] is None


def test_stream_month_without_rate_flags_it():
    inp = _inputs()
    inp.rates = []
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, inp)
    assert c["cost"] is None and c["margin"] is None and "no_rate" in c["flags"]


def test_incomplete_covers_privera_jan_to_oct_2025_only():
    assert controlling.incomplete_note("privera_invoice", 2025, 1)
    assert controlling.incomplete_note("privera_neuzugaenge", 2025, 10)
    assert controlling.incomplete_note("privera_invoice", 2025, 11) is None
    assert controlling.incomplete_note("compass", 2025, 5) is None


def test_month_payload_totals_matrix_and_extras():
    p = controlling.month_payload(2025, 9, _inputs())
    assert p["month"] == "2025-09" and p["prevMonth"] == "2025-08"
    assert [s["key"] for s in p["streams"]] == [s.key for s in STREAMS]
    t = p["totals"]["cur"]
    assert t["hours"] == 160.0
    assert t["cost"] == 13600.0
    assert t["invoiced"] == 20000.0 and t["margin"] == 7250.0
    assert t["streamsWithMargin"] == 1 and t["streams"] == 10
    assert t["workingDays"] == 22
    assert t["fte"] == pytest.approx(160 / (8.4 * 22), abs=0.01)
    pe = next(s for s in p["streams"] if s["key"] == "privera_posteingang")
    assert pe["delta"]["hours"] == 70.0
    m = p["tasks"]
    assert m["tasks"] == ["Validierung", "Zupfen"]
    assert m["taskTotals"] == [110.0, 50.0] and m["total"] == 160.0
    assert [i["id"] for i in p["unassigned"]] == [3]
    assert p["unmapped"] == [
        {"customer": "Aveniq", "package": "HAP", "task": "Support", "hours": 2.5}
    ]


def test_month_payload_attaches_invoice_lines():
    lines = [{"kind": "line", "text": "Scanning", "amount": 1.0, "unitPrice": 1.0, "total": 1.0}]
    p = controlling.month_payload(2025, 9, _inputs(), positions={1: lines})
    pe = next(s for s in p["streams"] if s["key"] == "privera_posteingang")
    assert pe["cur"]["invoices"][0]["positions"] == lines


def test_month_payload_at_the_first_month_has_no_previous():
    p = controlling.month_payload(2025, 1, _inputs())
    assert p["prevMonth"] is None and p["totals"]["prev"] is None
    assert all(s["prev"] is None and s["delta"] is None for s in p["streams"])


def test_trend_payload_series_and_totals():
    months = [(2025, 8), (2025, 9)]
    t = controlling.trend_payload(months, _inputs())
    assert t["months"] == ["2025-08", "2025-09"]
    pe = next(s for s in t["streams"] if s["key"] == "privera_posteingang")
    assert [p["hours"] for p in pe["points"]] == [80.0, 150.0]
    assert [x["margin"] for x in t["totals"]] == [9000.0 - 6800.0, 7250.0]


def test_parse_month_never_goes_before_the_bps_history():
    assert controlling.parse_month("2024-05", dt.date(2026, 10, 1)) == (2025, 1)
    assert controlling.parse_month("2026-03", dt.date(2026, 10, 1)) == (2026, 3)


def test_working_days():
    assert controlling.working_days(2025, 9) == 22
    assert controlling.working_days(2026, 2) == 20


# --------------------------------------------------------------------------
# Export and migration
# --------------------------------------------------------------------------


def test_export_workbook_has_four_sheets_and_states():
    from openpyxl import load_workbook

    labels = {
        k: k
        for k in (
            "title overview detail tasks volumes stream task hours rate cost invoiced invoiced_incl "
            "margin margin_pct delta_margin documents previous state total invoice date position "
            "quantity unit_price amount seconds_per_doc docs_per_hour chf_per_doc fte unassigned "
            "unmapped incomplete no_rate external converted no_hours cost_unknown"
        ).split()
    }
    labels["states"] = {INVOICED: "invoiced", MISSING: "no invoice", UNLINKED: "not linked"}
    payload = controlling.month_payload(2025, 9, _inputs())
    wb = load_workbook(io.BytesIO(controlling_export.workbook(payload, labels, "September 2025")))
    assert wb.sheetnames == ["overview", "detail", "tasks", "volumes"]
    values = [c for row in wb["overview"].iter_rows(values_only=True) for c in row if c is not None]
    assert "no invoice" in values  # Compass: state, not 0
    assert 7250 in values


def test_migration_maps_every_project_to_a_known_stream():
    sql = MIGRATION.read_text(encoding="utf-8")
    pairs = re.findall(r"\((\d+),\s*N'([a-z_]+)',", sql)
    assert pairs, "0145 no longer seeds dbo.ControllingStreamProjects"
    assert {key for _pid, key in pairs} <= set(STREAMS_BY_KEY)
    assert {key for _pid, key in pairs} == set(STREAMS_BY_KEY), "a stream has no Bexio project"
    assert len({pid for pid, _key in pairs}) == len(pairs)


# --------------------------------------------------------------------------
# Reconciliation rules (0146): overrides, exclusions, external costs
# --------------------------------------------------------------------------


def test_assign_invoices_override_beats_project_and_none_excludes():
    invoices = [
        _inv(707, "2025-10-08", 10, 4837.51),  # RE-26780: Posteingang project, Invoice support
        _inv(708, "2025-10-08", 10, 2092.51),
        _inv(709, "2025-10-06", 46, 3700.51),  # Postpakete: excluded, not unassigned
    ]
    by, unassigned = controlling.assign_invoices(
        invoices, PROJECTS | {46: None}, {294}, overrides={707: "privera_invoice"}
    )
    assert [i["id"] for i in by[("2025-09", "privera_invoice")]] == [707]
    assert [i["id"] for i in by[("2025-09", "privera_posteingang")]] == [708]
    assert unassigned == {}


def test_bill_costs_by_vendor_and_bill_month():
    bills = [
        {
            "id": "a",
            "nr": "00035",
            "vendor": "Digi-Texx DIGI-TEXX Vietnam Anna Building",
            "title": "BPO Services Juli 2026",
            "date": "2026-07-31",
            "net": 6414.45,
            "currency": "CHF",
            "status": "paid",
        },
        {
            "id": "b",
            "nr": "00063",
            "vendor": "digi-texx",
            "title": "BPO Services August 2026",
            "date": "2026-08-31",
            "net": 4917.0,
            "currency": "EUR",
            "status": "booked",
        },
        {
            "id": "c",
            "nr": "00064",
            "vendor": "Digi-Texx",
            "title": "draft",
            "date": "2026-08-31",
            "net": 1.0,
            "currency": "CHF",
            "status": "draft",
        },
        {
            "id": "d",
            "nr": "00096",
            "vendor": "WWZ Energie AG",
            "title": "Internet",
            "date": "2026-08-17",
            "net": 195.1,
            "currency": "CHF",
            "status": "paid",
        },
    ]
    items = controlling.bill_costs(bills, {"digi-texx": ("privera_invoice", "Digi-Texx")})
    assert [(i["month"], i["stream"], i["amount"]) for i in items] == [
        ("2026-07", "privera_invoice", Decimal("6414.45")),
        ("2026-08", "privera_invoice", None),  # foreign: no amount, flagged
    ]
    assert items[0]["label"] == "Digi-Texx · BPO Services Juli 2026"


def test_external_costs_add_to_the_cost():
    inp = _inputs()
    inp.costs = controlling.fold_costs(
        [
            {
                "id": 1,
                "month": "2025-09",
                "stream": "privera_posteingang",
                "label": "Digi-Texx",
                "amount": Decimal("1000"),
                "source": "manual",
            }
        ]
    )
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, inp)
    assert c["hoursCost"] == 12750.0 and c["cost"] == 13750.0 and c["margin"] == 6250.0
    assert c["externalCosts"][0]["amount"] == 1000.0


def test_an_external_cost_without_amount_leaves_the_cost_unknown():
    inp = _inputs()
    inp.costs = {
        ("2025-09", "privera_posteingang"): [
            {
                "id": "b",
                "month": "2025-09",
                "stream": "privera_posteingang",
                "label": "x",
                "amount": None,
                "source": "bexio",
            }
        ]
    }
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, inp)
    assert c["cost"] is None and c["margin"] is None and "cost_unknown" in c["flags"]


def test_invoiced_without_hours_is_flagged():
    inp = _inputs()
    inp.hours = Hours()
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, inp)
    assert c["hours"] == 0 and "no_hours" in c["flags"]


@pytest.mark.parametrize(
    "args, error",
    [
        (("2025-09", "nope", "x", 1), "Unknown stream."),
        (("2025-09", "compass", "  ", 1), "Name the cost."),
        (("Sep", "compass", "x", 1), "A month must read YYYY-MM."),
        (("2025-09", "compass", "x", "abc"), "The amount must be a number."),
        (("2025-09", "compass", "x", 0), "The amount must not be 0."),
    ],
)
def test_check_cost_rejects(args, error):
    with pytest.raises(ValueError, match=re.escape(error)):
        controlling.check_cost(*args)


def test_check_cost_allows_a_credit():
    assert controlling.check_cost("2025-09", "compass", " Gutschrift  x ", "-12.345") == (
        dt.date(2025, 9, 1),
        "compass",
        "Gutschrift x",
        Decimal("-12.35"),
    )


def test_0146_excludes_only_projects_0145_mapped():
    sql45 = MIGRATION.read_text(encoding="utf-8")
    sql46 = (MIGRATION.parent / "0146_controlling_reconciliation.sql").read_text(encoding="utf-8")
    mapped = {int(pid) for pid, _key in re.findall(r"\((\d+),\s*N'([a-z_]+)',", sql45)}
    excluded = {
        int(pid) for pid in re.findall(r"\((\d+),\s*N'[^']*(?:not BFH|not revenue)'\)", sql46)
    }
    assert excluded == {6, 46, 73}
    assert excluded <= mapped


# --------------------------------------------------------------------------
# Design review (0147): foreign currency, drafts, BPS down, snapshots, rates
# --------------------------------------------------------------------------


def test_foreign_and_draft_amounts_are_shown_not_summed():
    invs = [
        _inv(1, "2025-10-06", 43, 18260.67, 19739.78, currency="EUR"),
        _inv(2, "2025-10-08", 43, 585.02, 632.4),
    ]
    assert controlling.foreign_totals(invs) == [
        {"currency": "EUR", "excl": 18260.67, "total": 19739.78}
    ]
    drafts = [_inv(3, "2025-10-01", 10, 757.9, 819.29, status="draft")]
    assert controlling.draft_totals(drafts) == (Decimal("757.9"), Decimal("819.29"))
    inp = _inputs()
    inp.invoices[("2025-09", "privera_invoice")] = drafts
    c = controlling.stream_month(STREAMS_BY_KEY["privera_invoice"], 2025, 9, inp)
    assert c["state"] == DRAFT and c["draft"] == 757.9 and c["margin"] is None


def test_bps_down_leaves_hours_unknown_not_zero():
    inp = _inputs()
    inp.hours_error = "Could not read the BPS hours."
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, inp)
    assert c["hours"] is None and c["cost"] is None and c["margin"] is None
    assert c["invoiced"] == 20000.0 and "bps_error" in c["flags"]
    p = controlling.month_payload(2025, 9, inp)
    assert p["tasks"] is None and p["hoursError"] and p["totals"]["cur"]["hours"] is None


def test_override_rate_is_flagged():
    inp = _inputs()
    inp.rates = [*RATES, _rate(9, 72, "2025-01-01", stream="privera_posteingang")]
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, inp)
    assert c["rate"] == 72.0 and "override" in c["flags"]


def test_snapshot_freezes_figures_and_flags_live_drift():
    inp = _inputs()
    snap = controlling.snapshot(2025, 9, inp)
    assert controlling.snapshot_problems(snap) == []
    frozen = snap["streams"]["privera_posteingang"]
    assert frozen["invoiced"] == 20000.0 and frozen["invoiceIds"] == [1]
    assert "invoices" not in frozen
    # Later: a rate change and a corrected invoice do not move the closed month.
    later = _inputs()
    later.rates = [_rate(1, 99, "2025-01-01"), RATES[1]]
    later.invoices[("2025-09", "privera_posteingang")][0]["excl"] = 20500.0
    later.frozen = {"2025-09": snap}
    c = controlling.stream_month(STREAMS_BY_KEY["privera_posteingang"], 2025, 9, later)
    assert c["frozen"] and c["rate"] == 85.0 and c["invoiced"] == 20000.0
    assert "moved" in c["flags"] and c["live"]["invoiced"] == 20500.0
    assert c["invoices"][0]["excl"] == 20500.0  # the lines stay live


def test_snapshot_problems_name_the_sources_down():
    inp = _inputs(invoices=False)
    inp.hours_error = "x"
    assert controlling.snapshot_problems(controlling.snapshot(2025, 9, inp)) == ["BPS", "Bexio"]


def test_rate_overlap_per_scope():
    rates = [
        _rate(1, 85, "2025-01-01", "2025-12-01"),
        _rate(2, 72, "2026-01-01", stream="mediamarkt"),
    ]
    d = dt.date
    assert controlling.rate_overlap(rates, "hourly", None, d(2025, 6, 1), None)["id"] == 1
    assert controlling.rate_overlap(rates, "hourly", None, d(2026, 1, 1), None) is None
    assert (
        controlling.rate_overlap(rates, "hourly", "mediamarkt", d(2025, 1, 1), d(2025, 12, 1))
        is None
    )
    assert controlling.rate_overlap(rates, "hourly", "mediamarkt", d(2027, 1, 1), None)["id"] == 2
    # Editing a row does not clash with itself.
    assert controlling.rate_overlap(rates, "hourly", None, d(2025, 1, 1), None, ignore_id=1) is None


def test_stream_descriptor_carries_the_block_heading():
    d = controlling.stream_descriptor(STREAMS_BY_KEY["privera_neuzugaenge"])
    assert (d["client"], d["title"], d["nav"], d["unit"]) == (
        "Privera",
        "Neuzugänge",
        "Neuzugänge",
        "dossiers",
    )
    assert controlling.stream_descriptor(STREAMS_BY_KEY["compass"])["nav"] == "Compass"


# --------------------------------------------------------------------------
# QA review (#433): edge cases that used to 500
# --------------------------------------------------------------------------


@pytest.mark.parametrize("bad", ["NaN", "Infinity", "-inf"])
def test_non_finite_numbers_are_rejected_not_500(bad):
    with pytest.raises(ValueError, match=re.escape("The rate must be a number.")):
        controlling.check_rate("hourly", None, bad, "2025-01", None)
    with pytest.raises(ValueError, match=re.escape("The amount must be a number.")):
        controlling.check_cost("2025-09", "compass", "x", bad)


def test_an_older_snapshot_missing_keys_still_serves():
    inp = _inputs()
    inp.frozen = {
        "2025-09": {"streams": {"compass": {"hours": 10.0, "state": "invoiced", "invoiced": 100.0}}}
    }
    c = controlling.stream_month(STREAMS_BY_KEY["compass"], 2025, 9, inp)
    assert c["frozen"] and c["tasks"] == {} and c["foreign"] == [] and c["documents"] is None
    t = controlling.trend_payload([(2025, 9)], inp)
    assert t["totals"][0]["hours"] is not None


def test_export_without_bps_hours_still_builds():
    labels = {
        k: k
        for k in (
            "title overview detail tasks volumes stream task hours rate cost invoiced invoiced_incl "
            "margin margin_pct delta_margin documents previous state total invoice date position "
            "quantity unit_price amount seconds_per_doc docs_per_hour chf_per_doc fte unassigned "
            "unmapped incomplete no_rate external override moved bps_error no_hours cost_unknown"
        ).split()
    }
    labels["states"] = {}
    inp = _inputs()
    inp.hours_error = "down"
    payload = controlling.month_payload(2025, 9, inp)
    assert payload["tasks"] is None
    assert controlling_export.workbook(payload, labels, "September 2025")[:2] == b"PK"

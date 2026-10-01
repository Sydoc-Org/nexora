"""The BPS hours export of Sydoc Finance: one sheet per invoice, booked and billed hours (#408)."""

import io
from decimal import Decimal

from openpyxl import load_workbook

from nx_lib import finance_export as fx
from nx_lib.finance import BILLED_FIELD, SECTIONS_BY_KEY

SPEC = SECTIONS_BY_KEY["bps"].bookings
LABELS = {
    "title": "Sydoc BPS · billable hours",
    "overview": "Overview",
    "invoice": "Invoice",
    "bookings": "Bookings",
    "bookings_n": lambda n: f"{n} bookings",
    "hours": "Hours",
    "billed": "Billed (¼ h)",
    "total": "Total",
    "date": "Date",
    "package": "Package",
    "task": "Task",
    "person": "Person",
    "comment": "Comment",
    "empty": "No billable hours",
}


def _row(customer, package, hours, day="2026-09-02", person="Ben"):
    return [day, customer, package, "Change", person, hours, "Kommentar \N{EN DASH} ä €", None]


def _payload(*rows, columns=None):
    cols = columns or [{"field": f} for f in SPEC.columns] + [{"field": BILLED_FIELD}]
    groups: dict = {}
    for r in rows:
        groups.setdefault(r[1], []).append(r)
    return {
        "bookings": {"columns": cols, "groups": [{"key": k, "rows": v} for k, v in groups.items()]}
    }


def test_privera_is_split_per_stream_in_page_order_and_the_rest_per_customer():
    payload = _payload(
        _row("Aveniq", "BFH", 0.3333),
        _row("Aveniq", "ZHAW", 1),
        _row("Privera", "Tagesgeschäft Neuzugänge", 1),
        _row("Privera", "Neuzugänge", 0.5),
        _row("Privera", "Invoice", 0.25),
        _row("Privera", "Posteingang", 0.1),
        _row("Elektro Material", "Invoice", 2),
    )
    sheets = fx.sheets(payload, SPEC)
    assert [(s.key, s.title) for s in sheets] == [
        ("elektro-material", "Elektro-Material"),
        ("privera-posteingang", "Privera Posteingang"),
        ("privera-invoice", "Privera Invoice"),
        ("privera-neuzugaenge", "Privera Neuzugänge"),
        ("aveniq", "Aveniq"),
    ]
    neu = sheets[3]
    assert len(neu.rows) == 2 and neu.hours == Decimal("1.5")
    aveniq = sheets[4]
    assert aveniq.hours == Decimal("1.3333") and aveniq.billed == Decimal("1.5")


def test_billed_hours_are_recomputed_for_a_snapshot_without_them():
    """A month closed before the rounding has no billed column: the export still rounds."""
    cols = [{"field": f} for f in SPEC.columns]
    payload = _payload(_row("Bucherer", "EasyTax", 0.4)[:-1], columns=cols)
    (sheet,) = fx.sheets(payload, SPEC)
    assert sheet.rows[0].billed == Decimal("0.5")


def test_workbook_has_an_overview_and_one_sheet_per_invoice():
    payload = _payload(_row("Privera", "Invoice", 0.3333), _row("Frigemo", "fAPA", 1))
    data = fx.workbook(fx.sheets(payload, SPEC), LABELS, "September 2026")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Overview", "Privera Invoice", "Frigemo"]
    ws = wb["Privera Invoice"]
    assert [c.value for c in ws[5]] == [
        "Date",
        "Package",
        "Task",
        "Person",
        "Comment",
        "Hours",
        "Billed (¼ h)",
    ]
    assert [c.value for c in ws[6]][5:] == [0.3333, 0.5]
    assert ws["F7"].value == "=SUM(F6:F6)"


def test_one_invoice_has_no_overview_and_an_empty_month_still_downloads():
    one = fx.workbook(fx.sheets(_payload(_row("Frigemo", "fAPA", 1)), SPEC), LABELS, "Sep")
    assert load_workbook(io.BytesIO(one)).sheetnames == ["Frigemo"]
    empty = fx.workbook([], LABELS, "Sep")
    assert load_workbook(io.BytesIO(empty)).sheetnames == ["Overview"]


def test_sheet_names_are_excel_safe_and_unique():
    taken: set = set()
    assert fx._sheet_name("A/B: [x]", taken) == "A B   x"
    long = "x" * 40
    assert fx._sheet_name(long, taken) == "x" * 31
    assert fx._sheet_name(long, taken) == "x" * 27 + " (2)"


def test_pdf_prints_unicode_comments():
    payload = _payload(_row("Privera", "Invoice", 0.3333), _row("Aveniq", "BFH", 2))
    data = fx.pdf(fx.sheets(payload, SPEC), LABELS, "September 2026")
    assert data.startswith(b"%PDF")
    assert fx.pdf([], LABELS, "September 2026").startswith(b"%PDF")


def test_zip_holds_one_file_per_invoice():
    import zipfile

    data = fx.zipped([("a.pdf", b"1"), ("b.pdf", b"2")])
    assert zipfile.ZipFile(io.BytesIO(data)).namelist() == ["a.pdf", "b.pdf"]

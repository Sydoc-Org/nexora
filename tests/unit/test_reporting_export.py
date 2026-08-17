"""Unit tests for nx_lib.reporting.export — rows → xlsx / csv bytes."""

import csv
import io
import zipfile

from openpyxl import load_workbook

from nx_lib.reporting.export import _safe_cell, rows_to_csv, rows_to_xlsx


def test_rows_to_xlsx_uses_custom_headers_and_title():
    columns = [
        {"field": "doctype", "header": "Document type"},
        {"field": "pages", "header": "Pages"},
    ]
    rows = [["Invoice", 3], ["Letter", 1]]
    data = rows_to_xlsx(columns, rows, title="Q1 report")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "Q1 report"  # title now in A1
    assert ws["A4"].value == "Document type"  # header row at row 4
    assert ws["A4"].font.bold
    assert ws["B4"].value == "Pages"
    assert ws["A5"].value == "Invoice"
    assert ws["B6"].value == 1


def test_rows_to_xlsx_falls_back_to_field_when_no_header():
    columns = [{"field": "doctype", "header": None}]
    data = rows_to_xlsx(columns, [["x"]], title="t")
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A4"].value == "doctype"


def test_rows_to_xlsx_returns_bytes():
    data = rows_to_xlsx([{"field": "a", "header": "A"}], [], title="t")
    assert isinstance(data, bytes | bytearray) and len(data) > 0


def test_rows_to_xlsx_sanitizes_illegal_title_characters():
    data = rows_to_xlsx([{"field": "a", "header": "A"}], [], title="Q1/Q2: results")
    ws = load_workbook(io.BytesIO(data)).active
    assert not any(ch in ws.title for ch in r"\/?*[]:")
    assert ws["A1"].value == "Q1/Q2: results"  # full title in A1


def test_rows_to_xlsx_neutralizes_formula_injection():
    columns = [{"field": "name", "header": "=danger"}]
    rows = [["=1+1"]]
    data = rows_to_xlsx(columns, rows, title="t")
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A4"].value == "'=danger"  # header row at row 4
    assert ws["A4"].data_type == "s"
    assert ws["A5"].value == "'=1+1"  # data starts at row 5
    assert ws["A5"].data_type == "s"


def _parse_csv(data):
    """Decode export bytes (utf-8-sig strips the BOM) and parse the rows."""
    text = data.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def test_rows_to_csv_has_bom_and_header_and_rows():
    columns = [
        {"field": "doctype", "header": "Document type"},
        {"field": "pages", "header": "Pages"},
    ]
    rows = [["Invoice", 3], ["Letter", 1]]
    data = rows_to_csv(columns, rows)
    assert isinstance(data, bytes) and data[:3] == b"\xef\xbb\xbf"  # UTF-8 BOM
    parsed = _parse_csv(data)
    assert parsed[0] == ["Document type", "Pages"]
    assert parsed[1] == ["Invoice", "3"]
    assert parsed[2] == ["Letter", "1"]


def test_rows_to_csv_uses_crlf_line_terminator():
    data = rows_to_csv([{"field": "a", "header": "A"}], [["x"]])
    assert b"A\r\n" in data
    assert b"x\r\n" in data


def test_rows_to_csv_falls_back_to_field_when_no_header():
    data = rows_to_csv([{"field": "doctype", "header": None}], [["x"]])
    assert _parse_csv(data)[0] == ["doctype"]


def test_rows_to_csv_renders_none_as_empty():
    data = rows_to_csv([{"field": "a", "header": "A"}, {"field": "b", "header": "B"}], [[None, 0]])
    assert _parse_csv(data)[1] == ["", "0"]


def test_rows_to_csv_neutralizes_formula_injection():
    columns = [{"field": "name", "header": "=danger"}]
    rows = [["=1+1"], ["-2"], ["+3"], ["@cmd"], ["safe"]]
    parsed = _parse_csv(rows_to_csv(columns, rows))
    assert parsed[0] == ["'=danger"]
    assert parsed[1] == ["'=1+1"]
    assert parsed[2] == ["'-2"]
    assert parsed[3] == ["'+3"]
    assert parsed[4] == ["'@cmd"]
    assert parsed[5] == ["safe"]


# ---------------------------------------------------------------------------
# New tests: title block, styled header, optional chart embedding
# ---------------------------------------------------------------------------


def _make_tiny_png():
    """Generate a minimal valid PNG via PIL (avoids hard-coded byte fragility)."""
    from PIL import Image as PilImage

    buf = io.BytesIO()
    PilImage.new("RGBA", (1, 1)).save(buf, "PNG")
    return buf.getvalue()


def test_xlsx_has_title_block_and_styled_header():
    data = rows_to_xlsx([{"field": "a", "header": "A"}], [[1]], title="My Report")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "My Report"
    assert ws["A1"].font.bold
    # header row is NOT at row 1 anymore
    header_cell = next(
        c for row in ws.iter_rows() for c in row if c.value == "A" and c.font and c.font.bold
    )
    assert header_cell is not None
    assert ws.freeze_panes is not None


def test_xlsx_embeds_chart_png():
    tiny_png = _make_tiny_png()
    data = rows_to_xlsx(
        [{"field": "a", "header": "A"}], [[1]], title="With Chart", chart_png=tiny_png
    )
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        media = [n for n in z.namelist() if n.startswith("xl/media/")]
    assert media, "chart image not embedded in workbook"


def test_xlsx_without_chart_has_no_media():
    data = rows_to_xlsx([{"field": "a", "header": "A"}], [[1]], title="Plain")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert not [n for n in z.namelist() if n.startswith("xl/media/")]


# ---------------------------------------------------------------------------
# bytes / control-char cells (Task 52) — openpyxl raises on both; the csv
# path would otherwise write the Python b'...' repr literally.
# ---------------------------------------------------------------------------


def test_rows_to_xlsx_decodes_bytes_cell():
    columns = [{"field": "a", "header": "A"}]
    rows = [[b"caf\xc3\xa9"]]
    data = rows_to_xlsx(columns, rows, title="t")  # must not raise
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A5"].value == "café"


def test_rows_to_xlsx_strips_control_characters():
    columns = [{"field": "a", "header": "A"}]
    rows = [["bell\x07ringer"]]
    data = rows_to_xlsx(columns, rows, title="t")  # must not raise
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A5"].value == "bellringer"


def test_rows_to_csv_decodes_bytes_cell():
    columns = [{"field": "a", "header": "A"}]
    rows = [[b"caf\xc3\xa9"]]
    data = rows_to_csv(columns, rows)  # must not raise
    assert _parse_csv(data)[1] == ["café"]
    assert b"b'" not in data  # no Python bytes-repr leaking into the file


def test_rows_to_csv_strips_control_characters():
    columns = [{"field": "a", "header": "A"}]
    rows = [["bell\x07ringer"]]
    data = rows_to_csv(columns, rows)
    assert _parse_csv(data)[1] == ["bellringer"]
    assert b"\x07" not in data


def test_safe_cell_stringifies_non_primitive_types():
    from uuid import UUID

    val = UUID("12345678-1234-5678-1234-567812345678")
    assert _safe_cell(val) == str(val)


def test_safe_cell_passes_decimal_and_date_through_unchanged():
    from datetime import date
    from decimal import Decimal

    d = Decimal("1.50")
    dt = date(2020, 1, 1)
    assert _safe_cell(d) == d
    assert _safe_cell(dt) == dt


def test_xlsx_forecast_rows_styled_italic():
    from openpyxl import load_workbook

    columns = [
        {"field": "d", "header": "Date"},
        {"field": "n", "header": "Count"},
        {"field": "__forecast", "header": "Forecast"},
    ]
    rows = [["2025-01-01", 10, ""], ["2025-02-01", 12, ""], ["2025-03-01", 14.0, "forecast"]]
    data = rows_to_xlsx(columns, rows, title="T", forecast_start=2)
    ws = load_workbook(io.BytesIO(data)).active
    # header_row is 4 without a chart; data rows follow
    # openpyxl round-trips an empty-string cell value as None, not "" — pinned
    # via ws.iter_rows() against the actual saved/reloaded workbook.
    assert ws.cell(row=5, column=3).value is None
    fc_cell = ws.cell(row=7, column=1)
    assert fc_cell.font.italic
    assert ws.cell(row=7, column=3).value == "forecast"
